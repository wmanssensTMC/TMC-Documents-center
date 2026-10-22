#!/usr/bin/env python3
"""Yellow parts of an engine draft, for the on-screen check (phase 2c).

Reuses the engine's own helpers (tmc_common: same walk as yellow_runs, render_pdf); the
engine itself is not changed. A yellow mark is only removed for a part a person confirmed
on the check screen -- never automatically (BRIEF 7.2).

  python yellow.py --engine DIR list DOCX                     -> JSON list on stdout
  python yellow.py --engine DIR apply DOCX FIXES.json OUT.docx
  python yellow.py --engine DIR preview DOCX OUT.pdf          -> PDF stamped DRAFT
  python yellow.py --engine DIR changes DRAFT EDITED SOURCE.pdf -> JSON list of changes on stdout
  python yellow.py --engine DIR view DOCX OUT.json            -> the on-screen editor's view
  python yellow.py --engine DIR edit DOCX EDITS.json OUT.docx -> the editor's changes written in
  python yellow.py --engine DIR layout PDF DOCX               -> exit 1 if text covers the letterhead or is cut off
  python yellow.py --engine DIR check DOCX FORBID.json [SHOWN.json]
                                     -> exit 1 if unsafe, names the supplier, or does not
                                        print a value of its file name (SHOWN: [[label, value]])
"""
import argparse, io, json, re, shutil, sys, tempfile
from pathlib import Path


def runs_in_order(doc):
    """Highlighted runs exactly as the engine reads them (tmc_common.iter_highlighted):
    body, tables, nested tables, text boxes, headers and footers. One walk for both the
    check screen and the engine's --finalise, so they can never disagree."""
    from tmc_common import iter_highlighted
    return list(iter_highlighted(doc))


def cmd_list(docx):
    from docx import Document
    doc = Document(str(docx))
    items = []
    for i, (r, p) in enumerate(runs_in_order(doc)):
        # Over 500 characters: too long for the on-screen editor, check it in Word (review M1).
        items.append({"id": f"y{i}", "text": r.text[:500], "context": (p.text if p is not None else r.text)[:300], "truncated": len(r.text) > 500})
    json.dump(items, sys.stdout, ensure_ascii=False)


def cmd_apply(docx, fixes_path, out):
    from docx import Document
    from tmc_common import clear_mark
    fixes = json.loads(Path(fixes_path).read_text(encoding="utf-8"))
    doc = Document(str(docx))
    found = runs_in_order(doc)
    ids = {f"y{i}" for i in range(len(found))}
    unknown = set(fixes) - ids
    if unknown:
        print(f"unknown yellow part(s): {sorted(unknown)}", file=sys.stderr)
        sys.exit(1)
    for i, (r, _p) in enumerate(found):
        fix = fixes.get(f"y{i}")
        if fix is None:
            continue  # not confirmed: stays yellow, finalise will refuse
        if len(r.text) > 500:
            print(f"y{i}: too long for the on-screen check; check it in Word", file=sys.stderr)
            sys.exit(1)
        value = str(fix.get("value", ""))
        if any(ord(ch) < 32 or invisible(ch) for ch in value) or len(value) > 500:
            print(f"y{i}: value refused", file=sys.stderr)
            sys.exit(1)
        if value != r.text:          # "OK as is" never touches the text (review M1)
            r.text = value
        clear_mark(r)                # highlight and yellow shading
    doc.save(str(out))


# Uploaded Word files are untrusted (reviews M2/N1/N2). Checked on the raw file, before
# python-docx or LibreOffice read it.
MAX_UNZIPPED = 100 * 1024 * 1024
MAX_ENTRIES = 2000
FORBIDDEN_PARTS = ("vbaproject.bin", "/embeddings/", "/activex/")
# Field codes that make Word or LibreOffice fetch or include outside content.
FIELD_RE = re.compile(r"\b(INCLUDEPICTURE|INCLUDETEXT|INCLUDE|LINK|DDEAUTO|DDE|IMPORT)\b", re.I)
REL_HYPERLINK = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"
W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def read_all(z, info, budget):
    """Real bytes, not the size the zip claims (a crafted file can understate it)."""
    out = bytearray()
    with z.open(info) as f:
        while chunk := f.read(65536):
            out += chunk
            budget[0] -= len(chunk)
            if budget[0] < 0:
                sys.exit("refused: the Word file unpacks too large")
    return bytes(out)


A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def xml_text(root):
    """All text of a part, per paragraph: Word text (w:t), deleted text (w:delText) and text
    in charts and diagrams (a:t) -- review N1."""
    kinds = (f"{W_NS}t", f"{W_NS}delText", f"{A_NS}t")
    lines = []
    for p in list(root.iter(f"{W_NS}p")) + list(root.iter(f"{A_NS}p")):
        lines.append("".join(e.text or "" for e in p.iter() if e.tag in kinds))
    return "\n".join(lines)


def cmd_check(docx, forbid_path, engine_scripts, shown_path=None):
    """Refuses a Word file that is not safe to process, that names the supplier (rule 3), or
    that does not print a value its file name is made of (name and content agree, Willem 19/10)."""
    import xml.etree.ElementTree as ET
    import zipfile
    sys.path.insert(0, str(engine_scripts))
    from provenance import norm               # the engine's own normalisation

    texts = []
    budget = [MAX_UNZIPPED]
    try:
        with zipfile.ZipFile(docx) as z:
            infos = z.infolist()
            if len(infos) > MAX_ENTRIES:
                sys.exit("refused: the Word file has too many parts")
            for i in infos:
                name = "/" + i.filename.lower()
                if any(f in name for f in FORBIDDEN_PARTS):
                    sys.exit("refused: the Word file contains macros or embedded objects")
                data = read_all(z, i, budget)
                if name.endswith(".rels"):
                    for rel in ET.fromstring(data):
                        mode = next((v for k, v in rel.attrib.items() if k.lower() == "targetmode"), "")
                        if mode.lower() == "external" and rel.attrib.get("Type") != REL_HYPERLINK:
                            sys.exit("refused: the Word file links to outside content")
                        texts.append(rel.attrib.get("Target", ""))   # link addresses too (review F2)
                elif name.startswith("/word/") and name.endswith(".xml"):
                    root = ET.fromstring(data)
                    # Tracked changes keep deleted text that a PDF may still show (review N1).
                    if next(root.iter(f"{W_NS}ins"), None) is not None or next(root.iter(f"{W_NS}del"), None) is not None:
                        sys.exit("refused: the Word file has tracked changes; accept or reject all changes first")
                    if b"instrText" in data or b"fldSimple" in data:
                        codes = [t.text or "" for t in root.iter(f"{W_NS}instrText")]
                        codes += [e.get(f"{W_NS}instr", "") for e in root.iter(f"{W_NS}fldSimple")]
                        if any(FIELD_RE.search(c) for c in codes):
                            sys.exit("refused: the Word file includes outside content through a field")
                    # body, headers, footers, footnotes, endnotes, comments, text boxes, charts
                    texts.append(xml_text(root))
                    # field codes (HYPERLINK "...") and bookmark-link targets are not shown as text
                    texts.extend(t.text or "" for t in root.iter(f"{W_NS}instrText"))
                    texts.extend(h.get(f"{W_NS}anchor", "") for h in root.iter(f"{W_NS}hyperlink"))
                elif name in ("/docprops/core.xml", "/docprops/app.xml"):
                    root = ET.fromstring(data)
                    texts.extend(e.text or "" for e in root.iter())
    except (zipfile.BadZipFile, ET.ParseError):
        sys.exit("refused: not a valid Word file")

    haystack = norm("\n".join(texts))
    compact = haystack.replace(" ", "")
    for supplier in json.loads(Path(forbid_path).read_text(encoding="utf-8")):
        n = norm(supplier)
        if len(n) < 3:
            continue
        if f" {n} " in f" {haystack} " or (len(n.replace(' ', '')) >= 6 and n.replace(" ", "") in compact):
            sys.exit("refused: the document still names the supplier")

    for label, value in json.loads(Path(shown_path).read_text(encoding="utf-8")) if shown_path else []:
        n = norm(value)
        # norm() keeps Latin letters and digits only: a value with other letters, or too
        # short to mean anything, cannot be checked and is refused (review M2)
        shortest = 3 if label == "product name" else 2      # review R3
        if not latin_only(value) or len(n.replace(" ", "")) < shortest:
            sys.exit(f"refused: the {label} '{value}' cannot be checked against the document; "
                     "use Latin letters and digits")
        # whole words only: "Acetone" must not match inside another name
        if f" {n} " not in f" {haystack} ":
            sys.exit(f"refused: the {label} '{value}' is not printed in the document; "
                     "the file name must match the document, so change one of them")


def latin_only(value):
    """True when every letter or digit is plain Latin once accents are taken off."""
    import unicodedata
    t = unicodedata.normalize("NFKD", value)
    return all(ord(ch) < 128 or unicodedata.combining(ch) or not ch.isalnum() for ch in t)


MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"


def paragraph_texts(docx):
    """Every paragraph's text in reading order: body (tables, text boxes), then headers and
    footers. A text box is stored twice by Word; only the shown version is read."""
    import xml.etree.ElementTree as ET
    import zipfile
    P, T = f"{W_NS}p", f"{W_NS}t"
    BREAKS = {f"{W_NS}br", f"{W_NS}cr", f"{W_NS}tab"}

    def walk(el, out):
        for c in el:
            if c.tag == MC_FALLBACK:
                continue
            if c.tag == P:
                own, nested = [], []
                collect(c, own, nested)
                out.append("".join(own))
                out.extend(nested)
            else:
                walk(c, out)

    def collect(el, own, nested):
        for c in el:
            if c.tag == MC_FALLBACK:
                continue
            if c.tag == T:
                own.append(c.text or "")
            elif c.tag in BREAKS:
                own.append(" ")                 # a line break or tab keeps words apart
            elif c.tag == P:
                walk(_wrap(c), nested)          # a text box inside this paragraph
            else:
                collect(c, own, nested)

    def _wrap(p):
        holder = ET.Element("holder")
        holder.append(p)
        return holder

    out = []
    with zipfile.ZipFile(docx) as z:
        names = z.namelist()
        parts = ["word/document.xml"] + sorted(n for n in names if re.fullmatch(r"word/(header|footer)\d*\.xml", n))
        # also footnotes, endnotes and comments (review F7)
        parts += ["word/footnotes.xml", "word/endnotes.xml", "word/comments.xml"]
        for n in parts:
            if n in names:
                walk(ET.fromstring(z.read(n)), out)
    return [t.strip() for t in out if t.strip()]


# A word or a number as written: "99.6", "0,4", "SB-2026-114" stay whole, so 99.6 -> 6.99 is a change.
WORD = re.compile(r"[^\W_]+(?:[.,/:-][^\W_]+)*")


def words(text):
    import unicodedata
    return {w.lower() for w in WORD.findall(unicodedata.normalize("NFKC", text))}


def pictures(docx):
    """Fingerprints of the pictures in a Word file."""
    import hashlib
    import zipfile
    with zipfile.ZipFile(docx) as z:
        return {hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.lower().startswith("word/media/")}


def cmd_changes(draft, edited, source_pdf):
    """What a person changed on screen, as JSON: [{before, after, unknown}]. unknown = new words
    or numbers that are not in the supplier original: a warning, the person may go on (Willem
    19/10). Nothing is refused here; the safety checks are in cmd_check."""
    import difflib
    import unicodedata
    import pdfplumber

    try:
        with pdfplumber.open(source_pdf) as pdf:
            source = "\n".join(p.extract_text() or "" for p in pdf.pages[:200])
    except Exception:
        source = ""                                # unreadable original: every new value warns
    known = words(source)

    # capped so a huge file cannot hold the worker for minutes (review F6)
    full_a, full_b = paragraph_texts(draft), paragraph_texts(edited)
    a, b = full_a[:20000], full_b[:20000]
    changes = []
    if len(full_a) > 20000 or len(full_b) > 20000:
        changes.append({"before": "", "after": "[very long document: only the first 20000 paragraphs were compared]",
                        "unknown": ["check the rest yourself"]})
    # a picture that was not in the draft: it could be a supplier logo (review F2)
    for _ in range(len(pictures(edited) - pictures(draft))):
        changes.append({"before": "", "after": "[new picture added]", "unknown": ["picture: check it is not a supplier logo"]})
    for tag, i1, i2, j1, j2 in difflib.SequenceMatcher(None, a, b, autojunk=False).get_opcodes():
        if tag == "equal":
            continue
        before, after = "\n".join(a[i1:i2]), "\n".join(b[j1:j2])
        had = words(before)
        unknown = []
        for word in WORD.findall(unicodedata.normalize("NFKC", after)):
            if word.lower() not in had and word.lower() not in known and word not in unknown:
                unknown.append(word)
        changes.append({"before": before[:500], "after": after[:500], "unknown": unknown[:20]})
        if len(changes) >= 200:
            break
    json.dump(changes, sys.stdout, ensure_ascii=False)


# ---- Own on-screen editor (Willem 20/10/2026: no SuperDoc, no outside licence) ---------------
# The draft is shown in the browser as a simple view in which every piece of text (a Word
# "run") has a number. The browser sends back only the runs that changed, and they are written
# into the engine's own Word file: tables, pictures and the TMC template cannot break. The same
# walk gives the numbers for the view and for writing in, so they always agree; each change also
# carries the run's old text, and a run that no longer reads that is refused.

W_URI = W_NS[1:-1]
R_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
# A run whose content is only text, tabs and line breaks can be retyped without losing anything.
TEXT_ONLY = {f"{W_NS}{t}" for t in ("rPr", "t", "tab", "br", "cr", "noBreakHyphen", "softHyphen", "lastRenderedPageBreak")}
NOT_SHOWN = {f"{W_NS}{t}" for t in ("del", "moveFrom", "pPr")} | {MC_FALLBACK}
IMAGE_TYPES = {"image/png", "image/jpeg", "image/gif"}
MAX_IMAGE = 600 * 1024          # one picture in the view
MAX_IMAGES = 6 * 1024 * 1024    # all pictures in the view
MAX_RUNS = 20000
MAX_PART = 5000                 # characters in one piece of text
MAX_TEXT = 300_000              # characters in all changes together
SAFE_FONT = re.compile(r"^[A-Za-z0-9 ._-]{1,40}$")
HEX = re.compile(r"^[0-9A-Fa-f]{6}$")


def _containers(doc):
    """Body, then every header and footer once (as iter_highlighted reads them)."""
    out = [("body", doc.element.body, doc)]
    seen = []
    for s in doc.sections:
        for attr, kind in (("first_page_header", "header"), ("header", "header"), ("even_page_header", "header"),
                           ("first_page_footer", "footer"), ("footer", "footer"), ("even_page_footer", "footer")):
            hf = getattr(s, attr)
            if hf.is_linked_to_previous:
                continue
            el = hf._element
            if any(el is x for x in seen):
                continue
            seen.append(el)
            out.append((kind, el, hf))
    return out


def _numbered(doc):
    """Every run and paragraph with its number (r0.., p0..) and the part it belongs to."""
    runs, paras = [], []
    for _kind, el, parent in _containers(doc):
        runs += [(r, parent) for r in el.xpath('.//w:r[not(ancestor::*[local-name()="Fallback"])]')]
        paras += [(p, parent) for p in el.xpath('.//w:p[not(ancestor::*[local-name()="Fallback"])]')]
    return runs, paras


def _tables(doc):
    """Every table with its number (t0..) and the part it belongs to (same walk as the runs)."""
    out = []
    for _kind, el, parent in _containers(doc):
        out += [(t, parent) for t in el.xpath('.//w:tbl[not(ancestor::*[local-name()="Fallback"])]')]
    return out


# ---- rows and columns (Willem 22/10) ---------------------------------------------------------
# A table may get rows and columns only when its grid is plain: every row fills the grid, no
# cell merged with the row below, no cell inside a content control. A new row or column copies
# the formatting of its neighbour, without text and without yellow; widths are shared out so
# the table keeps its width on the TMC page.

def _cells(tr):
    return tr.xpath("./w:tc")


def _span(tc):
    v = tc.xpath("./w:tcPr/w:gridSpan/@w:val")
    return int(v[0]) if v and v[0].isdigit() else 1


def _set_span(tc, n):
    from docx.oxml import OxmlElement
    tcPr = tc.find(f"{W_NS}tcPr")
    if tcPr is None:
        tcPr = OxmlElement("w:tcPr")
        tc.insert(0, tcPr)
    g = tcPr.find(f"{W_NS}gridSpan")
    if n == 1:
        if g is not None:
            tcPr.remove(g)
        return
    if g is None:
        g = OxmlElement("w:gridSpan")
        after = tcPr.find(f"{W_NS}tcW")
        if after is not None:
            after.addnext(g)
        else:
            tcPr.insert(0, g)
    g.set(f"{W_NS}val", str(n))


ALIGNS = {"left": "LEFT", "center": "CENTER", "right": "RIGHT", "justify": "JUSTIFY"}


def _removed_text(tbl, op):
    """The text a row or column change takes out of the table (for the warning), or ''."""
    rows = tbl.xpath("./w:tr")
    at = op.get("at")
    if not isinstance(at, int) or isinstance(at, bool) or at < 0:
        return ""
    if op.get("op") == "delRow" and at < len(rows):
        return " | ".join("".join(t.text or "" for t in tc.iter(f"{W_NS}t")) for tc in _cells(rows[at])).strip(" |")
    if op.get("op") == "delCol":
        out = []
        for tr in rows:
            col = 0
            for tc in _cells(tr):
                if col <= at < col + _span(tc) and _span(tc) == 1:
                    out.append("".join(t.text or "" for t in tc.iter(f"{W_NS}t")))
                col += _span(tc)
        return " | ".join(x for x in out if x)
    return ""


def plain_grid(tbl):
    """True when rows and columns may be added or removed on screen."""
    cols = tbl.xpath("./w:tblGrid/w:gridCol")
    rows = tbl.xpath("./w:tr")
    if not cols or not rows or tbl.xpath("./w:tr/w:sdt | ./w:tr/w:tc/w:tcPr/w:vMerge"):
        return False
    if tbl.xpath("./w:tr/w:trPr/w:gridBefore | ./w:tr/w:trPr/w:gridAfter"):
        return False
    allowed = {f"{W_NS}{t}" for t in ("trPr", "tblPrEx", "tc", "bookmarkStart", "bookmarkEnd")}
    if any(c.tag not in allowed for tr in rows for c in tr):
        return False                                # cells inside other markup: copied with their text
    return all(sum(_span(tc) for tc in _cells(tr)) == len(cols) for tr in rows)


NOT_COPIED = {f"{W_NS}{t}" for t in ("highlight", "shd", "ins", "del", "rPrChange", "moveFrom", "moveTo",
                                        "vanish", "specVanish", "webHidden")}
MAX_TABLE_OPS, MAX_ROWS, MAX_COLS = 500, 500, 40


def _sid(e):
    """The id of one change, only when it is text."""
    v = e.get("id") if isinstance(e, dict) else None
    return v if isinstance(v, str) else None


def _blank(tc_src, templates):
    """A copy of a cell without its text, nested tables or yellow; remembers the letter to use."""
    import copy
    tc = copy.deepcopy(tc_src)
    from tmc_common import YELLOW_FILLS
    for el in tc.xpath("./w:tcPr/w:vMerge | ./w:tcPr/w:tcPrChange | ./w:tcPr/w:cellIns | ./w:tcPr/w:cellDel"):
        el.getparent().remove(el)
    for shd in tc.xpath("./w:tcPr/w:shd | ./w:p/w:pPr/w:shd"):
        if (shd.get(f"{W_NS}fill") or "").upper() in YELLOW_FILLS:
            shd.getparent().remove(shd)
    first = tc.find(f"{W_NS}p")
    rPr = None
    if first is not None:
        model = first.xpath(".//w:r[w:t]/w:rPr")
        rPr = copy.deepcopy(model[0]) if model else None
    for c in list(tc):
        if c.tag != f"{W_NS}tcPr" and c is not first:
            tc.remove(c)
    if first is None:
        from docx.oxml import OxmlElement
        first = OxmlElement("w:p")
        tc.append(first)
    for c in list(first):
        if c.tag != f"{W_NS}pPr":
            first.remove(c)
    for el in first.xpath("./w:pPr/w:pPrChange | ./w:pPr/w:rPr"):
        el.getparent().remove(el)
    if rPr is not None:
        for c in list(rPr):
            if c.tag in NOT_COPIED:
                rPr.remove(c)
    templates[tc] = rPr
    return tc


def _widths(tbl):
    out = []
    for g in tbl.xpath("./w:tblGrid/w:gridCol"):
        v = g.get(f"{W_NS}w")
        out.append(int(v) if v and v.isdigit() else 1000)
    return out


def _set_widths(tbl, widths):
    """Grid widths, and each cell's width as the sum of the columns it covers."""
    from docx.oxml import OxmlElement
    grid = tbl.find(f"{W_NS}tblGrid")
    for g in list(grid):
        grid.remove(g)
    for w in widths:
        g = OxmlElement("w:gridCol")
        g.set(f"{W_NS}w", str(w))
        grid.append(g)
    for tr in tbl.xpath("./w:tr"):
        col = 0
        for tc in _cells(tr):
            span = _span(tc)
            tcW = tc.find(f"{W_NS}tcPr/{W_NS}tcW")
            if tcW is not None:
                tcW.set(f"{W_NS}w", str(sum(widths[col:col + span])))
                tcW.set(f"{W_NS}type", "dxa")
            col += span


def table_op(tbl, op, templates, created):
    """One row or column change on a plain table. Returns nothing; refuses what does not fit."""
    import copy
    kind = op.get("op")
    rows = tbl.xpath("./w:tr")
    widths = _widths(tbl)
    n = len(widths)

    def take_ids(count_or_flags):
        ids = op.get("ids")
        if not isinstance(ids, list) or len(ids) != len(count_or_flags):
            sys.exit("refused: the new cells do not match the table; open the editor again")
        for i, need in zip(ids, count_or_flags):
            ok = (i is None) if not need else (isinstance(i, str) and re.fullmatch(r"n\d{1,5}", i) and i not in created)
            if not ok:
                sys.exit("refused: the new cells do not match the table; open the editor again")
        return ids

    def index(name, top):
        v = op.get(name)
        if not isinstance(v, int) or isinstance(v, bool) or not 0 <= v <= top:
            sys.exit(f"refused: table change with a wrong {name}")
        return v

    if kind == "addRow":
        at, src = index("at", len(rows)), index("from", len(rows) - 1)
        tr = copy.deepcopy(rows[src])
        for h in tr.xpath("./w:trPr/w:tblHeader | ./w:trPr/w:ins | ./w:trPr/w:del | ./w:trPr/w:trPrChange"):
            h.getparent().remove(h)
        cells = _cells(tr)
        ids = take_ids([True] * len(cells))
        for old, i in zip(cells, ids):
            new = _blank(old, templates)
            old.addprevious(new)
            tr.remove(old)
            created[i] = new
        (rows[at].addprevious if at < len(rows) else rows[-1].addnext)(tr)
    elif kind == "delRow":
        at = index("at", len(rows) - 1)
        if len(rows) == 1:
            sys.exit("refused: a table keeps at least one row")
        tbl.remove(rows[at])
    elif kind == "addCol":
        at = index("at", n)
        src = at - 1 if at > 0 else 0
        half = max(widths[src] // 2, 200)
        widths[src] = max(widths[src] - half, 200)
        widths.insert(at, half)
        plan = []
        for tr in rows:
            col, action = 0, None
            for tc in _cells(tr):
                span = _span(tc)
                if col < at < col + span:
                    action = ("widen", tc)
                elif col == at:
                    action = action or ("before", tc)
                elif col + span == at:
                    action = ("after", tc)
                col += span
            plan.append(action)
        ids = take_ids([a[0] != "widen" for a in plan])
        for (how, tc), i in zip(plan, ids):
            if how == "widen":
                _set_span(tc, _span(tc) + 1)
                continue
            new = _blank(tc, templates)
            _set_span(new, 1)
            (tc.addprevious if how == "before" else tc.addnext)(new)
            created[i] = new
        _set_widths(tbl, widths)
    elif kind == "delCol":
        at = index("at", n - 1)
        if n == 1:
            sys.exit("refused: a table keeps at least one column")
        for tr in rows:
            col = 0
            for tc in _cells(tr):
                span = _span(tc)
                if col <= at < col + span:
                    if span > 1:
                        _set_span(tc, span - 1)
                    else:
                        tr.remove(tc)
                    break
                col += span
            if not _cells(tr):
                sys.exit("refused: a row would lose its last cell")
        gone = widths.pop(at)
        widths[at - 1 if at > 0 else 0] += gone
        _set_widths(tbl, widths)
    else:
        sys.exit("refused: unknown table change")


def _editable(r_el):
    if any(c.tag not in TEXT_ONLY for c in r_el):
        return False
    if any(b.get(f"{W_NS}type") not in (None, "textWrapping") for b in r_el.iter(f"{W_NS}br")):
        return False                              # a page or column break: layout, not text
    return r_el.find(f"{W_NS}t") is not None and not r_el.xpath("ancestor::w:del | ancestor::w:moveFrom")


def _owner_run(el):
    """The nearest run around an element (pictures are not python-docx elements: no xpath prefixes)."""
    el = el.getparent()
    while el is not None and el.tag != f"{W_NS}r":
        el = el.getparent()
    return el


def _run_obj(r_el, parent):
    from docx.text.paragraph import Paragraph
    from docx.text.run import Run
    p_els = r_el.xpath("ancestor::w:p[1]")
    para = Paragraph(p_els[0], parent) if p_els else None
    return Run(r_el, para if para is not None else parent), para


def _chain(style):
    while style is not None:
        yield style
        style = style.base_style


def _effective(run, para, attr):
    """Bold, italic or underline as Word shows it: the run, its character style, its paragraph style."""
    v = getattr(run.font, attr)
    if v is None:
        try:
            styles = list(_chain(run.style)) + (list(_chain(para.style)) if para is not None else [])
        except Exception:
            styles = []
        v = next((getattr(s.font, attr) for s in styles if getattr(s.font, attr) is not None), None)
    return bool(v)


def _look(run, para):
    """Font, size and colour for the view only (never written back)."""
    try:
        styles = list(_chain(run.style)) + (list(_chain(para.style)) if para is not None else [])
    except Exception:
        styles = []
    fonts = [run.font] + [s.font for s in styles]
    out = {}
    size = next((f.size for f in fonts if f.size is not None), None)
    if size is not None and 4 <= size.pt <= 96:
        out["size"] = round(size.pt, 1)
    name = next((f.name for f in fonts if f.name), None)
    if name and SAFE_FONT.match(name):
        out["font"] = name
    for f in fonts:
        try:
            rgb = f.color.rgb if f.color is not None and f.color.type is not None else None
        except Exception:
            rgb = None
        if rgb is not None and HEX.match(str(rgb)):
            out["color"] = f"#{rgb}"
            break
    return out


def _mm(twips):
    try:
        return round(int(twips) / 56.7, 1)
    except (TypeError, ValueError):
        return None


class _View:
    def __init__(self, doc):
        from tmc_common import is_marked
        self.is_marked = is_marked
        self.runs, self.paras = _numbered(doc)
        if len(self.runs) > MAX_RUNS:
            sys.exit(f"refused: too many text pieces for the on-screen editor ({len(self.runs)})")
        self.rid = {r: f"r{i}" for i, (r, _) in enumerate(self.runs)}
        self.rparent = {r: parent for r, parent in self.runs}
        self.pid = {p: f"p{i}" for i, (p, _) in enumerate(self.paras)}
        self.pparent = {p: parent for p, parent in self.paras}
        self.tid = {t: f"t{i}" for i, (t, _) in enumerate(_tables(doc))}
        self.tparent = {t: parent for t, parent in _tables(doc)}
        self.doc = doc
        self.media, self.media_bytes, self.media_ids = {}, 0, {}
        self.styles = doc.styles.element

    def picture(self, part, rel_id, extent):
        import base64
        import hashlib
        try:
            target = part.related_parts[rel_id]
        except (KeyError, AttributeError):
            return None
        if target.content_type not in IMAGE_TYPES or len(target.blob) > MAX_IMAGE:
            return {"t": "img", "m": None, **extent}
        digest = hashlib.sha256(target.blob).hexdigest()
        if digest not in self.media_ids:
            if self.media_bytes + len(target.blob) > MAX_IMAGES:
                return {"t": "img", "m": None, **extent}
            self.media_bytes += len(target.blob)
            self.media_ids[digest] = f"m{len(self.media_ids)}"
            self.media[self.media_ids[digest]] = f"data:{target.content_type};base64," + base64.b64encode(target.blob).decode()
        return {"t": "img", "m": self.media_ids[digest], **extent}

    def run_items(self, r_el):
        run, para = _run_obj(r_el, self.rparent[r_el])
        items = []
        if r_el.find(f"{W_NS}instrText") is None:    # a field's code is not shown, only its result
            fmt = {"b": _effective(run, para, "bold"), "i": _effective(run, para, "italic"),
                   "u": _effective(run, para, "underline"), "hl": bool(self.is_marked(run)), **_look(run, para)}
            text = run.text
            if _editable(r_el):
                items.append({"t": "r", "id": self.rid[r_el], "text": text, **fmt})
            elif text:
                items.append({"t": "r", "id": None, "text": text, **fmt})
        part = self.rparent[r_el].part
        # pictures and text boxes of this run itself (not of a run inside its text box)
        for el in r_el.xpath('.//*[(local-name()="blip" or local-name()="imagedata" or local-name()="txbxContent")'
                             ' and not(ancestor::*[local-name()="Fallback"])]'):
            if _owner_run(el) is not r_el:
                continue
            if el.tag == f"{W_NS}txbxContent":
                items.append({"t": "box", "blocks": self.blocks(el)})
                continue
            rel = el.get(R_EMBED) or el.get(R_ID)
            ext = el.xpath('ancestor::*[local-name()="inline" or local-name()="anchor"][1]/*[local-name()="extent"]')
            extent = {}
            if ext:
                try:
                    extent = {"w": round(int(ext[0].get("cx")) / 36000, 1), "h": round(int(ext[0].get("cy")) / 36000, 1)}
                except (TypeError, ValueError):
                    pass
            pic = self.picture(part, rel, extent) if rel else None
            if pic:
                items.append(pic)
        return items

    def para(self, p_el):
        from docx.text.paragraph import Paragraph
        items = []

        def walk(el):
            for c in el:
                if c.tag in NOT_SHOWN:
                    continue
                if c.tag == f"{W_NS}r":
                    items.extend(self.run_items(c))
                elif c.tag != f"{W_NS}p":           # hyperlinks, fields, inserted text, content controls
                    walk(c)

        walk(p_el)
        out = {"t": "p", "id": self.pid[p_el], "items": items}
        para = Paragraph(p_el, self.pparent[p_el])
        try:
            fmt = [para.paragraph_format] + [s.paragraph_format for s in _chain(para.style)]
        except Exception:
            fmt = [para.paragraph_format]
        align = next((f.alignment for f in fmt if f.alignment is not None), None)
        if align is not None:
            name = {0: "left", 1: "center", 2: "right", 3: "justify"}.get(int(align))
            if name:
                out["align"] = name
        for key, attr in (("before", "space_before"), ("after", "space_after")):
            v = next((getattr(f, attr) for f in fmt if getattr(f, attr) is not None), None)
            if v is not None and 0 <= v.pt <= 200:
                out[key] = round(v.pt, 1)
        indent = next((f.left_indent for f in fmt if f.left_indent is not None), None)
        if indent is not None and 0 < indent.mm < 200:
            out["indent"] = round(indent.mm, 1)
        if not any(i["t"] == "r" and i["id"] for i in items):
            out["empty"] = True                     # text typed here is added at its end
        return out

    def borders(self, tbl):
        """Does the table draw lines? Its own setting, else its table style (followed up)."""
        def has(tblPr):
            b = tblPr.find(f"{W_NS}tblBorders") if tblPr is not None else None
            return None if b is None else any(e.get(f"{W_NS}val") not in ("nil", "none") for e in b)
        found = has(tbl.find(f"{W_NS}tblPr"))
        style = tbl.xpath("./w:tblPr/w:tblStyle/@w:val")
        sid = style[0] if style else None
        for _ in range(10):
            if found is not None or not sid:
                break
            el = self.styles.xpath(f'./w:style[@w:styleId="{sid}"]') if re.fullmatch(r"[\w-]{1,60}", sid) else []
            if not el:
                break
            found = has(el[0].find(f"{W_NS}tblPr"))
            base = el[0].xpath("./w:basedOn/@w:val")
            sid = base[0] if base else None
        return bool(found)

    def table(self, tbl):
        cols = [_mm(g.get(f"{W_NS}w")) or 20 for g in tbl.xpath("./w:tblGrid/w:gridCol")]
        rows, active = [], {}
        for tr in tbl.xpath("./w:tr"):
            row, col = [], 0
            for tc in tr.xpath("./w:tc | ./w:sdt/w:sdtContent/w:tc"):
                span = tc.xpath("./w:tcPr/w:gridSpan/@w:val")
                span = int(span[0]) if span and span[0].isdigit() else 1
                vm = tc.find(f"{W_NS}tcPr/{W_NS}vMerge")
                if vm is not None and vm.get(f"{W_NS}val") != "restart" and col in active:
                    active[col]["rowspan"] = active[col].get("rowspan", 1) + 1
                    col += span
                    continue
                cell = {"blocks": self.blocks(tc)}
                if span > 1:
                    cell["span"] = span
                fill = tc.xpath("./w:tcPr/w:shd/@w:fill")
                if fill and HEX.match(fill[0]):
                    cell["shade"] = f"#{fill[0]}"
                if vm is not None:
                    active[col] = cell
                else:
                    active.pop(col, None)
                row.append(cell)
                col += span
            rows.append(row)
        plain = self.tparent.get(tbl) is self.doc and plain_grid(tbl)     # never in a header or footer
        return {"t": "table", "id": self.tid.get(tbl), "plain": plain, "cols": cols, "lines": self.borders(tbl), "rows": rows}

    def blocks(self, el):
        out = []
        for c in el:
            if c.tag == f"{W_NS}p":
                out.append(self.para(c))
            elif c.tag == f"{W_NS}tbl":
                out.append(self.table(c))
            elif c.tag in (f"{W_NS}sdt", f"{W_NS}sdtContent", f"{W_NS}customXml"):
                out.extend(self.blocks(c))
        return out


def cmd_view(docx, out):
    """The editor's view of a draft as JSON: page size, headers, body, footers, pictures."""
    from docx import Document
    doc = Document(str(docx))
    v = _View(doc)
    sect = doc.sections[0]
    page = {"width": round(sect.page_width.mm, 1) if sect.page_width else 210}
    for key in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        m = getattr(sect, key)
        page[key.split("_")[0]] = round(m.mm, 1) if m is not None else 20
    parts = [{"kind": kind, "blocks": v.blocks(el)} for kind, el, _ in _containers(doc)]
    order = {"header": 0, "body": 1, "footer": 2}
    parts.sort(key=lambda p: order[p["kind"]])     # shown top to bottom; numbers do not change
    Path(out).write_text(json.dumps({"v": 1, "page": page, "parts": parts, "media": v.media}, ensure_ascii=False),
                         encoding="utf-8")


SHOWS_NOTHING = {0x034F, 0x115F, 0x1160, 0x17B4, 0x17B5, 0x3164, 0xFFA0}


def invisible(ch):
    """Control, format and other characters that show nothing (tab and line break excepted):
    they could split a supplier name so the check misses it (review M1; same set as
    web/lib/documents/edits.ts)."""
    import unicodedata
    o = ord(ch)
    if ch in "\t\n":
        return False
    return (unicodedata.category(ch) in ("Cc", "Cf", "Zl", "Zp") or o in SHOWS_NOTHING
            or 0x180B <= o <= 0x180F or 0xFE00 <= o <= 0xFE0F or 0xE0100 <= o <= 0xE01EF)


def _parts(raw, where):
    """The new pieces of one run or paragraph, checked: text and four on/off switches."""
    if not isinstance(raw, list) or not 1 <= len(raw) <= 50:
        sys.exit(f"refused: {where}: wrong list of pieces")
    out = []
    for p in raw:
        if not isinstance(p, dict) or not isinstance(p.get("text"), str) or len(p["text"]) > MAX_PART:
            sys.exit(f"refused: {where}: text missing or too long")
        if any(invisible(ch) for ch in p["text"]):
            sys.exit(f"refused: {where}: control or invisible characters")
        if any(not isinstance(p.get(k), bool) for k in ("b", "i", "u", "hl")):
            sys.exit(f"refused: {where}: formatting must be true or false")
        size = p.get("size")
        if size is not None and (isinstance(size, bool) or not isinstance(size, (int, float))
                                 or not 4 <= size <= 96 or size * 2 != int(size * 2)):
            sys.exit(f"refused: {where}: letter size")
        out.append(p)
    return out


def _before(run, para):
    """A piece's formatting as the view showed it, to change only what the person changed."""
    from tmc_common import is_marked
    return {"b": _effective(run, para, "bold"), "i": _effective(run, para, "italic"),
            "u": _effective(run, para, "underline"), "hl": bool(is_marked(run)), "size": _look(run, para).get("size")}


def _restyle(run, para, part, before):
    """Sets bold, italic, underline, letter size and the yellow mark only where the person changed them."""
    from docx.shared import Pt
    from tmc_common import clear_mark, mark_yellow
    for key, attr in (("b", "bold"), ("i", "italic"), ("u", "underline")):
        if part[key] != before[key]:
            setattr(run.font, attr, part[key])
    if part.get("size") is not None and part["size"] != before.get("size"):
        run.font.size = Pt(part["size"])
    if part["hl"] != before["hl"]:
        if part["hl"]:
            mark_yellow(run)
        else:
            clear_mark(run)          # a person removed it on screen (BRIEF 7.2: never automatic)


def cmd_edit(docx, edits_path, out):
    """Writes the editor's changes into the draft: {runs: [{id, was, parts}], paras: [{id, parts}]}.
    Refuses the whole set when anything does not fit; nothing is written then."""
    import copy
    from docx import Document
    from tmc_common import is_marked
    data = json.loads(Path(edits_path).read_text(encoding="utf-8"))
    if not isinstance(data, dict) or not isinstance(data.get("runs", []), list) or not isinstance(data.get("paras", []), list):
        sys.exit("refused: unreadable changes")
    run_edits, para_edits = data.get("runs", []), data.get("paras", [])
    if len(run_edits) + len(para_edits) > 5000:
        sys.exit("refused: too many changes")
    doc = Document(str(docx))
    runs, paras = _numbered(doc)
    total, seen, jobs = 0, set(), []

    for e in run_edits:
        rid = _sid(e)
        m = re.fullmatch(r"r(\d{1,6})", rid or "")
        if not m or rid in seen or int(m[1]) >= len(runs):
            sys.exit(f"refused: unknown or repeated text piece {rid!r}")
        seen.add(rid)
        r_el, parent = runs[int(m[1])]
        run, para = _run_obj(r_el, parent)
        if not _editable(r_el):
            sys.exit(f"refused: {rid} cannot be edited on screen")
        if e.get("was") != run.text:
            sys.exit(f"refused: {rid} does not match the draft; open the editor again")
        parts = _parts(e.get("parts"), rid)
        total += sum(len(p["text"]) for p in parts)
        jobs.append(("run", r_el, parent, parts))

    for e in para_edits:
        pid = _sid(e)
        m = re.fullmatch(r"p(\d{1,6})", pid or "")
        if not m or pid in seen or int(m[1]) >= len(paras):
            sys.exit(f"refused: unknown or repeated paragraph {pid!r}")
        seen.add(pid)
        p_el, parent = paras[int(m[1])]
        own = [r for r in p_el.xpath('.//w:r[not(ancestor::*[local-name()="Fallback"])]') if r.xpath("ancestor::w:p[1]")[0] is p_el]
        if any(_editable(r) for r in own):
            sys.exit(f"refused: {pid} has text of its own; change that text instead")
        parts = _parts(e.get("parts"), pid)
        total += sum(len(p["text"]) for p in parts)
        jobs.append(("para", p_el, parent, parts))
    layout, tables, cells = data.get("layout", []), data.get("tables", []), data.get("cells", [])
    if not all(isinstance(x, list) for x in (layout, tables, cells)) or len(layout) + len(tables) + len(cells) > 5000:
        sys.exit("refused: unreadable changes")
    aligns = []
    for e in layout:
        pid = _sid(e)
        m = re.fullmatch(r"p(\d{1,6})", pid or "")
        if not m or int(m[1]) >= len(paras) or e.get("align") not in ALIGNS:
            sys.exit(f"refused: unknown paragraph or alignment {pid!r}")
        if paras[int(m[1])][1] is not doc:
            sys.exit(f"refused: {pid} is in the TMC header or footer; its layout stays as it is")
        aligns.append((paras[int(m[1])], e["align"]))
    all_tables, table_jobs = _tables(doc), []
    if sum(len(e.get("ops", [])) for e in tables if isinstance(e, dict) and isinstance(e.get("ops"), list)) > MAX_TABLE_OPS:
        sys.exit("refused: too many table changes")
    for e in tables:
        tid = _sid(e)
        m = re.fullmatch(r"t(\d{1,5})", tid or "")
        if not m or int(m[1]) >= len(all_tables) or not isinstance(e.get("ops"), list) or not 1 <= len(e["ops"]) <= 200:
            sys.exit(f"refused: unknown table {tid!r}")
        tbl, parent = all_tables[int(m[1])]
        if parent is not doc:
            sys.exit(f"refused: {tid} is in the TMC header or footer; its rows and columns stay as they are")
        if not plain_grid(tbl):
            sys.exit(f"refused: {tid} has merged cells; change its rows and columns in Word")
        table_jobs.append((tbl, e["ops"]))
    texts = {}
    for e in cells:
        cid = _sid(e)
        if not re.fullmatch(r"n\d{1,5}", cid or "") or cid in texts:
            sys.exit(f"refused: unknown or repeated new cell {cid!r}")
        texts[cid] = _parts(e.get("parts"), cid)
        total += sum(len(p["text"]) for p in texts[cid])
    if total > MAX_TEXT:
        sys.exit("refused: the changes are too long")

    for what, el, parent, parts in jobs:
        if what == "run":
            run, para = _run_obj(el, parent)
            before = _before(run, para)
            template, prev = copy.deepcopy(el), el
            for k, part in enumerate(parts):
                r_el = el
                if k:
                    r_el = copy.deepcopy(template)   # same font, size and colour as the original piece
                    prev.addnext(r_el)
                r, _ = _run_obj(r_el, parent)
                if r.text != part["text"]:
                    r.text = part["text"]
                _restyle(r, para, part, before)
                prev = r_el
        else:
            from docx.text.paragraph import Paragraph
            para = Paragraph(el, parent)
            mark = el.find(f"{W_NS}pPr/{W_NS}rPr")
            for part in parts:
                r = para.add_run(part["text"])
                if mark is not None:                 # the paragraph's own letter, without marks or tracking
                    rPr = r._r.get_or_add_rPr()
                    for c in mark:
                        if c.tag not in (f"{W_NS}highlight", f"{W_NS}shd", f"{W_NS}ins", f"{W_NS}del",
                                         f"{W_NS}rPrChange", f"{W_NS}moveFrom", f"{W_NS}moveTo"):
                            rPr.append(copy.deepcopy(c))
                before = _before(r, para)
                _restyle(r, para, part, before)

    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.text.paragraph import Paragraph
    for (p_el, parent), align in aligns:
        Paragraph(p_el, parent).alignment = getattr(WD_ALIGN_PARAGRAPH, ALIGNS[align])

    # Rows and columns, in the order the person made them. A row or column taken out is reported
    # with its text: nothing from the supplier original may be left out unnoticed (BRIEF 7.1).
    notes, templates, created = [], {}, {}
    for tbl, ops in table_jobs:
        for op in ops:
            if not isinstance(op, dict):
                sys.exit("refused: unknown table change")
            gone = _removed_text(tbl, op)
            table_op(tbl, op, templates, created)
            if len(tbl.xpath("./w:tr")) > MAX_ROWS or len(tbl.xpath("./w:tblGrid/w:gridCol")) > MAX_COLS:
                sys.exit("refused: the table would get too many rows or columns")
            if gone:
                notes.append({"before": gone[:500], "after": "", "unknown": [f"{op['op'] == 'delRow' and 'row' or 'column'} deleted: check nothing from the supplier original is left out"]})
    for cid, parts in texts.items():
        tc = created.get(cid)
        if tc is None:
            sys.exit(f"refused: unknown new cell {cid!r}")
        if not tc.xpath("ancestor::w:body | ancestor::w:hdr | ancestor::w:ftr"):
            continue                                  # its row or column was taken out again
        p_el = tc.find(f"{W_NS}p")
        parent = next(parent for t, parent in all_tables if any(t is a for a in tc.iterancestors(f"{W_NS}tbl")))
        para = Paragraph(p_el, parent)
        for part in parts:
            r = para.add_run(part["text"])
            if templates.get(tc) is not None:
                r._r.insert(0, copy.deepcopy(templates[tc]))
            _restyle(r, para, part, _before(r, para))
    doc.save(str(out))
    json.dump(notes, sys.stdout, ensure_ascii=False)


# ---- layout check on the PDF (Willem 22/10: a layout problem blocks Final) -----------------
# The TMC letterhead (banner at the top, address bar and stamp at the bottom) must stay clear,
# and every word of the Word file must show in the PDF: a long product name that no longer
# fits its place is cut off by LibreOffice without any warning.
PAGE_NO = re.compile(r"^(page\s*)?\d+(\s*(of|/)\s*\d+)?$", re.I)


def tokens(text):
    """Letters and digits only, so a word the PDF breaks at a hyphen still counts."""
    import unicodedata
    return [t.lower() for t in re.findall(r"[^\W_]+", unicodedata.normalize("NFKC", text))]


def _counts_matter(t):
    """Words of 3 letters or more and numbers of 2 digits or more (a cut-off value counts too)."""
    return (len(t) >= 3 and any(ch.isalpha() for ch in t)) or (len(t) >= 2 and t.isdigit())


SHOWN_TEXT = ('.//w:t[not(ancestor::*[local-name()="Fallback"]) and not(ancestor::w:fldSimple)'
              ' and not(ancestor::w:del) and not(ancestor::w:r[w:rPr/w:vanish])]'
              ' | .//w:fldChar[not(ancestor::*[local-name()="Fallback"])]'
              ' | .//w:p | .//w:br | .//w:tab')


def shown_text(docx):
    """The text Word prints: body, and the headers and footers that are in use. Field codes and
    field results (page numbers), hidden and deleted text are left out."""
    from docx import Document
    doc = Document(str(docx))
    roots, seen = [doc.element.body], []
    even = doc.settings.odd_and_even_pages_header_footer
    for sec in doc.sections:
        parts = [sec.header, sec.footer]
        if sec.different_first_page_header_footer:
            parts += [sec.first_page_header, sec.first_page_footer]
        if even:
            parts += [sec.even_page_header, sec.even_page_footer]
        for hf in parts:
            if hf.is_linked_to_previous or any(hf._element is x for x in seen):
                continue
            seen.append(hf._element)
            roots.append(hf._element)
    out = []
    for root in roots:
        field = 0                                   # 0 outside a field, 1 its code, 2 its result
        for e in root.xpath(SHOWN_TEXT):
            if e.tag == f"{W_NS}fldChar":
                field = {"begin": 1, "separate": 2, "end": 0}.get(e.get(f"{W_NS}fldCharType"), field)
            elif e.tag == f"{W_NS}t":
                if field == 0:
                    out.append(e.text or "")
            else:
                out.append("\n")
    return "".join(out)


def layout_problems(pdf_path, docx):
    import pymupdf
    from collections import Counter
    problems, shown = [], Counter()
    with pymupdf.open(pdf_path) as pdf:
        for n, page in enumerate(pdf, 1):
            w, h = page.rect.width, page.rect.height
            sheet = [pymupdf.Rect(i["bbox"]) for i in page.get_image_info()]
            sheet = [r for r in sheet if r.width >= 0.4 * w and (r.y1 <= 0.25 * h or r.y0 >= 0.75 * h)]
            for block in page.get_text("dict")["blocks"]:
                for line in block.get("lines", []):
                    for span in line["spans"]:
                        text = span["text"].strip()
                        if not text:
                            continue
                        box = pymupdf.Rect(span["bbox"])
                        if box.x0 < -1 or box.y0 < -1 or box.x1 > w + 1 or box.y1 > h + 1:
                            problems.append(f"page {n}: text runs off the page: '{text[:60]}'")
                        if PAGE_NO.match(text) or not box.get_area():
                            continue
                        for r in sheet:
                            over = box & r
                            if not over.is_empty and over.get_area() >= 0.5 * box.get_area():
                                problems.append(f"page {n}: text over the TMC letterhead: '{text[:60]}'")
                                break
            shown.update(tokens(page.get_text()))
    wanted = Counter(tokens(shown_text(docx)))
    missing = sorted(t for t, n in wanted.items() if n > shown[t] and _counts_matter(t))
    if missing:
        problems.append("text that does not show in the PDF (cut off?): " + " ".join(missing[:12]))
    return problems


def cmd_layout(pdf_path, docx):
    """Exit 1 with the problems when the PDF's layout is not right; nothing is changed."""
    problems = layout_problems(pdf_path, docx)
    if problems:
        for p in problems[:10]:
            print(f"  {p}", file=sys.stderr)
        sys.exit("refused: layout problem in the PDF; make the text shorter or the letters smaller, then save again")


def cmd_preview(docx, out, engine_scripts):
    sys.path.insert(0, str(engine_scripts))
    from tmc_common import render_pdf          # the engine's own Word -> PDF
    from pypdf import PdfReader, PdfWriter
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    with tempfile.TemporaryDirectory() as d:
        copy = Path(d) / "draft.docx"
        shutil.copy(docx, copy)
        pdf = render_pdf(copy)
        reader = PdfReader(str(pdf))
        writer = PdfWriter()
        for page in reader.pages:
            w, h = float(page.mediabox.width), float(page.mediabox.height)
            buf = io.BytesIO()
            c = canvas.Canvas(buf, pagesize=(w, h) if w and h else A4)
            c.setFillColorRGB(0.85, 0.1, 0.1, alpha=0.25)
            c.setFont("Helvetica-Bold", 54)
            c.translate(w / 2, h / 2)
            c.rotate(45)
            c.drawCentredString(0, 0, "DRAFT - NOT FOR USE")
            c.save()
            buf.seek(0)
            page.merge_page(PdfReader(buf).pages[0])
            writer.add_page(page)
        with open(out, "wb") as f:
            writer.write(f)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--engine", required=True, help="engine folder (contains scripts/)")
    ap.add_argument("command", choices=["list", "apply", "preview", "check", "changes", "view", "edit", "layout"])
    ap.add_argument("paths", nargs="+")
    a = ap.parse_args()
    scripts = Path(a.engine) / "scripts"
    sys.path.insert(0, str(scripts))
    if a.command == "list" and len(a.paths) == 1:
        cmd_list(a.paths[0])
    elif a.command == "apply" and len(a.paths) == 3:
        cmd_apply(*a.paths)
    elif a.command == "preview" and len(a.paths) == 2:
        cmd_preview(a.paths[0], a.paths[1], scripts)
    elif a.command == "changes" and len(a.paths) == 3:
        cmd_changes(a.paths[0], a.paths[1], a.paths[2])
    elif a.command == "layout" and len(a.paths) == 2:
        cmd_layout(a.paths[0], a.paths[1])
    elif a.command == "view" and len(a.paths) == 2:
        cmd_view(a.paths[0], a.paths[1])
    elif a.command == "edit" and len(a.paths) == 3:
        cmd_edit(*a.paths)
    elif a.command == "check" and len(a.paths) in (2, 3):
        cmd_check(a.paths[0], a.paths[1], scripts, a.paths[2] if len(a.paths) == 3 else None)
    else:
        ap.error("wrong number of paths")


if __name__ == "__main__":
    main()
