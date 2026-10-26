"""Tests for web/worker/yellow.py. Run:  python web/worker/test_yellow.py
Needs python-docx (engine requirements). The preview test also needs LibreOffice."""
import json, shutil, subprocess, sys, tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
ENGINE = HERE.parent.parent / "engine"
sys.path.insert(0, str(ENGINE / "scripts"))
from docx import Document                      # noqa: E402
from docx.enum.text import WD_COLOR_INDEX      # noqa: E402
from tmc_common import yellow_runs             # noqa: E402  (the engine's own check)


def run(*args):
    r = subprocess.run([sys.executable, "-I", str(HERE / "yellow.py"), "--engine", str(ENGINE), *map(str, args)],
                       capture_output=True, text=True)
    return r.returncode, r.stdout, r.stderr


def draft(path):
    doc = Document()
    p = doc.add_paragraph("Product: ")
    p.add_run("PRODUCTNAME").font.highlight_color = WD_COLOR_INDEX.YELLOW
    t = doc.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "Expiry date"
    r = t.cell(0, 1).paragraphs[0].add_run("01/09/2O29")   # OCR doubt
    r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    m = doc.add_table(rows=1, cols=2)
    merged = m.cell(0, 0).merge(m.cell(0, 1))          # appears twice in python-docx
    merged.paragraphs[0].add_run("footer sentence kept").font.highlight_color = WD_COLOR_INDEX.YELLOW
    doc.save(path)


ok = bad = 0
def check(name, cond, info=""):
    global ok, bad
    print(("  ok  " if cond else "  XX  ") + name + ("" if cond else f"\n{info}"))
    ok, bad = ok + bool(cond), bad + (not cond)


with tempfile.TemporaryDirectory() as d:
    d = Path(d)
    draft(d / "d.docx")

    rc, out, err = run("list", d / "d.docx")
    items = json.loads(out) if rc == 0 else []
    check("lists every yellow part once (merged cell counted once)", [i["text"] for i in items] ==
          ["PRODUCTNAME", "01/09/2O29", "footer sentence kept"], out + err)
    check("gives the surrounding text", items and items[0]["context"] == "Product: PRODUCTNAME", out)

    (d / "some.json").write_text(json.dumps({"y0": {"value": "Sodium Benzoate"}, "y1": {"value": "01/09/2029"}}))
    rc, out, err = run("apply", d / "d.docx", d / "some.json", d / "some.docx")
    left = yellow_runs(d / "some.docx")
    # unique texts: older engine versions listed a merged cell once per grid column
    check("only confirmed parts lose their yellow; the rest stays", rc == 0 and sorted(set(left)) == ["footer sentence kept"], err + str(left))
    text = "\n".join(p.text for p in Document(str(d / "some.docx")).paragraphs)
    check("corrections are written in", "Product: Sodium Benzoate" in text, text)

    (d / "all.json").write_text(json.dumps({"y0": {"value": "Sodium Benzoate"}, "y1": {"value": "01/09/2029"},
                                            "y2": {"value": "footer sentence kept"}}))
    rc, out, err = run("apply", d / "d.docx", d / "all.json", d / "all.docx")
    check("all confirmed: the engine's own check finds no yellow", rc == 0 and yellow_runs(d / "all.docx") == [], err)

    (d / "bad.json").write_text(json.dumps({"y9": {"value": "x"}}))
    rc, _, err = run("apply", d / "d.docx", d / "bad.json", d / "bad.docx")
    check("unknown part refused, nothing written", rc == 1 and not (d / "bad.docx").exists(), err)

    (d / "ctrl.json").write_text(json.dumps({"y0": {"value": "a\u0007b"}}))
    rc, _, err = run("apply", d / "d.docx", d / "ctrl.json", d / "ctrl.docx")
    check("control characters refused", rc == 1, err)

    (d / "same.json").write_text(json.dumps({"y0": {"value": "PRODUCTNAME"}, "y1": {"value": "01/09/2O29"}, "y2": {"value": "footer sentence kept"}}))
    run("apply", d / "d.docx", d / "same.json", d / "same.docx")
    a = [r.text for p in Document(str(d / "d.docx")).paragraphs for r in p.runs]
    b = [r.text for p in Document(str(d / "same.docx")).paragraphs for r in p.runs]
    check("OK as is never changes the text", a == b, f"{a} != {b}")

    long_doc = Document(); long_doc.add_paragraph().add_run("x" * 600).font.highlight_color = WD_COLOR_INDEX.YELLOW
    long_doc.save(d / "long.docx")
    _, out, _ = run("list", d / "long.docx")
    check("parts over 500 characters are flagged for Word", json.loads(out)[0]["truncated"] is True, out)
    (d / "long.json").write_text(json.dumps({"y0": {"value": "x" * 500}}))
    rc, _, _ = run("apply", d / "long.docx", d / "long.json", d / "long_out.docx")
    check("and cannot be finalised on screen", rc == 1)

    (d / "forbid.json").write_text(json.dumps(["Acme Chemicals GmbH"]))
    rc, _, err = run("check", d / "all.docx", d / "forbid.json")
    check("check: a clean Word file passes", rc == 0, err)
    named = Document(); named.add_paragraph("Made by ACME   chemicals gmbh, Berlin"); named.save(d / "named.docx")
    rc, _, err = run("check", d / "named.docx", d / "forbid.json")
    check("check: refused when the supplier is named (any case or spacing)", rc == 1 and "supplier" in err, err)
    hf = Document(); hf.sections[0].footer.paragraphs[0].text = "Acme Chemicals GmbH"; hf.save(d / "hf.docx")
    rc, _, err = run("check", d / "hf.docx", d / "forbid.json")
    check("check: also looks in headers and footers", rc == 1, err)
    import zipfile
    with zipfile.ZipFile(d / "all.docx") as src, zipfile.ZipFile(d / "ext.docx", "w") as dst:
        for i in src.infolist():
            data = src.read(i)
            if i.filename == "word/_rels/document.xml.rels":
                data = data.replace(b"</Relationships>", b'<Relationship Id="rX" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="http://evil.example/x.png" TargetMode="External"/></Relationships>')
            dst.writestr(i, data)
    rc, _, err = run("check", d / "ext.docx", d / "forbid.json")
    check("check: refused when it links to outside content", rc == 1 and "outside" in err, err)
    def patched(src_name, out_name, part, old, new):
        with zipfile.ZipFile(d / src_name) as src, zipfile.ZipFile(d / out_name, "w") as dst:
            for i in src.infolist():
                data = src.read(i)
                if i.filename == part:
                    data = data.replace(old, new)
                dst.writestr(i, data)

    patched("all.docx", "ext2.docx", "word/_rels/document.xml.rels", b"</Relationships>",
            b"<Relationship Id='rY' Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/image' "
            b"Target='file:///etc/passwd' targetmode='EXTERNAL'/></Relationships>")
    rc, _, err = run("check", d / "ext2.docx", d / "forbid.json")
    check("check: outside link refused however it is written (N1)", rc == 1 and "outside" in err, err)

    patched("all.docx", "link.docx", "word/_rels/document.xml.rels", b"</Relationships>",
            b'<Relationship Id="rZ" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" '
            b'Target="https://www.transmare.com" TargetMode="External"/></Relationships>')
    rc, _, err = run("check", d / "link.docx", d / "forbid.json")
    check("check: a normal web link is allowed", rc == 0, err)

    patched("all.docx", "field.docx", "word/document.xml", b"</w:body>",
            b'<w:p><w:r><w:instrText xml:space="preserve"> INCLUDEPICTURE "http://evil.example/x.png" </w:instrText></w:r></w:p></w:body>')
    rc, _, err = run("check", d / "field.docx", d / "forbid.json")
    check("check: a field that includes outside content is refused (N1)", rc == 1 and "field" in err, err)

    props = Document(); props.add_paragraph("clean text"); props.core_properties.author = "Acme Chemicals GmbH"
    props.save(d / "props.docx")
    rc, _, err = run("check", d / "props.docx", d / "forbid.json")
    check("check: supplier in the document properties is refused (N2)", rc == 1 and "supplier" in err, err)

    punct = Document(); punct.add_paragraph("Produced by ACME-Chemicals, G.m.b.H."); punct.save(d / "punct.docx")
    rc, _, err = run("check", d / "punct.docx", d / "forbid.json")
    check("check: supplier written with punctuation is refused (N2)", rc == 1 and "supplier" in err, err)

    # security review R3 (26/10/2026): "OK as is" removes the yellow but not the neutralizer's note
    note = Document(); note.add_paragraph("[results table may be incomplete - compare with the original]")
    note.save(d / "note.docx")
    rc, _, err = run("check", d / "note.docx", d / "forbid.json")
    check("check: a neutralizer note left in the text is refused, yellow or not (R3)", rc == 1 and "note" in err, err)

    other = Document(); other.add_paragraph("Acme is a common word; chemicals too."); other.save(d / "other.docx")
    rc, _, err = run("check", d / "other.docx", d / "forbid.json")
    check("check: parts of the name alone do not trigger", rc == 0, err)

    # name and content agree (Willem 19/10): each value of the file name is printed
    (d / "shown.json").write_text(json.dumps([["product name", "Sodium Benzoate"], ["batch", "SB-1"]]))
    shown = Document(); shown.add_paragraph("Product: SODIUM BENZOATE"); shown.add_paragraph("Lot SB-1"); shown.save(d / "shown.docx")
    rc, _, err = run("check", d / "shown.docx", d / "forbid.json", d / "shown.json")
    check("check: passes when every value of the file name is printed", rc == 0, err)
    (d / "notshown.json").write_text(json.dumps([["product name", "Benzoate Pure"]]))
    rc, _, err = run("check", d / "shown.docx", d / "forbid.json", d / "notshown.json")
    check("check: refused when the file name has a value the document does not print", rc == 1 and "not printed" in err, err)
    (d / "part.json").write_text(json.dumps([["product name", "Sodium Benzo"]]))
    rc, _, err = run("check", d / "shown.docx", d / "forbid.json", d / "part.json")
    check("check: part of a printed word does not count (review M2)", rc == 1 and "not printed" in err, err)
    (d / "cyr.json").write_text(json.dumps([["product name", "Ацетон"]]))
    rc, _, err = run("check", d / "shown.docx", d / "forbid.json", d / "cyr.json")
    check("check: a value in non-Latin letters cannot slip through unchecked (review M2)", rc == 1 and "cannot be checked" in err, err)

    patched("all.docx", "tracked.docx", "word/document.xml", b"</w:body>",
            b'<w:p><w:del w:id="9" w:author="x"><w:r><w:delText>Acme Chemicals GmbH</w:delText></w:r></w:del></w:p></w:body>')
    rc, _, err = run("check", d / "tracked.docx", d / "forbid.json")
    check("check: tracked changes are refused, deleted supplier text included (N1)", rc == 1 and "tracked" in err, err)

    # on-screen edits (Willem 19/10): what changed, and what is not in the supplier original
    from reportlab.pdfgen import canvas as pdfcanvas
    src = pdfcanvas.Canvas(str(d / "source.pdf")); src.drawString(72, 720, "Assay 99.6 % Loss on drying 0.4 %"); src.save()
    before = Document(); before.add_paragraph("Assay 99.6 %"); t = before.add_table(rows=1, cols=1)
    t.cell(0, 0).text = "Loss on drying 0.4 %"; before.save(d / "before.docx")
    after = Document(); after.add_paragraph("Assay 99.9 %"); t = after.add_table(rows=1, cols=1)
    t.cell(0, 0).text = "Loss on drying 0.4 %"; after.add_paragraph("Loss on drying 0.4 %"); after.save(d / "after.docx")
    rc, out, err = run("changes", d / "before.docx", d / "after.docx", d / "source.pdf")
    got = json.loads(out) if rc == 0 else []
    check("changes: each change listed with before and after", rc == 0 and len(got) == 2 and got[0]["before"] == "Assay 99.6 %"
          and got[0]["after"] == "Assay 99.9 %", out + err)
    check("changes: a number not in the original is a warning", got and got[0]["unknown"] == ["99.9"], out)
    check("changes: text taken from the original is not", len(got) > 1 and got[1]["unknown"] == [], out)
    rc, out, err = run("changes", d / "before.docx", d / "before.docx", d / "source.pdf")
    check("changes: nothing changed, nothing listed", rc == 0 and json.loads(out) == [], out + err)
    swap = Document(); swap.add_paragraph("Assay 6.99 %"); t = swap.add_table(rows=1, cols=1)
    t.cell(0, 0).text = "Loss on drying 0.4 %"; swap.save(d / "swap.docx")
    rc, out, err = run("changes", d / "before.docx", d / "swap.docx", d / "source.pdf")
    check("changes: digits moved around is still a new number (6.99 is not 99.6)", rc == 0 and json.loads(out)[0]["unknown"] == ["6.99"], out + err)

    # ---- own on-screen editor (Willem 20/10): a view with numbered text, changes written in ----
    from docx.oxml import parse_xml
    from docx.shared import Pt
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
    ed = Document()
    ed.sections[0].header.paragraphs[0].add_run("TMC header")
    p = ed.add_paragraph("Product: ")
    p.add_run("PRODUCTNAME").font.highlight_color = WD_COLOR_INDEX.YELLOW
    big = p.add_run(" grade"); big.font.size = Pt(14); big.bold = True
    t = ed.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Assay"
    t.cell(0, 1).paragraphs[0].add_run("99.6").font.highlight_color = WD_COLOR_INDEX.YELLOW
    t.cell(1, 0).merge(t.cell(1, 1)).paragraphs[0].add_run("Merged row")
    inner = t.cell(0, 0).add_table(rows=1, cols=1); inner.cell(0, 0).text = "Nested"
    box = '<w:p><w:r><w:t>Box text</w:t></w:r></w:p>'
    ed.add_paragraph()._p.append(parse_xml(
        f'<w:r xmlns:w="{W}" xmlns:mc="{MC}"><mc:AlternateContent><mc:Choice Requires="wps"><w:txbxContent>{box}'
        f'</w:txbxContent></mc:Choice><mc:Fallback><w:txbxContent>{box}</w:txbxContent></mc:Fallback></mc:AlternateContent></w:r>'))
    ed.add_paragraph()                                   # an empty line a person may type in
    ed.add_paragraph()._p.append(parse_xml(f'<w:r xmlns:w="{W}"><w:instrText xml:space="preserve"> PAGE </w:instrText></w:r>'))
    ed.sections[0].footer.paragraphs[0].add_run("Footer line")
    ed.save(d / "ed.docx")

    def view_of(path):
        rc, _, err = run("view", path, d / "view.json")
        return (json.loads((d / "view.json").read_text(encoding="utf-8")) if rc == 0 else None), err

    def flat(blocks, out):
        for b in blocks:
            if b["t"] == "p":
                out.append(b)
                for i in b["items"]:
                    if i["t"] == "box":
                        flat(i["blocks"], out)
            else:
                for row in b["rows"]:
                    for cell in row:
                        flat(cell["blocks"], out)
        return out

    v, err = view_of(d / "ed.docx")
    paras = flat([b for part in v["parts"] for b in part["blocks"]], []) if v else []
    shown = [i for p in paras for i in p["items"] if i["t"] == "r"]
    runs = {i["text"]: i for i in shown}
    check("view: body, table, nested table, merged cell, text box, header and footer all shown",
          all(k in runs for k in ("Product: ", "PRODUCTNAME", "Assay", "99.6", "Nested", "Merged row", "Box text", "TMC header", "Footer line")),
          err + str(sorted(runs)))
    check("view: the text box is shown once (its fallback copy is skipped)", [i["text"] for i in shown].count("Box text") == 1)
    check("view: yellow parts and formatting carried over",
          runs.get("PRODUCTNAME", {}).get("hl") is True and runs.get(" grade", {}).get("b") is True and runs.get(" grade", {}).get("size") == 14,
          str(runs.get(" grade")))
    check("view: a merged cell spans its columns",
          v and any(c.get("span") == 2 for b in v["parts"][1]["blocks"] if b["t"] == "table" for r in b["rows"] for c in r))
    check("view: a field's code is not shown", " PAGE " not in runs, str(sorted(runs)))
    check("view: header first, footer last", v and [p["kind"] for p in v["parts"]] == ["header", "body", "footer"])
    empty = [p for p in paras if p.get("empty")]
    check("view: an empty line is offered for typing", len(empty) >= 1)

    def edit(edits, out="ed_out.docx"):
        (d / "edits.json").write_text(json.dumps(edits), encoding="utf-8")
        if (d / out).exists():
            (d / out).unlink()
        return run("edit", d / "ed.docx", d / "edits.json", d / out)

    def piece(text, b=False, i=False, u=False, hl=False):
        return {"text": text, "b": b, "i": i, "u": u, "hl": hl}

    pn, assay = runs["PRODUCTNAME"], runs["99.6"]
    rc, _, err = edit({"runs": [
        {"id": pn["id"], "was": "PRODUCTNAME", "parts": [piece("Sodium "), piece("Benzoate", b=True)]},
        {"id": assay["id"], "was": "99.6", "parts": [piece("99.6")]},
        {"id": runs["Box text"]["id"], "was": "Box text", "parts": [piece("Box text\nsecond line", i=True)]},
        {"id": runs["TMC header"]["id"], "was": "TMC header", "parts": [piece("TMC header", u=True)]},
    ], "paras": [{"id": empty[0]["id"], "parts": [piece("typed on an empty line")]}]})
    out = Document(str(d / "ed_out.docx")) if rc == 0 else None
    check("edit: the changes are written in", out is not None and "Product: Sodium Benzoate grade" in [x.text for x in out.paragraphs], err)
    if out is not None:
        body = [r for x in out.paragraphs for r in x.runs]
        bold = [r.text for r in body if r.bold]
        check("edit: part of a piece made bold, the rest not", "Benzoate" in bold and "Sodium " not in bold, str(bold))
        check("edit: yellow goes only where the person removed it", yellow_runs(d / "ed_out.docx") == [], str(yellow_runs(d / "ed_out.docx")))
        check("edit: the size of an untouched piece is kept", any(r.text == " grade" and r.font.size == Pt(14) for r in body))
        check("edit: line break and italic in the shown text box",
              out.element.body.xpath('.//*[local-name()="Choice"]//w:r[w:rPr/w:i]/w:br'), out.element.xml[-1200:])
        check("edit: header underlined", out.sections[0].header.paragraphs[0].runs[0].underline is True)
        every = ["".join(t.text or "" for t in x.iter(f"{{{W}}}t")) for x in out.element.body.iter(f"{{{W}}}p")]
        check("edit: text typed on an empty line is added there", "typed on an empty line" in every, str(every))
        check("edit: tables and rows unchanged", len(out.tables) == 1 and len(out.tables[0].rows) == 2)
    rc, _, err = edit({"runs": [{"id": pn["id"], "was": "something else", "parts": [piece("x")]}]})
    check("edit: refused when a piece no longer reads what was shown", rc == 1 and "does not match" in err and not (d / "ed_out.docx").exists(), err)
    rc, _, err = edit({"runs": [{"id": "r999", "was": "", "parts": [piece("x")]}]})
    check("edit: unknown piece refused", rc == 1, err)
    rc, _, err = edit({"runs": [{"id": pn["id"], "was": "PRODUCTNAME", "parts": [piece("a\u0007b")]}]})
    check("edit: control characters refused", rc == 1, err)
    for cp in (0xFE0F, 0x85, 0x202E, 0x200B, 0x3164, 0xE0100, 0x2028):
        rc, _, err = edit({"runs": [{"id": pn["id"], "was": "PRODUCTNAME", "parts": [piece("BA" + chr(cp) + "SF")]}]})
        check(f"edit: invisible character U+{cp:04X} refused (it could split a supplier name, review M1)", rc == 1, err)
    rc, _, err = edit({"runs": [{"id": pn["id"], "was": "PRODUCTNAME", "parts": [piece("Assay\t99.6\nnext")]}]})
    check("edit: tab and line break are text", rc == 0, err)
    (d / "inv.json").write_text(json.dumps({"y0": {"value": "BA" + chr(0xFE0F) + "SF"}}))
    rc, _, err = run("apply", d / "d.docx", d / "inv.json", d / "inv.docx")
    check("apply: invisible characters refused in the yellow-parts list too", rc == 1, err)
    rc, _, err = edit({"runs": [{"id": pn["id"], "was": "PRODUCTNAME", "parts": [{"text": "x", "b": "yes", "i": False, "u": False, "hl": False}]}]})
    check("edit: formatting must be true or false", rc == 1, err)
    rc, _, err = edit({"runs": [{"id": pn["id"], "was": "PRODUCTNAME", "parts": [piece("x")]}, {"id": pn["id"], "was": "PRODUCTNAME", "parts": [piece("y")]}]})
    check("edit: the same piece twice refused", rc == 1, err)
    with_text = next(p for p in paras if not p.get("empty"))
    rc, _, err = edit({"paras": [{"id": with_text["id"], "parts": [piece("x")]}]})
    check("edit: a paragraph with its own text is changed through that text", rc == 1, err)
    rc, _, err = edit({"runs": [], "paras": []})
    check("edit: no changes keeps every yellow part", rc == 0 and sorted(yellow_runs(d / "ed_out.docx")) == ["99.6", "PRODUCTNAME"], err)

    # ---- rows, columns, alignment and letter size (Willem 22/10) ----
    body_tables = [b for b in v["parts"][1]["blocks"] if b["t"] == "table"]
    tb = body_tables[0]
    check("view: a table has its number and may get rows and columns (merged across only)", tb["id"] == "t0" and tb["plain"] is True, str(tb)[:200])

    def edit_out(edits):
        rc, out, err = edit(edits)
        return rc, (json.loads(out) if rc == 0 and out.strip() else []), err

    def widths(t):
        return [int(g.get(f"{{{W}}}w")) for g in t._tbl.xpath("./w:tblGrid/w:gridCol")]

    before_w = sum(widths(Document(str(d / "ed.docx")).tables[0]))
    rc, notes, err = edit_out({"runs": [], "paras": [], "tables": [{"id": "t0", "ops": [
        {"op": "addRow", "at": 1, "from": 0, "ids": ["n0", "n1"]},
        {"op": "addCol", "at": 1, "ids": ["n2", "n3", None]},
    ]}], "cells": [{"id": "n0", "parts": [piece("Density")]}, {"id": "n2", "parts": [piece("Method", b=True)]}]})
    out = Document(str(d / "ed_out.docx")) if rc == 0 else None
    if out is not None:
        t0 = out.tables[0]
        rows = [[c.text for c in r.cells] for r in t0.rows]
        check("table: a new row and a new column, with the text typed in them", len(t0.rows) == 3 and rows[1][0] == "Density" and rows[0][1] == "Method", str(rows))
        check("table: the merged row grows with the new column instead of getting a cell", len(t0.rows[2]._tr.xpath("./w:tc")) == 1, str(rows))
        check("table: the table keeps its width on the page", abs(sum(widths(t0)) - before_w) <= 1, f"{widths(t0)} vs {before_w}")
        check("table: a new cell has no yellow", "Density" not in yellow_runs(d / "ed_out.docx"))
        check("table: new cell text keeps the bold that was asked for", any(r.bold for r in t0.rows[0].cells[1].paragraphs[0].runs))
    else:
        check("table: rows and columns added", False, err)
    rc, _, err = edit({"runs": [], "paras": [], "tables": [{"id": "t0", "ops": [{"op": "addCol", "at": 1, "ids": ["n2", "n3", "n4"]}]}]})
    check("table: refused when the new cells do not match the table", rc == 1, err)
    rc, notes, err = edit_out({"runs": [], "paras": [], "tables": [{"id": "t0", "ops": [{"op": "delRow", "at": 0}]}]})
    check("table: a deleted row is a warning with its text (nothing left out unnoticed)",
          rc == 0 and notes and "Assay" in notes[0]["before"] and "deleted" in notes[0]["unknown"][0], err + str(notes))
    rc, notes, err = edit_out({"runs": [], "paras": [], "tables": [{"id": "t0", "ops": [{"op": "delCol", "at": 1}]}]})
    out = Document(str(d / "ed_out.docx")) if rc == 0 else None
    check("table: a deleted column; the merged row keeps one cell", out is not None and len(out.tables[0].columns) == 1
          and notes and "99.6" in notes[0]["before"], err + str(notes))
    rc, _, err = edit({"runs": [], "paras": [], "tables": [{"id": "t0", "ops": [{"op": "delCol", "at": 0}, {"op": "delCol", "at": 0}]}]})
    check("table: the last column cannot be deleted", rc == 1, err)
    merged = Document(); mt = merged.add_table(rows=2, cols=2); mt.cell(0, 0).merge(mt.cell(1, 0)); mt.cell(0, 1).text = "x"
    merged.save(d / "merged.docx")
    mv, _ = view_of(d / "merged.docx")
    check("view: a table with cells merged down is not offered for rows and columns",
          mv and [b for part in mv["parts"] for b in part["blocks"] if b["t"] == "table"][0]["plain"] is False)
    (d / "edits.json").write_text(json.dumps({"runs": [], "paras": [], "tables": [{"id": "t0", "ops": [{"op": "delRow", "at": 0}]}]}))
    rc, _, err = run("edit", d / "merged.docx", d / "edits.json", d / "merged_out.docx")
    check("table: refused in a table with cells merged down", rc == 1 and "merged" in err, err)

    title_id = next(p["id"] for p in paras if any(i.get("text") == "Product: " for i in p["items"]))
    rc, _, err = edit({"runs": [{"id": runs["Product: "]["id"], "was": "Product: ", "parts": [{**piece("Product: "), "size": 14}]}], "paras": [],
                       "layout": [{"id": title_id, "align": "center"}]})
    out = Document(str(d / "ed_out.docx")) if rc == 0 else None
    check("layout: alignment and letter size written in", out is not None and out.paragraphs[0].alignment == 1
          and out.paragraphs[0].runs[0].font.size == Pt(14), err)
    rc, _, err = edit({"runs": [{"id": runs["Product: "]["id"], "was": "Product: ", "parts": [{**piece("Product: "), "size": 500}]}], "paras": []})
    check("layout: an impossible letter size is refused", rc == 1, err)
    rc, _, err = edit({"runs": [], "paras": [], "layout": [{"id": title_id, "align": "diagonal"}]})
    check("layout: an unknown alignment is refused", rc == 1, err)

    # ---- layout check on the PDF (Willem 22/10): letterhead clear, nothing cut off ----
    from reportlab.lib.pagesizes import A4
    logo = ENGINE / "assets" / "logo.png"
    word = Document(); word.add_paragraph("Safety Data Sheet Sodium Benzoate"); word.save(d / "lay.docx")

    def sheet(name, title_y, title="Safety Data Sheet Sodium Benzoate", footer_text="Page 1 of 1"):
        c = pdfcanvas.Canvas(str(d / name), pagesize=A4)
        w, h = A4
        c.drawImage(str(logo), 0, h - 75, width=w, height=75)        # banner at the top
        c.drawImage(str(logo), 40, 10, width=w - 80, height=40)       # address bar at the bottom
        c.setFont("Helvetica", 12)
        if title:
            c.drawString(60, title_y, title)
        c.drawString(w / 2 - 30, 25, footer_text)                     # the page number may sit there
        c.save()

    sheet("ok.pdf", A4[1] - 120)
    rc, _, err = run("layout", d / "ok.pdf", d / "lay.docx")
    check("layout: a clean page passes (page number on the address bar is fine)", rc == 0, err)
    sheet("over.pdf", A4[1] - 40)
    rc, _, err = run("layout", d / "over.pdf", d / "lay.docx")
    check("layout: text over the TMC banner is refused", rc == 1 and "letterhead" in err, err)
    sheet("cut.pdf", A4[1] - 120, title="Safety Data Sheet Sodium")
    rc, _, err = run("layout", d / "cut.pdf", d / "lay.docx")
    check("layout: a product name cut off in the PDF is refused", rc == 1 and "benzoate" in err, err)
    vals = Document(); vals.add_paragraph("Assay 99.6 Loss on drying 0.4"); vals.add_paragraph("Water 0.4")
    vals.save(d / "vals.docx")
    sheet("vals.pdf", A4[1] - 120, title="Assay 99.6 Loss on drying 0.4")   # the second 0.4 is cut off
    rc, _, err = run("layout", d / "vals.pdf", d / "vals.docx")
    check("layout: a cut-off value is refused, also when the same number is printed elsewhere (review M2)",
          rc == 1 and "water" in err, err)
    fld = Document(); fp = fld.add_paragraph("Safety Data Sheet Sodium Benzoate ")
    fp._p.append(parse_xml(f'<w:fldSimple xmlns:w="{W}" w:instr="NUMPAGES"><w:r><w:t>17</w:t></w:r></w:fldSimple>'))
    fld.save(d / "fld.docx")
    rc, _, err = run("layout", d / "ok.pdf", d / "fld.docx")
    check("layout: a page-number field's old value is not taken for cut-off text", rc == 0, err)

    # rows, columns and alignment never in the TMC header or footer (review M3)
    hdr = Document(); ht = hdr.sections[0].header.add_table(rows=1, cols=2, width=Pt(300)); ht.cell(0, 0).text = "TMC"
    hdr.add_paragraph("Body text"); hdr.save(d / "hdr.docx")
    hv, _ = view_of(d / "hdr.docx")
    head_tables = [b for part in hv["parts"] if part["kind"] == "header" for b in part["blocks"] if b["t"] == "table"]
    check("view: a table in the header is not offered for rows and columns", head_tables and head_tables[0]["plain"] is False, str(head_tables)[:200])
    (d / "edits.json").write_text(json.dumps({"runs": [], "paras": [], "tables": [{"id": head_tables[0]["id"], "ops": [{"op": "delRow", "at": 0}]}]}))
    rc, _, err = run("edit", d / "hdr.docx", d / "edits.json", d / "hdr_out.docx")
    check("table: refused in the TMC header", rc == 1 and "header" in err, err)
    head_p = next(p["id"] for part in hv["parts"] if part["kind"] == "header" for p in flat(part["blocks"], []))
    (d / "edits.json").write_text(json.dumps({"runs": [], "paras": [], "layout": [{"id": head_p, "align": "right"}]}))
    rc, _, err = run("edit", d / "hdr.docx", d / "edits.json", d / "hdr_out.docx")
    check("layout: alignment refused in the TMC header", rc == 1 and "header" in err, err)
    (d / "edits.json").write_text(json.dumps({"runs": [{"id": 5, "was": "", "parts": []}], "paras": []}))
    rc, _, err = run("edit", d / "hdr.docx", d / "edits.json", d / "hdr_out.docx")
    check("edit: an id that is not text is refused without a crash (review L5)", rc == 1 and "Traceback" not in err, err)

    (d / "zip.docx").write_bytes(b"PK\x03\x04 not really a zip")
    rc, _, err = run("check", d / "zip.docx", d / "forbid.json")
    check("check: refused when it is not a real Word file", rc == 1, err)

    if shutil.which("soffice") or shutil.which("libreoffice"):
        rc, _, err = run("preview", d / "d.docx", d / "p.pdf")
        pdf = (d / "p.pdf").read_bytes() if (d / "p.pdf").exists() else b""
        check("preview PDF made with the DRAFT stamp", rc == 0 and pdf.startswith(b"%PDF"), err)
    else:
        print("  skip preview (no LibreOffice)")

print(f"\n{ok} passed, {bad} failed")
sys.exit(1 if bad else 0)
