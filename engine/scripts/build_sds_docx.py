"""Render neutralised SDS elements into a TMC-branded .docx."""
import os, re, shutil, tempfile
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BLACK    = RGBColor(0, 0, 0)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
TMC_BLUE = RGBColor(0x2B, 0x3A, 0x67)

BAND_DARK  = "2B3A67"   # SECTION bar: TMC navy, white text
BAND_LIGHT = "D8DEEA"   # sub-section bar: navy tint, navy text
BAND_TABLE = "D8DEEA"   # table header row

# label | colon | value, as the source sets it
LABEL_W, COLON_W, VALUE_W = Cm(6.3), Cm(0.4), Cm(10.7)


def _clear_body(doc):
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def _style(run, size, bold=False, color=BLACK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Arial"
    rPr = run._r.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.append(rf)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(attr), "Arial")
    return run


def _para(doc, text, size=8, bold=False, space_after=1, color=BLACK, indent=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(0), Pt(space_after), 1.0
    if indent is not None:
        pf.left_indent = Cm(indent)
    _style(p.add_run(text), size, bold, color)
    return p


def _cell_text(cell, text, size=8, bold=False, centre=False, color=BLACK):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    p.paragraph_format.space_after = Pt(1.5)
    if centre:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _style(p.add_run(text), size, bold, color)


def _shade(cell, fill):
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(sh)


def _borders(tbl, color="9AA4BF"):
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), color)
        b.append(e)
    tbl._tbl.tblPr.append(b)




def _fixed_layout(tbl, widths):
    """Word treats cell widths as advisory unless the table layout is fixed and
    a grid is declared, which is why a label column can silently grow."""
    tblPr = tbl._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    grid = tbl._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        tbl._tbl.remove(grid)
    grid = OxmlElement("w:tblGrid")
    for w in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(w.twips)))
        grid.append(col)
    tbl._tbl.insert(list(tbl._tbl).index(tblPr) + 1, grid)
    for row in tbl.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = w


def _band(doc, text, fill, color, size, bold=True, indent=0.15):
    """A full-width filled bar, the way the source sets its headings."""
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    _fixed_layout(t, [Cm(17.4)])
    cell = t.rows[0].cells[0]
    _shade(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1.5)
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.left_indent = Cm(indent)
    _style(p.add_run(text), size, bold, color)
    _row_props(t.rows[0])
    return t


def _row_props(row, header=False):
    """Keep a row intact across pages; optionally repeat it as a header."""
    trPr = row._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    trPr.append(cant)
    if header:
        th = OxmlElement("w:tblHeader")
        trPr.append(th)


_BAND = re.compile(r"^\d+(?:\.\d+)*\.?\s")


def _is_band(row):
    """A numbered sub-section row inside a table (e.g. '14.3. Transport ...'),
    shown in the source as a shaded full-width band."""
    filled = [c for c in row if c.strip()]
    return len(filled) == 1 and bool(_BAND.match(filled[0].strip()))


PT_CM = 2.54 / 72          # one point in cm
HEADER_WIDTH_CM = 17.4     # the header text width, under the logo
BOLD_CHAR_EM = 0.6         # average width of a bold character, in em (on the safe side)


def _header_name_fit(name, largest=16.0, smallest=11.0):
    """-> (font size, number of lines) for the product name in the page header."""
    def per_line(size):
        return max(1, int(HEADER_WIDTH_CM / (size * BOLD_CHAR_EM * PT_CM)))
    size = largest
    while size > smallest and len(name) > per_line(size):
        size -= 0.5
    return size, max(1, -(-len(name) // per_line(size)))


def _set_header_footer(doc, meta):
    sec = doc.sections[0]
    sec.different_first_page_header_footer = False
    hdr = sec.header

    # A long product name must never be cut off under the logo (22/10/2026, Willem): the font
    # shrinks from 16 pt to at least 11 pt, and a name that still needs more lines gets room by
    # moving the page text down.
    name = meta.get("product_name", "") or ""
    size, lines = _header_name_fit(name)
    if lines > 1:
        sec.top_margin = sec.top_margin + Cm((lines - 1) * size * 1.25 * PT_CM)
    p = hdr.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(0)
    _style(p.add_run(name), size, True, TMC_BLUE)

    p2 = hdr.add_paragraph()
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(1)
    _style(p2.add_run(meta.get("doc_kind", "Safety Data Sheet")), 12)

    for line in meta.get("framing", []):
        pf = hdr.add_paragraph()
        pf.paragraph_format.space_before = Pt(0)
        pf.paragraph_format.space_after = Pt(0)
        _style(pf.add_run(line), 6.5)

    # prepend the supplier's document-control line to the footer's page-number line
    if meta.get("control_line"):
        pg = sec.footer.paragraphs[-1]
        run = pg.add_run(meta["control_line"] + "   -   ")
        _style(run, 6.5)
        pg._p.insert(list(pg._p).index(pg.runs[0]._r), run._r)


def build(els, meta, template, output):
    tmp_dir = tempfile.mkdtemp(prefix="tmc_sds_")
    tmp = os.path.join(tmp_dir, "w.docx")
    shutil.copy(str(template), tmp)
    os.chmod(tmp, 0o644)
    doc = Document(tmp)
    _clear_body(doc)
    _set_header_footer(doc, meta)

    if meta.get("version_line"):
        _para(doc, meta["version_line"], 7, space_after=2)
    if meta.get("file_nr"):
        _para(doc, "TMC file nr: " + meta["file_nr"], 7, bold=True, space_after=2)

    i = 0
    while i < len(els):
        e = els[i]
        kind = e["kind"]

        if kind == "kv":                      # group consecutive rows into one table
            run_end = i
            while run_end < len(els) and els[run_end]["kind"] == "kv":
                run_end += 1
            block = els[i:run_end]
            t = doc.add_table(rows=len(block), cols=3)
            t.alignment = WD_TABLE_ALIGNMENT.LEFT
            t.autofit = False
            _fixed_layout(t, [LABEL_W, COLON_W, VALUE_W])
            for ri, kv in enumerate(block):
                c0, cc, c1 = t.rows[ri].cells
                _cell_text(c0, kv["label"])
                _cell_text(cc, ":" if kv["label"] or kv["value"] else "")
                _cell_text(c1, kv["value"])
                if kv.get("yellow"):            # a person checks this row (security review M3)
                    from tmc_common import mark_yellow
                    for cell in (c0, c1):
                        for p in cell.paragraphs:
                            for r in p.runs:
                                mark_yellow(r)
                if kv.get("images"):
                    pic = c1.add_paragraph()
                    pic.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    pic.paragraph_format.space_before = Pt(2)
                    pic.paragraph_format.space_after = Pt(0)
                    for img in kv["images"]:
                        pic.add_run().add_picture(
                            img["path"], width=Cm(img["w_pt"] / 72.0 * 2.54))
                        _style(pic.add_run(" "), 5)
            _para(doc, "", 2, space_after=0)
            i = run_end
            continue

        if kind == "section":
            _para(doc, "", 3, space_after=0)
            _band(doc, e["text"], BAND_DARK, WHITE, 10)
            _para(doc, "", 3, space_after=0)

        elif kind == "sub":
            lvl = e.get("level", 1)
            if lvl == 1:
                _para(doc, "", 2, space_after=0)
                _band(doc, e["text"], BAND_LIGHT, TMC_BLUE, 8.5)
                _para(doc, "", 2, space_after=0)
            else:
                p = _para(doc, e["text"], 8, bold=True, space_after=2,
                          color=TMC_BLUE)
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.keep_with_next = True

        elif kind == "para":
            p = _para(doc, e["text"], 8, indent=(0.6 if e.get("indent") else None))
            if e.get("yellow"):                 # a person checks this line (22/10/2026)
                from tmc_common import mark_yellow
                for r in p.runs:
                    mark_yellow(r)

        elif kind == "image" and e.get("col_groups"):
            groups = e["col_groups"]
            t = doc.add_table(rows=1, cols=len(groups))
            t.alignment = WD_TABLE_ALIGNMENT.LEFT
            t.autofit = False
            _fixed_layout(t, [Cm(17.4 / len(groups))] * len(groups))
            _borders(t)
            _row_props(t.rows[0])
            for ci, group in enumerate(groups):
                cell = t.rows[0].cells[ci]
                cell.text = ""
                par = cell.paragraphs[0]
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                par.paragraph_format.space_before = Pt(1)
                par.paragraph_format.space_after = Pt(1)
                for img in group:
                    par.add_run().add_picture(
                        img["path"], width=Cm(img["w_pt"] / 72.0 * 2.54))
                    _style(par.add_run(" "), 5)
            _para(doc, "", 2, space_after=0)

        elif kind == "image":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_together = True
            for img in e["images"]:
                run = p.add_run()
                # keep the source's own size: GHS pictograms have a legal
                # minimum size, so they are never scaled to fit
                run.add_picture(img["path"], width=Cm(img["w_pt"] / 72.0 * 2.54))
                _style(p.add_run("  "), 8)

        elif kind == "table":
            rows = e["rows"]
            ncol = max(len(r) for r in rows)
            t = doc.add_table(rows=len(rows), cols=ncol)
            t.alignment = WD_TABLE_ALIGNMENT.LEFT
            t.autofit = False
            _fixed_layout(t, [Cm(17.4 / ncol)] * ncol)
            _borders(t)
            centred = ncol > 2          # the source centres multi-column tables
            # Only the first row is marked as a repeating header. Marking two
            # makes LibreOffice drop a body row when the table breaks across a
            # page, so the second header row is handled in the verifier instead
            # of being forced into the layout.
            head_rows = 1
            for ri, row in enumerate(rows):
                band = _is_band(row)
                _row_props(t.rows[ri], header=(ri < head_rows))
                for ci in range(ncol):
                    cell = t.rows[ri].cells[ci]
                    _cell_text(cell, row[ci] if ci < len(row) else "",
                               7.5, ri < head_rows or band,
                               centre=(centred and not band),
                               color=(TMC_BLUE if (band or ri < head_rows) else BLACK))
                    if ri < head_rows or band:
                        _shade(cell, BAND_TABLE)
            _para(doc, "", 2, space_after=0)

        i += 1

    doc.save(tmp)
    out = Path(output)
    if out.exists():
        out.unlink()
    shutil.copy2(tmp, output)
    os.chmod(output, 0o644)
    shutil.rmtree(tmp_dir, ignore_errors=True)
    return output
