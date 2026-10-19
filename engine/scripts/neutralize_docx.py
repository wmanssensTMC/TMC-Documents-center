#!/usr/bin/env python3
"""
Transmare Chemie NV — Document Neutralizer v2 (DOCX output)
============================================================
Reads supplier COA PDFs and one or more commercial reference documents
(PFI, Order Confirmation, Contract de Vente, Invoice), then produces
one TMC-branded COA .docx per supplier COA using the official TMC template.

Key rules:
- Product name is ALWAYS taken from the commercial document, not the COA.
- Only fills fields that are actually present in the supplier COA.
- Interprets the specification string into Limit min / Limit max columns.
- Adds as many test result rows as needed.
- Quality standard carries the specification exactly as the supplier printed it
  (Willem, 16/09/2026); Limit min / max hold numbers only; Method only when the
  COA has its own method column.
- v0.6.0: entity choice (--entity), provenance gate, yellow review marks for
  required fields that stayed empty (draft only, exit 2, PDF after --finalise).
- Scanned / image-based PDFs are automatically OCR'd before processing.
- All output text is forced to explicit black (never white or invisible).
- Test table detection is format-agnostic: works with any column layout
  and any language (EN/NL/FR/DE), not just the exact phrase "TEST PARAMETER".
- Text-based extraction fallback when pdfplumber finds no structured tables.
"""

import sys, os, re, argparse, copy, shutil, subprocess, tempfile, atexit
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor, as_completed

SKILL_DIR    = Path(__file__).parent.parent
TEMPLATE_DOC = SKILL_DIR / "assets" / "TMC_COA_template.docx"

try:
    import pdfplumber
except ImportError:
    sys.exit("Missing: pdfplumber\nRun: pip install pdfplumber --break-system-packages")

try:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    sys.exit("Missing: python-docx\nRun: pip install python-docx --break-system-packages")

sys.path.insert(0, str(Path(__file__).parent))
from tmc_common import (load_entity, entity_code, apply_entity_footer, close_hyphens,  # noqa: E402
                        mark_yellow, finalise, non_clobbering, EXIT_REVIEW)
import provenance  # noqa: E402

# Fields that must always be on a certificate (platform NEU_MOETVELDEN). Empty means
# something did not come over, and he has to see it: yellow, Final locked.
NO_PRODUCT = "[product name not given - check]"
REQUIRED = (("lot", "Lot / batch nr"), ("fabricagedatum", "Production date"),
            ("vervaldatum", "Expiry date"))

_OCR_THRESHOLD = 80

def _is_text_sparse(text):
    return len(text.strip()) < _OCR_THRESHOLD

def _ocr_to_searchable_pdf(path):
    lo_bin = shutil.which("ocrmypdf")
    if not lo_bin:
        return None
    tmp = tempfile.mktemp(suffix="_ocr.pdf", prefix="tmc_ocr_")
    try:
        r = subprocess.run(
            [lo_bin, "--skip-text", "--rotate-pages", "--deskew",
             "--optimize", "0", "--jobs", "2", "--output-type", "pdf", path, tmp],
            capture_output=True, timeout=180)
        if r.returncode == 0 and os.path.exists(tmp) and os.path.getsize(tmp) > 0:
            return tmp
    except Exception:
        pass
    try:
        if os.path.exists(tmp): os.remove(tmp)
    except OSError:
        pass
    return None

def _ocr_to_text_pytesseract(path):
    try:
        from pdf2image import convert_from_path
        import pytesseract
        return "\n".join(pytesseract.image_to_string(img)
                         for img in convert_from_path(path, dpi=200))
    except Exception:
        return ""

COA_KEYWORDS = [
    "certificate of analysis", "certificaat van analyse",
    "certificat d'analyse", "analysenzeugnis", "analysecertificaat",
    "batch certificate", "lot certificate", "quality certificate",
    # Suppliers title the same document many ways. Seen in Willem's own set:
    # Zhejiang Jiahua uses "PRODUCT ANALYSIS REPORT".
    "analysis report", "analysis certificate", "analytical report",
    "certificate of quality", "inspection certificate", "inspection report",
    "test report", "test certificate", "rapport d'analyse",
    "pruefzeugnis", "prüfzeugnis", "qualitaetszeugnis", "qualitätszeugnis",
]
COMMERCIAL_KEYWORDS = {
    "ORDER_CONFIRMATION": ["order confirmation","orderbevestiging","confirmation de commande","sales order"],
    "PFI":               ["pro forma invoice","proforma invoice","facture proforma","pro-forma"],
    "SALES_CONTRACT":    ["contract de vente","sales contract","verkoopcontract","sale and purchase"],
    "INVOICE":           ["commercial invoice","invoice no","invoice number","facture commerciale"],
}

def classify_file(text):
    t = text.lower()
    if any(kw in t for kw in COA_KEYWORDS):
        return "COA"
    best, best_score = "UNKNOWN", 0
    for dtype, kws in COMMERCIAL_KEYWORDS.items():
        score = sum(1 for kw in kws if kw in t)
        if score > best_score:
            best, best_score = dtype, score
    return best if best_score else "UNKNOWN"

def _cell_str(v):
    return str(v).strip() if v is not None else ""

# ---------------------------------------------------------------------------
# PDF repair shim
# Some supplier PDFs (e.g. KG/Ligamed) are valid but written in a way pdfminer
# rejects with "No /Root object!" (startxref on the same line as its offset).
# pypdf reads them fine, so rewrite such a file to a clean temporary PDF and
# parse that instead. The supplier's original file is never modified.
# ---------------------------------------------------------------------------
_REPAIRED = {}

def _repair_pdf(path):
    """Return a path pdfplumber can open, repairing via pypdf when needed."""
    if path in _REPAIRED:
        return _REPAIRED[path]
    try:
        with pdfplumber.open(path) as _pdf:
            _pdf.pages[0]
        _REPAIRED[path] = path
        return path
    except Exception as exc:
        try:
            from pypdf import PdfReader, PdfWriter
        except ImportError:
            raise RuntimeError(
                f"{Path(path).name}: unreadable by pdfminer ({exc}) and pypdf is not "
                f"installed. Run: pip install pypdf --break-system-packages") from exc
        reader = PdfReader(path, strict=False)
        writer = PdfWriter()
        for page in reader.pages:
            writer.add_page(page)
        fd, tmp = tempfile.mkstemp(suffix="_repaired.pdf")
        os.close(fd)
        with open(tmp, "wb") as fh:
            writer.write(fh)
        print(f"  i  Repaired unreadable PDF for parsing: {Path(path).name}")
        _REPAIRED[path] = tmp
        return tmp

def _cleanup_repaired():
    for orig, tmp in _REPAIRED.items():
        if tmp != orig:
            try: os.remove(tmp)
            except OSError: pass
atexit.register(_cleanup_repaired)

def _extract_pdf_raw(path):
    pages, tables, layout = [], [], []
    with pdfplumber.open(_repair_pdf(path)) as pdf:
        for i, page in enumerate(pdf.pages):
            text = page.extract_text() or ""
            pages.append({"page": i+1, "text": text})
            # v0.6.0: the text fallback reads columns by character position, which
            # only works on text that keeps the page's horizontal layout
            try:
                layout.append(page.extract_text(layout=True) or "")
            except Exception:
                layout.append(text)
            for tbl in (page.extract_tables() or []):
                if tbl:
                    tables.append({"page": i+1, "rows": tbl})
    return {"source": path, "pages": pages, "tables": tables,
            "full_text": "\n".join(p["text"] for p in pages),
            "layout_text": "\n".join(layout)}

def extract_pdf(path):
    data = _extract_pdf_raw(path)
    if _is_text_sparse(data["full_text"]):
        ocr_path = _ocr_to_searchable_pdf(path)
        if ocr_path:
            try:
                ocr_data = _extract_pdf_raw(ocr_path)
                if not _is_text_sparse(ocr_data["full_text"]):
                    ocr_data["source"] = path
                    return ocr_data
            finally:
                try: os.remove(ocr_path)
                except OSError: pass
        tess_text = _ocr_to_text_pytesseract(path)
        if len(tess_text.strip()) > len(data["full_text"].strip()):
            return {"source": path, "pages": [{"page":1,"text":tess_text}],
                    "tables": [], "full_text": tess_text}
    return data

def extract_pdf_text_only(path):
    pages = []
    with pdfplumber.open(_repair_pdf(path)) as pdf:
        for i, page in enumerate(pdf.pages):
            pages.append({"page": i+1, "text": page.extract_text() or ""})
    full_text = "\n".join(p["text"] for p in pages)
    if _is_text_sparse(full_text):
        ocr_path = _ocr_to_searchable_pdf(path)
        if ocr_path:
            try:
                ocr_pages = []
                with pdfplumber.open(ocr_path) as pdf:
                    for i, page in enumerate(pdf.pages):
                        ocr_pages.append({"page":i+1,"text":page.extract_text() or ""})
                ocr_text = "\n".join(p["text"] for p in ocr_pages)
                if not _is_text_sparse(ocr_text):
                    return {"source":path,"pages":ocr_pages,"tables":[],"full_text":ocr_text}
            finally:
                try: os.remove(ocr_path)
                except OSError: pass
        tess_text = _ocr_to_text_pytesseract(path)
        if len(tess_text.strip()) > len(full_text.strip()):
            return {"source":path,"pages":[{"page":1,"text":tess_text}],"tables":[],"full_text":tess_text}
    return {"source":path,"pages":pages,"tables":[],"full_text":full_text}

def _quick_classify(path):
    data  = extract_pdf_text_only(path)
    dtype = classify_file(data["full_text"])
    return (path, data, dtype)

# ═══════════════════════════════════════════════════════════════════════════
# TEST TABLE DETECTION & COLUMN MAPPING  (format-agnostic)
# ═══════════════════════════════════════════════════════════════════════════
_CELL_EMPTY = {"-","—","–","\u2013","\u2014","n/a","n.a.","nil","none","/","--","---"}

_HDR_SETS = {
    "param": {
        "parameter","parameters","test parameter","test parameters",
        "characteristic","characteristics","property","properties",
        "description","item","items","test item","analysis","analyte",
        "test","tests","component","specification name","name",
        "kenmerk","kenmerken","eigenschap","beschrijving","omschrijving",
        "caractéristique","caractéristiques","paramètre","paramètres",
        "eigenschaft","merkmal","prüfmerkmal","prüfparameter",
    },
    "result": {
        "result","results","value","values","found","actual","obtained",
        "measured","test result","test results","found value",
        "measured value","analyse result","analysis result","reported value",
        "resultaat","résultat","ergebnis","waarde","gevonden",
    },
    "spec": {
        "specification","specifications","spec","specs",
        "requirement","requirements","limit","limits",
        "standard","standards","criteria","criterion","norm","norms",
        "acceptance criteria","required","expected","target",
        "acceptance criterion","required value","expected value",
        "spécification","specificatie","eis","grenzwert","vereiste",
    },
    "lim_min": {"min","minimum","min.","lower limit","limit min","lim min","lower","min value","lower bound","min spec"},
    "lim_max": {"max","maximum","max.","upper limit","limit max","lim max","upper","max value","upper bound","max spec"},
    "unit":    {"unit","units","uom","u/m","u.o.m.","eenheid","unité","unity","measure","um","unit of measure"},
    "method":  {"method","methods","test method","test methods","norm","reference","methode","prüfmethode","méthode","analysis method","standard method","testing method"},
}

def _classify_hdr(text):
    t = text.strip().lower()
    if not t: return None
    # Pass 1: exact match across all roles (highest priority — prevents "limit min"
    # matching "spec" via substring before reaching "lim_min" exact match)
    for role, kw_set in _HDR_SETS.items():
        if t in kw_set:
            return role
    # Pass 2: substring containment
    for role, kw_set in _HDR_SETS.items():
        for kw in kw_set:
            if kw in t:
                return role
    return None

def _find_test_header_row(rows):
    for row_idx, row in enumerate(rows[:6]):
        if not row: continue
        cells = [_cell_str(c) for c in row]
        col_map = {}
        for col_idx, cell in enumerate(cells):
            role = _classify_hdr(cell)
            if role and role not in col_map:
                col_map[role] = col_idx
        if "param" in col_map and {"result","spec","lim_min","lim_max"} & col_map.keys():
            return (row_idx, col_map)
    return (-1, {})

def _build_test_row(cells, col_map):
    def get(role):
        idx = col_map.get(role)
        if idx is not None and idx < len(cells):
            v = cells[idx].strip()
            return "" if v.lower() in _CELL_EMPTY else v
        return ""

    param  = get("param")
    result = get("result")
    uom    = get("unit")
    method = get("method")

    if "spec" in col_map:
        spec = get("spec")
    elif "lim_min" in col_map or "lim_max" in col_map:
        lmin, lmax = get("lim_min"), get("lim_max")
        if lmin and lmax:   spec = f"{lmin} - {lmax}"
        elif lmin:          spec = f"{lmin} Min."
        elif lmax:          spec = f"{lmax} Max."
        else:               spec = ""
    else:
        p_idx = col_map.get("param", 0)
        r_idx = col_map.get("result", -1)
        others = [cells[i].strip() for i in range(len(cells))
                  if i != p_idx and i != r_idx
                  and cells[i].strip() and cells[i].strip().lower() not in _CELL_EMPTY]
        spec = others[0] if others else ""

    param = param.strip()
    if not param: return None
    if not any(v.strip() for v in [spec, result, uom]): return None

    p_lower = param.lower()
    skip_starts = ("note:","this product","these results","guaranteed","remarks:","remark:",
                   "all specifications","the above","tested by","approved by","authorised",
                   "authorized","certif","the analysis","declaration","per above","the goods",
                   "all tests","we hereby")
    if any(p_lower.startswith(s) for s in skip_starts): return None

    if _classify_hdr(param) is not None:
        if not any(re.search(r'\d', c) for c in cells if c.strip()):
            return None

    return [param, spec, result, uom, method]

# ═══════════════════════════════════════════════════════════════════════════
# COA FIELD PARSER
# ═══════════════════════════════════════════════════════════════════════════
_DEFAULT_COL_MAP = {"param": 0, "unit": 1, "spec": 2, "result": 3, "method": 4}

_KV_MAP = [
    (lambda t: t == "product",                                          "product_name"),
    (lambda t: "product name" in t,                                     "product_name"),
    (lambda t: t == "goods description",                                 "grade"),
    (lambda t: "product grade" in t,                                     "grade"),
    (lambda t: "batch" in t,                                             "batch_lot"),
    (lambda t: "lot" in t and "ballot" not in t and "slot" not in t,   "batch_lot"),
    (lambda t: t in ("production date","manufacturing date","manufacture date",
                     "mfg. date","mfg date","date of manufacture","date of production"), "manufacture_date"),
    (lambda t: t in ("expiry date","retest date","best before","best before date",
                     "use before","exp. date","expiration date","expire date"),          "expiry_date"),
    (lambda t: "cas" in t and len(t) < 20,                              "cas_number"),
    (lambda t: t in ("country of origin","origin","country of manufacture","country of mfg"), "origin"),
]

def _extract_kv_row(cells, fields):
    if len(cells) < 2 or not (cells[1] or "").strip(): return
    key = cells[0].lower().strip().rstrip(":")
    val = cells[1].strip()
    for pred, fkey in _KV_MAP:
        if fkey not in fields and pred(key):
            fields[fkey] = val
            return

def _extract_tests_from_text(full_text):
    lines   = full_text.split("\n")
    results = []

    section_start = -1
    for i, line in enumerate(lines[:80]):
        ll = line.strip().lower()
        if not ll: continue
        has_param  = any(k in ll for k in ["parameter","characteristic","test","description","analysis"])
        has_result = any(k in ll for k in ["result","found","value","actual"])
        has_spec   = any(k in ll for k in ["specification","spec","requirement","limit","standard"])
        if has_param and (has_result or has_spec):
            section_start = i
            break

    if section_start == -1:
        return results

    hline = lines[section_start].strip()
    col_pos_map = {}
    for m in re.finditer(r'\S+', hline):
        word = m.group().strip()
        role = _classify_hdr(word)
        if role and role not in col_pos_map:
            col_pos_map[role] = m.start()

    data_lines = lines[section_start + 1:]

    if len(col_pos_map) >= 2 and "param" in col_pos_map:
        sorted_cols = sorted(col_pos_map.items(), key=lambda x: x[1])
        col_starts  = [pos  for _, pos  in sorted_cols]
        col_roles   = [role for role, _ in sorted_cols]
        col_map_idx = {role: idx for idx, role in enumerate(col_roles)}

        for line in data_lines:
            if not line.strip() or re.match(r'^\s*[-=_*]{3,}', line): continue
            row_cells = []
            for k, start in enumerate(col_starts):
                end  = col_starts[k+1] if k+1 < len(col_starts) else len(line)+1
                cell = line[start:end].strip() if start <= len(line) else ""
                row_cells.append(cell)
            row_data = _build_test_row(row_cells, col_map_idx)
            if row_data: results.append(row_data)

    if not results:
        skip_starts = ("note","remark","this product","the results","page","date:",
                       "approved","tested","authorized","authorised","certif","all results")
        for line in data_lines:
            stripped = line.strip()
            if not stripped or re.match(r'^\s*[-=_*]{3,}', stripped): continue
            parts = re.split(r'\s{2,}|\t+', stripped)
            parts = [p.strip() for p in parts if p.strip()]
            if len(parts) < 2: continue
            param = parts[0]
            if len(param) < 2: continue
            if any(param.lower().startswith(s) for s in skip_starts): continue
            if _classify_hdr(param) is not None: continue
            if   len(parts) == 2: results.append([param,"",parts[1],"",""])
            elif len(parts) == 3:
                is_result = (re.match(r'^[\d\.,]+', parts[2])
                             or parts[2].upper() in {"PASS","FAIL","OK","COMPLIES","CONFORMS","CONFORM","PASSES","FAILS"})
                if is_result: results.append([param, parts[1], parts[2], "", ""])
                else:         results.append([param, "", parts[1], parts[2], ""])
            else: results.append([param, parts[1], parts[2], parts[3], parts[4] if len(parts)>4 else ""])

    return results

def parse_coa_fields(data):
    text  = data["full_text"]
    lines = [l.strip() for l in text.split("\n") if l.strip()]
    fields        = {}
    all_test_rows = []

    for tbl_data in (data.get("tables") or []):
        rows = tbl_data.get("rows") or []
        if not rows: continue

        hdr_row_idx, col_map = _find_test_header_row(rows)

        if hdr_row_idx >= 0:
            for row in rows[hdr_row_idx+1:]:
                if not row: continue
                cells   = [_cell_str(c) for c in row]
                matched = sum(1 for c in cells if c.strip() and _classify_hdr(c) is not None)
                if matched >= 2:
                    _, new_map = _find_test_header_row([row])
                    if new_map: col_map.update(new_map)
                    continue
                row_data = _build_test_row(cells, col_map)
                if row_data: all_test_rows.append(row_data)
        else:
            in_test         = False
            current_col_map = dict(_DEFAULT_COL_MAP)

            for row_idx, row in enumerate(rows):
                if not row: continue
                cells = [_cell_str(c) for c in row]
                if not cells: continue

                if not in_test:
                    _, this_map = _find_test_header_row([row])
                    if this_map:
                        in_test = True
                        current_col_map = this_map
                        continue

                    first_up = cells[0].upper().strip()
                    is_marker = (first_up in {
                        "TEST PARAMETER","TEST PARAMETERS","PARAMETER","PARAMETERS",
                        "CHARACTERISTIC","CHARACTERISTICS","ANALYSIS ITEM","TEST ITEM","ANALYSES",
                    } or "TEST PARAMETER" in first_up or "TEST RESULT" in first_up)
                    if is_marker:
                        in_test = True
                        look_ahead = [row] + ([rows[row_idx+1]] if row_idx+1 < len(rows) else [])
                        _, detected = _find_test_header_row(look_ahead)
                        if detected: current_col_map = detected
                        continue

                    _extract_kv_row(cells, fields)
                else:
                    _, new_map = _find_test_header_row([row])
                    if new_map:
                        current_col_map = new_map
                        continue
                    if any("(ADDITIONAL)" in c.upper() for c in cells): continue
                    non_empty = [c for c in cells if c.strip() and c.strip().lower() not in _CELL_EMPTY]
                    if len(non_empty) < 2: continue
                    row_data = _build_test_row(cells, current_col_map)
                    if row_data: all_test_rows.append(row_data)

    if not all_test_rows:
        all_test_rows = _extract_tests_from_text(data.get("layout_text") or text)

    if all_test_rows:
        fields["results_table"] = all_test_rows

    def _rex(pat):
        m = re.search(pat, text, re.IGNORECASE)
        return m.group(1).strip() if m else None

    if "product_name" not in fields:
        fields["product_name"] = _rex(r"(?:product(?:\s+name)?|chemical\s+name)\s*[:\-]?\s*(.+)")
    if "batch_lot" not in fields:
        fields["batch_lot"] = (
            _rex(r"(?:batch|lot)\s+(?:no\.?|number|#|nr\.?)\s*[:\-]?\s*([A-Z0-9][A-Z0-9\-\/\.]+)")
            or _rex(r"(?:batch|lot)\s*[:\-#]\s*([A-Z0-9][A-Z0-9\-\/\.]+)"))
    fields["batch_lot"] = clean_identifier(fields.get("batch_lot"))
    if "manufacture_date" not in fields:
        fields["manufacture_date"] = _rex(
            r"(?:production\s+date|manufacturing\s+date|manufacture\s+date"
            r"|mfg\.?\s*date|date\s+of\s+manufacture|date\s+of\s+production)"
            r"\s*[:\-]?\s*(.+?)(?:\n|$)")
    if "expiry_date" not in fields:
        fields["expiry_date"] = _rex(
            r"(?:expiry\s+date|retest\s+date|best\s+before|use\s+before"
            r"|exp\.?\s*date|expiration\s+date)\s*[:\-]?\s*(.+?)(?:\n|$)")
    # No guessing: an absent product name stays absent. It comes from the
    # platform (product_naam) or from the commercial document, never from a
    # stray line of the supplier's certificate.
    if _looks_like_label(fields.get("product_name")):
        fields["product_name"] = ""

    fields["_source"] = data.get("source", "")
    return fields

# ═══════════════════════════════════════════════════════════════════════════
# REFERENCE DOCUMENT PARSER
# ═══════════════════════════════════════════════════════════════════════════
def parse_ref_fields(data):
    text   = data["full_text"]
    fields = {}

    for pat in [
        r"(?:pro\s*forma\s+invoice|proforma\s+invoice|invoice\s*(?:no\.?|#))\s*[:\-]?\s*([A-Z0-9\-\.\/]+)",
        r"(?:order\s*(?:no\.?|number|#))\s*[:\-]?\s*([A-Z0-9\-\.\/]+)",
        r"(?:contract\s*(?:no\.?|number|#))\s*[:\-]?\s*([A-Z0-9\-\.\/]+)",
        r"\b(\d{3,}\.[A-Z0-9]{3,})\b",
    ]:
        m = re.search(pat, text, re.IGNORECASE)
        if m:
            fields["doc_ref"] = m.group(1).strip()
            break

    dtype  = classify_file(text)
    labels = {"ORDER_CONFIRMATION":"Order Confirmation","PFI":"Pro Forma Invoice",
              "SALES_CONTRACT":"Contract de Vente","INVOICE":"Commercial Invoice"}
    fields["_doc_type"] = labels.get(dtype, "Reference Document")

    products = []
    for line in text.split("\n"):
        m = re.search(r"\b(\d[\d,\.]+)\s*(KG|MT|LT)\b", line, re.IGNORECASE)
        if not m: continue
        qty_val   = f"{m.group(1)} {m.group(2).upper()}"
        prod_line = re.sub(r"^\d+\s+\S+-\S+\s+", "", line).strip()
        prod_m    = re.match(r"([A-Za-z][^\d]+?)(?:\s+\d)", prod_line)
        if prod_m:
            name = prod_m.group(1).strip()
            if len(name) > 3:
                products.append({"name": name, "quantity": qty_val})
    fields["_products"] = products

    m = re.search(r"(?:country\s*of\s*origin|origine|herkomst)\s*[:\-]?\s*(.+?)(?:\n|$)", text, re.IGNORECASE)
    if m: fields["origin"] = m.group(1).strip()

    return fields

def match_pfi_product(coa_fields, ref_products):
    if not ref_products:
        return ""
    coa_candidates = [(coa_fields.get("product_name") or "").lower(),
                      (coa_fields.get("grade") or "").lower()]
    best_name, best_score = ref_products[0]["name"], 0
    for prod in ref_products:
        ref_words = set(re.findall(r"[a-z0-9]+", prod["name"].lower()))
        for cand in coa_candidates:
            cand_words = set(re.findall(r"[a-z0-9]+", cand))
            ref_sig    = {w for w in ref_words  if len(w) > 2 and not w.isdigit()}
            cand_sig   = {w for w in cand_words if len(w) > 2 and not w.isdigit()}
            score = len(ref_sig & cand_sig)
            if score > best_score:
                best_score, best_name = score, prod["name"]
    if best_score == 0:
        return ""
    return best_name

def match_best_ref_doc(coa_fields, ref_docs):
    if not ref_docs: return {}
    if len(ref_docs) == 1: return ref_docs[0]
    coa_candidates = [(coa_fields.get("product_name") or "").lower(),
                      (coa_fields.get("grade") or "").lower()]
    best_ref, best_score = ref_docs[0], 0
    for ref in ref_docs:
        for prod in ref.get("_products", []):
            ref_words = set(re.findall(r"[a-z0-9]+", prod["name"].lower()))
            ref_sig   = {w for w in ref_words if len(w) > 2 and not w.isdigit()}
            for cand in coa_candidates:
                cand_words = set(re.findall(r"[a-z0-9]+", cand))
                cand_sig   = {w for w in cand_words if len(w) > 2 and not w.isdigit()}
                score = len(ref_sig & cand_sig)
                if score > best_score:
                    best_score, best_ref = score, ref
    return best_ref

# ═══════════════════════════════════════════════════════════════════════════
# SPECIFICATION PARSER  (extended)
# ═══════════════════════════════════════════════════════════════════════════
def _try_parse_spec_str(s):
    m = re.match(r"^([\d\.]+)\s*(?:[-\u2013\u2014~]|to)\s*([\d\.]+)", s)
    if m: return (m.group(1), m.group(2))
    m = re.match(r"^([\d\.]+)\s*[Mm]in\.?\s*$", s)
    if m: return (m.group(1), "")
    m = re.match(r"^[Mm]in\.?\s*([\d\.]+)\s*$", s)
    if m: return (m.group(1), "")
    m = re.match(r"^([\d\.]+)\s*[Mm]ax\.?\s*$", s)
    if m: return ("", m.group(1))
    m = re.match(r"^[Mm]ax\.?\s*([\d\.]+)\s*$", s)
    if m: return ("", m.group(1))
    m = re.match(r"^(?:NLT|not\s+less\s+than)\s*([\d\.]+)", s, re.IGNORECASE)
    if m: return (m.group(1), "")
    m = re.match(r"^(?:NMT|not\s+more\s+than)\s*([\d\.]+)", s, re.IGNORECASE)
    if m: return ("", m.group(1))
    m = re.match(r"^[\u2265>]=?\s*([\d\.]+)\s*$", s)
    if m: return (m.group(1), "")
    m = re.match(r"^[\u2264<]=?\s*([\d\.]+)\s*$", s)
    if m: return ("", m.group(1))
    m = re.match(r"^([\d\.]+)\s*$", s)
    if m: return (m.group(1), m.group(1))
    return ("", "")

_UNIT_RE = r'(?:%|ppm|ppb|[µu]?g(?:/[mMlL]+)?|m[gGlL]|[kK][gG]|[Mm][Ll]|cP|mPa\.?s|[Gg]/[cC][mM]3?)'

def parse_spec(spec_str):
    s = spec_str.strip()
    if not s or s.lower() in ("-","\u2014","n/a","n.a.","nil","none"):
        return ("", "")
    # Try original string
    result = _try_parse_spec_str(s)
    if result != ("", ""):
        return result
    # Strip unit from END: "6.5 g/mL Max." -> "6.5 Max."
    s_trail = re.sub(r'\s*' + _UNIT_RE + r'\s*$', "", s).strip()
    if s_trail and s_trail != s:
        result = _try_parse_spec_str(s_trail)
        if result != ("", ""):
            return result
    # Strip unit BETWEEN number and Min/Max qualifier: "99.5 % Min." -> "99.5 Min."
    s_mid = re.sub(r'([\d\.]+)\s*' + _UNIT_RE + r'\s*', r'\1 ', s).strip()
    if s_mid and s_mid != s:
        result = _try_parse_spec_str(s_mid)
        if result != ("", ""):
            return result
    return ("", "")

# ═══════════════════════════════════════════════════════════════════════════
# MULTI-BATCH SPLITTER, FILENAME BUILDER, PDF CONVERTER
# ═══════════════════════════════════════════════════════════════════════════
def split_coa_by_batch(data):
    pages  = data.get("pages", [])
    tables = data.get("tables", [])
    if len(pages) <= 1:
        return [parse_coa_fields(data)]
    page_fields = []
    for pg in pages:
        pg_tables = [t for t in tables if t["page"] == pg["page"]]
        pg_data   = {"source": data["source"], "pages": [pg],
                     "tables": pg_tables, "full_text": pg["text"]}
        page_fields.append(parse_coa_fields(pg_data))
    batches_seen = []
    for f in page_fields:
        b = f.get("batch_lot")
        if b and b not in batches_seen:
            batches_seen.append(b)
    if len(batches_seen) <= 1:
        return [parse_coa_fields(data)]
    base   = parse_coa_fields(data)
    result = []
    for fields in page_fields:
        if not fields.get("batch_lot") and not fields.get("results_table"): continue
        merged = dict(base)
        merged.update({k: v for k, v in fields.items() if v})
        result.append(merged)
    return result if result else [base]


# Side-by-side batches: one result column per batch (19/10/2026, coa_batches /
# coa_consensus wired in). A value the scan readings do not agree on is never
# chosen: it is printed as DOUBT and marked yellow for a person.
DOUBT = "[unclear in the scan - check the original]"
_DATE_MADE = re.compile(r"manufactur\w*\s*date|production\s*date|prod\.?\s*date|mfg\.?\s*date"
                        r"|date\s*of\s*manufactur", re.I)
_DATE_EXP = re.compile(r"expir\w*|best\s*before|use\s*(by|before)", re.I)


def _column_reading(path):
    """-> (agreed, doubts, issues) or None when the COA has no side-by-side batches.
    agreed maps (page, lot) -> {param: (value, unit, label)}, as coa_consensus."""
    import pdfplumber
    import coa_batches as CB
    with pdfplumber.open(path) as pdf:
        scan = _is_text_sparse("".join(p.extract_text() or "" for p in pdf.pages))
        if not scan:
            from coa_consensus import _param_key
            agreed = {}
            for pno, page in enumerate(pdf.pages):
                found = CB.find_column_batches(page)
                if not found:
                    continue
                lots, centres, top = found
                header = CB.find_header_columns(page)
                per = CB.split_columns(page, lots, centres, top, header["unit_centre"] if header else None)
                for lot, rows in per.items():
                    agreed[(pno, lot)] = {_param_key(r[0]): (r[2], r[3], r[0]) for r in rows}
            return (agreed, [], []) if agreed else None
    from coa_consensus import read_with_consensus
    doubts = []
    agreed, issues = read_with_consensus(path, doubts)
    if not agreed and not any(label is None for _p, _l, label, _u, _r in doubts):
        return None
    return agreed, doubts, issues


def column_batches(data, path):
    """One coa_fields per lot for a side-by-side COA, or None for any other layout.

    -> (batches, held_back, notes, readings). held_back: lots no document is made for
    (conflicting analysis results, or only one scan reading found the lot); a person
    checks them against the original. readings: the texts read from the source, for
    the provenance gate."""
    found = _column_reading(path)
    if not found:
        return None
    agreed, doubts, issues = found
    from coa_consensus import merge_duplicate_lots, _param_key
    per_lot, notes, conflicts = merge_duplicate_lots(agreed)
    held = list(conflicts)
    for _p, lot, label, _u, _r in doubts:
        if label is None and lot not in per_lot and not any(w == f"lot {lot}" for w, _ in held):
            held.append((f"lot {lot}", "only one of the two scan readings found this batch"))
    notes = list(notes) + [(w, m) for w, m in issues if w == "ocr"]

    base = parse_coa_fields(data)
    batches, readings = [], []
    for lot, fields in per_lot.items():
        f = dict(base)
        f.update(batch_lot=lot, manufacture_date=None, expiry_date=None)
        lot_doubts = [(label, unit, reading) for _p, l, label, unit, reading in doubts if l == lot and label]
        unverified = {_param_key(label) for label, _u, reading in lot_doubts if reading is not None}
        rows, yellow = [], set()
        for key, (value, unit, label) in fields.items():
            readings += [label, value, unit]
            if key not in unverified:
                if _DATE_MADE.search(label) and not f["manufacture_date"]:
                    f["manufacture_date"] = value
                    continue
                if _DATE_EXP.search(label) and not f["expiry_date"]:
                    f["expiry_date"] = value
                    continue
            else:
                yellow.add(len(rows))
            rows.append([label, "", value, unit, ""])
        shown = set(fields)
        for label, unit, reading in lot_doubts:
            if reading is None and _param_key(label) not in shown:
                shown.add(_param_key(label))
                readings += [label, unit]
                yellow.add(len(rows))
                rows.append([label, "", DOUBT, unit, ""])
        f["results_table"] = rows
        f["_yellow_results"] = yellow
        batches.append(f)
    return batches, held, notes, [r for r in readings if r]

# ---------------------------------------------------------------------------
# Never invent an identity.
# Product name, TMC file number and lot normally arrive from the platform. When
# they are absent the field is LEFT BLANK for Willem to complete at review --
# it is never guessed from the supplier's own wording, because that is how a
# supplier's company name ends up printed as the TMC product name.
# ---------------------------------------------------------------------------
_LABEL_WORDS = {
    "production","data","date","dates","batch","lot","no","nr","number","expiry",
    "manufacture","manufacturing","mfg","retest","page","certificate","analysis",
    "product","name","result","results","specification","spec","code","item",
    "features","quality","standard","method","unit","parameter","test","sample",
}

def _looks_like_label(value):
    """True when the captured text is a form label or prose, not a real value."""
    v = (value or "").strip()
    if not v:
        return True
    words = re.findall(r"[A-Za-z]+", v.lower())
    if words and all(w in _LABEL_WORDS for w in words):
        return True
    return False

def clean_identifier(value):
    """A lot / batch number: keep only a plausible code, else blank."""
    v = (value or "").strip().strip(":-\u2013 ")
    if _looks_like_label(v):
        return ""
    if len(v) < 3 or len(v) > 40:
        return ""
    if not re.search(r"\d", v):          # every real lot we have seen carries a digit
        return ""
    if len(re.findall(r"[A-Za-z]+", v)) > 3:   # a sentence, not a code
        return ""
    return v

def make_output_stem(doc_ref, product_name, batch_lot, source_path="", code="TMC"):
    """TMC_COA_<TMC file nr>_<product name>_<batch> -- fixed house convention.

    Confirmed by Willem on 16/09/2026, e.g.
        TMC_COA_840.S0850_Betaine_AI_test_BT-3877
    Underscores throughout, no brackets. A part that is not present in the input
    is left out; it is never invented to complete the pattern.
    """
    def sanitize(s):
        s = (s or "").strip()
        s = re.sub(r"[^\w\s\-\.]", "", s)
        s = re.sub(r"\s+", "_", s)
        return s[:40]
    parts = [code or "TMC", "COA"]
    if doc_ref:      parts.append(sanitize(doc_ref))
    if product_name: parts.append(sanitize(product_name))
    if batch_lot:    parts.append(sanitize(batch_lot))
    if len(parts) == 2 and source_path:
        # Nothing identifying was supplied or found. Rather than produce a second
        # bare "TMC_COA" that collides with the last one in the COA folder, fall
        # back to the source file's own name. This identifies the draft without
        # inventing a product name -- review replaces it anyway.
        parts.append(sanitize(Path(source_path).stem))
    return "_".join(parts)

def build_pdf_libreoffice(docx_path, output_dir):
    import shutil as _shutil
    lo_bin = _shutil.which("libreoffice") or _shutil.which("soffice")
    if not lo_bin:
        raise FileNotFoundError("LibreOffice not found in PATH")
    subprocess.run([lo_bin,"--headless","--convert-to","pdf","--outdir",
                    str(output_dir), str(docx_path)], check=True, capture_output=True)
    return str(Path(output_dir) / (Path(docx_path).stem + ".pdf"))

# ═══════════════════════════════════════════════════════════════════════════
# DOCX TEMPLATE FILLER
# ═══════════════════════════════════════════════════════════════════════════
def clone_row(table, source_row):
    new_tr = copy.deepcopy(source_row._tr)
    table._tbl.append(new_tr)
    return table.rows[-1]

def set_c(cell, val, bold=None):
    for p in cell.paragraphs:
        for r in p.runs:
            r.text = ""
    para = cell.paragraphs[0]
    run  = para.runs[0] if para.runs else para.add_run()
    run.text = str(val or "")
    if bold is not None:
        run.bold = bold
    rPr = run._r.get_or_add_rPr()
    for existing_color in rPr.findall(qn("w:color")):
        rPr.remove(existing_color)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "000000")
    rPr.append(color)

def _yellow_cell(cell, text):
    set_c(cell, text)
    for r in cell.paragraphs[0].runs:
        if r.text:
            mark_yellow(r)

def build_docx(coa_fields, ref_fields, product_name, output_path, entity=None, review=()):
    if not TEMPLATE_DOC.exists():
        sys.exit(f"Template not found: {TEMPLATE_DOC}")

    tmp_dir  = tempfile.mkdtemp(prefix="tmc_coa_")
    tmp_path = os.path.join(tmp_dir, "working.docx")
    shutil.copy(str(TEMPLATE_DOC), tmp_path)
    os.chmod(tmp_path, 0o644)
    doc = Document(tmp_path)

    for para in doc.paragraphs:
        if "PRODUCTNAME" in para.text:
            for run in para.runs:
                if "PRODUCTNAME" in run.text:
                    run.text = run.text.replace("PRODUCTNAME", product_name or "")
            if not (product_name or "").strip():
                # No product name yet: leave the title clean rather than
                # trailing a dash into empty space, and mark the gap yellow
                # (v0.6.0: product name is a required field, as on the platform).
                for run in para.runs:
                    run.text = re.sub(r"[\s\u2013\u2014-]+$", "", run.text)
                r = para.add_run("  " + NO_PRODUCT)
                mark_yellow(r)
            break

    tbl0 = doc.tables[0]
    # The template says "Production data"; that is a typo for "Production date"
    # (the platform corrected it in its own export on the same day).
    set_c(tbl0.rows[1].cells[0], "Production date", bold=True)
    vals = [("lot", coa_fields.get("batch_lot")), ("fabricagedatum", coa_fields.get("manufacture_date")),
            ("vervaldatum", coa_fields.get("expiry_date")), ("dossier", ref_fields.get("doc_ref"))]
    for i, (key, val) in enumerate(vals):
        set_c(tbl0.rows[i].cells[1], val or "")
        if key in review:
            # an empty required field gets a visible yellow marker, never a guess
            _yellow_cell(tbl0.rows[i].cells[1], val or "[not on the source - check]")

    tbl1 = doc.tables[1]
    template_row = tbl1.rows[1] if len(tbl1.rows) >= 2 else tbl1.rows[0]
    for row in list(tbl1.rows)[1:]:
        tbl1._tbl.remove(row._tr)

    test_data = coa_fields.get("results_table") or []
    unsure    = coa_fields.get("_yellow_results") or ()
    for n, (param, spec, result, uom, method) in enumerate(test_data):
        lim_min, lim_max = parse_spec(spec)
        unit              = uom if uom.lower() not in _CELL_EMPTY else ""
        # Quality standard = the specification exactly as printed (rule of 16/09).
        # House rule 2: a spaced hyphen inside a number range is written closed up.
        values            = [param, close_hyphens(lim_min), close_hyphens(lim_max), result,
                             unit, method, close_hyphens(spec)]
        new_row           = clone_row(tbl1, template_row)
        for i, val in enumerate(values):
            if i < len(new_row.cells):
                set_c(new_row.cells[i], val)
        if n in unsure and len(new_row.cells) > 3:
            # the scan readings did not agree on this value: a person checks it
            _yellow_cell(new_row.cells[3], result or DOUBT)

    if not test_data:
        clone_row(tbl1, template_row)

    if entity:
        apply_entity_footer(doc, entity)
    doc.save(tmp_path)
    out = Path(output_path)
    try:
        if out.exists(): out.unlink()
    except OSError:
        pass
    shutil.copy2(tmp_path, output_path)
    os.chmod(output_path, 0o644)
    shutil.rmtree(tmp_dir, ignore_errors=True)
    return output_path

# ═══════════════════════════════════════════════════════════════════════════
# MAIN
# ═══════════════════════════════════════════════════════════════════════════
def main():
    ap = argparse.ArgumentParser(description="Transmare Chemie NV -- DOCX COA Neutralizer v2")
    ap.add_argument("files", nargs="+")
    ap.add_argument("--output-dir", default=None)
    ap.add_argument("--product-name", default=None,
                    help="Authoritative product name from the platform (product_naam). Overrules the documents.")
    ap.add_argument("--file-nr", default=None,
                    help="Authoritative TMC file number from the platform (dossier_nummer). Overrules the documents.")
    ap.add_argument("--batch", default=None,
                    help="Authoritative lot / batch number. Overrules the documents.")
    ap.add_argument("--entity", default=None,
                    help="JSON entity record (platform 'entiteiten'); default Transmare Chemie NV.")
    ap.add_argument("--finalise", action="store_true",
                    help="FILES are checked .docx drafts: render their PDF once no yellow is left.")
    args = ap.parse_args()

    if args.finalise:
        bad = [f for f in args.files if not finalise(f)]
        sys.exit(1 if bad else 0)
    entity = load_entity(args.entity)
    code, _known = entity_code(entity)

    print("\nTransmare Document Neutralizer v2 (DOCX)")
    print("=" * 45)

    ocr_available = bool(shutil.which("ocrmypdf"))
    try:
        import pytesseract  # noqa
        tess_available = True
    except ImportError:
        tess_available = False

    if ocr_available:       print("  i  OCR engine : ocrmypdf")
    elif tess_available:    print("  i  OCR engine : pytesseract")
    else:
        print("  !  No OCR engine — scanned PDFs may produce empty output.")
        print("     Run: pip install ocrmypdf --break-system-packages")

    valid_files = [f for f in args.files if os.path.isfile(f)]
    for f in args.files:
        if not os.path.isfile(f):
            print(f"  x  Not found: {f}")

    print(f"\nClassifying {len(valid_files)} file(s) ...")
    classify_results = {}
    with ThreadPoolExecutor(max_workers=min(len(valid_files), 8)) as ex:
        futures = {ex.submit(_quick_classify, f): f for f in valid_files}
        for fut in as_completed(futures):
            fpath, text_data, dtype = fut.result()
            classify_results[fpath] = (text_data, dtype)
            print(f"  {Path(fpath).name}  ->  {dtype}")

    coa_paths = [f for f in valid_files if classify_results[f][1] == "COA"]
    ref_paths = [f for f in valid_files if classify_results[f][1] != "COA"]

    if not coa_paths:
        print("\nFAIL  No certificate of analysis detected in the input.")
        print("      Nothing was written. Classified as:")
        for fpath in valid_files:
            print(f"        {Path(fpath).name}  ->  {classify_results[fpath][1]}")
        print("      If one of these IS a COA, its text could not be read "
              "(scan without OCR, or an unsupported layout).")
        sys.exit(1)

    print(f"\nExtracting table data from {len(coa_paths)} COA(s) ...")
    coa_data_map = {}
    with ThreadPoolExecutor(max_workers=min(len(coa_paths), 8)) as ex:
        futures = {ex.submit(extract_pdf, f): f for f in coa_paths}
        for fut in as_completed(futures):
            fpath               = futures[fut]
            coa_data_map[fpath] = fut.result()

    ref_docs = []
    for fpath in ref_paths:
        text_data, _ = classify_results[fpath]
        ref_docs.append(parse_ref_fields(text_data))

    out_dir   = Path(args.output_dir) if args.output_dir else None
    jobs      = []
    refused   = []          # whole input refused: nothing of it is written
    held_back = []          # one batch held back for a person; the other batches go on
    readings  = {}          # source path -> texts the column reading took from the source
    import coa_batches as CB
    import pdfplumber
    for fpath in coa_paths:
        data = coa_data_map[fpath]
        name = Path(fpath).name
        # "Page 1 of 3" in a one-page file: the input is incomplete (coa_batches).
        with pdfplumber.open(fpath) as pdf:
            real_pages = len(pdf.pages)
        ok, msg = CB.page_marker_check([p.get("text", "") for p in data.get("pages", [])], real_pages)
        if not ok:
            print(f"  x  {name}: {msg}")
            refused.append((name, msg))
            continue
        if msg:
            print(f"  !  {name}: {msg}")

        held  = []
        found = column_batches(data, fpath)
        if found:
            batch_list, held, notes, readings[fpath] = found
            print(f"  i  {name}: {len(batch_list) + len(held)} batch(es) side by side")
            for what, why in notes:
                print(f"  i  {name} {what}: {why}")
            for what, why in held:
                print(f"  x  HELD BACK {name} {what}: {why}. No document made; check the original.")
                held_back.append((f"{name} {what}", why))
        else:
            batch_list = split_coa_by_batch(data)
        # A typed batch number on a COA with several batches would give every batch
        # the same number: refused (19/10/2026). Leave Batch empty instead.
        n_batches = len(batch_list) + len(held)
        if args.batch is not None and args.batch.strip() and n_batches > 1:
            lots = ", ".join([f.get("batch_lot") or "?" for f in batch_list] + [w for w, _ in held])
            msg = (f"this COA covers {n_batches} batches ({lots}); "
                   "leave Batch empty so each batch gets its own document")
            print(f"  x  {name}: {msg}")
            refused.append((name, msg))
            continue
        best_ref   = match_best_ref_doc(parse_coa_fields(data), ref_docs)
        for coa_fields in batch_list:
            product_name = match_pfi_product(coa_fields, best_ref.get("_products", []))
            # The platform is the authority. Anything it supplies overrules the
            # documents; anything it does not supply stays blank for review.
            if args.product_name is not None:
                product_name = args.product_name.strip()
            if args.file_nr is not None:
                best_ref = dict(best_ref); best_ref["doc_ref"] = args.file_nr.strip()
            if args.batch is not None:
                coa_fields = dict(coa_fields); coa_fields["batch_lot"] = args.batch.strip()
            test_count   = len(coa_fields.get("results_table") or [])
            jobs.append((coa_fields, best_ref, product_name, test_count))

    if refused:
        print("\nREFUSED - nothing was written:")
        for name, why in refused:
            print(f"  x  {name}: {why}")
        sys.exit(1)
    if not jobs:
        print("\nFAIL  No batch could be made; every batch was held back (see above).")
        sys.exit(1)

    print(f"\nGenerating {len(jobs)} {code} COA(s) ...")
    pdf_failures = []
    drafts = []

    for coa_fields, best_ref, product_name, test_count in jobs:
        stem  = make_output_stem(best_ref.get("doc_ref",""), product_name,
                                 coa_fields.get("batch_lot",""), coa_fields.get("_source",""),
                                 code=code)
        od    = out_dir if out_dir else Path(coa_fields.get("_source",".")).parent
        od.mkdir(parents=True, exist_ok=True)
        newp, clash = non_clobbering(od / f"{stem}.docx")
        if clash:
            print(f"  !  {clash.name} already exists and is left untouched; written as {newp.name}")
        docx_path = str(newp)
        pdf_path  = str(newp.with_suffix(".pdf"))

        no_product = not (product_name or "").strip()
        review = {k for k, _ in REQUIRED
                  if not (({"lot": coa_fields.get("batch_lot"),
                            "fabricagedatum": coa_fields.get("manufacture_date"),
                            "vervaldatum": coa_fields.get("expiry_date")}[k]) or "").strip()}
        build_docx(coa_fields, best_ref, product_name, docx_path, entity=entity, review=review)

        # HARD RULE: nothing in the output may be invented (gate added to the COA
        # route in v0.6.0; the SDS route had it since 14/09).
        src = coa_data_map.get(coa_fields.get("_source")) or {}
        rows = [row for t in (src.get("tables") or []) for row in (t.get("rows") or []) if row]
        src_texts = [src.get("full_text", "")] + [p.get("text", "") for p in src.get("pages", [])]
        src_texts += [" ".join(c for c in row if c) for row in rows]
        src_texts += [c.replace("\n", " ") for row in rows for c in row if c]
        src_texts += readings.get(coa_fields.get("_source"), [])
        src_texts += [close_hyphens(t) for t in src_texts]
        allowed = [product_name or "", best_ref.get("doc_ref") or "", coa_fields.get("batch_lot") or "",
                   "Production date", "[not on the source - check]", NO_PRODUCT, DOUBT]
        allowed += list(entity.values())
        allowed += [r.get("doc_ref", "") for r in ref_docs] + [p for r in ref_docs for p in r.get("_products", [])]
        try:
            provenance.assert_clean(docx_path, src_texts, allowed, template_path=str(TEMPLATE_DOC),
                                    on_fail=lambda: Path(docx_path).unlink(missing_ok=True))
            print("  ok  provenance: every string traced to the source or a supplied value")
            # The batch this document is for, so the platform names and files it
            # exactly as printed (one document per batch; Willem 19/10).
            import json
            Path(docx_path).with_suffix(".meta.json").write_text(
                json.dumps({"batch": coa_fields.get("batch_lot") or ""}), encoding="utf-8")
        except RuntimeError as exc:
            print(f"  x  {exc}")
            pdf_failures.append((Path(docx_path).name, "provenance gate"))
            continue
        print(f"  ok  DOCX -> {Path(docx_path).name}  ({test_count} test row(s))")
        unsure = coa_fields.get("_yellow_results") or ()
        if review or no_product or unsure:
            names = (["Product name"] if no_product else []) + [lbl for k, lbl in REQUIRED if k in review]
            names += [f"{len(unsure)} unclear scan value(s)"] if unsure else []
            print(f"  !  REVIEW: marked yellow, PDF held back until checked: {', '.join(names)}")
            drafts.append(Path(docx_path).name)
            continue
        leeg = [n for n, v in (("product name", product_name),
                               ("TMC file nr", best_ref.get("doc_ref")),
                               ("lot / batch", coa_fields.get("batch_lot"))) if not (v or "").strip()]
        if leeg:
            print(f"  i  BLANK for review: {', '.join(leeg)}")

        try:
            build_pdf_libreoffice(docx_path, str(od))
            print(f"  ok  PDF  -> {Path(pdf_path).name}")
        except FileNotFoundError:
            pdf_failures.append((Path(docx_path).name, "LibreOffice not found in PATH"))
            print("  x  FAIL  PDF not produced: LibreOffice not found in PATH")
        except Exception as exc:
            pdf_failures.append((Path(docx_path).name, str(exc)))
            print(f"  x  FAIL  PDF not produced: {exc}")

        if os.path.exists(pdf_path) and os.path.getsize(pdf_path) > 0:
            print(f"  ok  PAIR verified: {Path(docx_path).name} + {Path(pdf_path).name}")
        elif not any(f[0] == Path(docx_path).name for f in pdf_failures):
            pdf_failures.append((Path(docx_path).name, "PDF missing or empty after conversion"))
            print("  x  FAIL  PDF missing or empty after conversion")

    if pdf_failures:
        print("\nINCOMPLETE - every COA/SDS must be delivered as .docx AND .pdf.")
        for name, why in pdf_failures:
            print(f"  x  {name}: {why}")
        print("Do not deliver the .docx on its own. Install LibreOffice and re-run.\n")
        sys.exit(1)

    if drafts:
        print(f"\nREVIEW - {len(drafts)} draft(s) wait for a check (yellow marks):")
        for d in drafts:
            print(f"  !  {d}")
        print("Check them in Word, remove the highlight, then run with --finalise <docx>.\n")
    if held_back:
        print(f"\nHELD BACK - {len(held_back)} batch(es) got no document; check the original:")
        for what, why in held_back:
            print(f"  x  {what}: {why}")
    if drafts or held_back:
        sys.exit(EXIT_REVIEW)
    print("\nDone. Every output delivered as .docx + .pdf pair.\n")

if __name__ == "__main__":
    main()
