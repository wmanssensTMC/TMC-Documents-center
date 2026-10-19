"""Shared helpers for the COA, TDS and SDS routes (added in v0.6.0, 15/10/2026).

Why one module: the platform (TMC Sonar Logistics, Documents page) and this skill
must produce the same document. Everything both routes need -- the entity block,
the house rules, the yellow review marks, the output naming, the PDF render -- lives
here once, so a rule changed for one document type is changed for all of them.
"""
import json
import re
import shutil
import subprocess
from pathlib import Path

from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

SKILL_DIR = Path(__file__).parent.parent
DEFAULT_ENTITY = SKILL_DIR / "entity.example.json"

# Exit code for "built, but parts are marked yellow and wait for a person".
# 0 = docx + pdf done, 1 = failure / nothing usable, 2 = draft waiting for review.
EXIT_REVIEW = 2

KNOWN_CODES = {
    "transmare chemie nv": "TMC",
    "transmare chemie": "TMC",
    "tmc gulf": "TMCGULF",
    "pharmalabo": "PHARMALABO",
}


# --------------------------------------------------------------------------- entity
def load_entity(path=None):
    """The TMC entity a document comes out under. Willem, 15/10/2026: the entity
    can be chosen for COA and TDS too, like the MSDS. Without --entity it is
    Transmare Chemie NV, as it always was."""
    p = Path(path) if path else DEFAULT_ENTITY
    ent = json.loads(p.read_text(encoding="utf-8"))
    return {k: v for k, v in ent.items() if isinstance(v, str)}


def entity_code(entity):
    code = (entity.get("code") or "").strip()
    if code:
        return code.upper(), True
    name = (entity.get("naam") or "").strip().lower()
    for key, val in KNOWN_CODES.items():
        if key == name or key in name:
            return val, True
    slug = re.sub(r"[^A-Za-z0-9]+", "", entity.get("naam") or "") or "ENTITY"
    return slug.upper(), False


def is_tmc_nv(entity):
    return entity_code(entity)[0] == "TMC"


def entity_footer_line(entity):
    parts = [entity.get("naam", ""), entity.get("adres", "")]
    tel = (entity.get("telefoon") or "").strip()
    if tel:
        parts.append("T " + tel)
    for k in ("email", "web"):
        if (entity.get(k) or "").strip():
            parts.append(entity[k].strip())
    if (entity.get("btw") or "").strip():
        parts.append("VAT: " + entity["btw"].strip())
    return " · ".join(p.strip() for p in parts if p and p.strip())


def apply_entity_footer(doc, entity):
    """The template footer is a picture of the Transmare Chemie NV address bar plus
    its round approval stamp. For another entity both pictures go (that address and
    that stamp are not theirs) and the entity's own block is written as text.
    For Transmare Chemie NV nothing changes."""
    if is_tmc_nv(entity):
        return False
    for sec in doc.sections:
        for part in (sec.footer, sec.first_page_footer):
            for d in part._element.xpath(".//w:drawing"):
                run = d.getparent()
                run.getparent().remove(run)
        p = sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph()
        r = p.add_run(entity_footer_line(entity))
        r.font.size = Pt(7.5)
        r.font.color.rgb = RGBColor(0x21, 0x33, 0x71)
    return True


# ------------------------------------------------------------------------ house rules
# The coordinator's own rules from the platform (meta/neutraliseer_lessen), copied
# 15/10/2026. They refine the reading rules; they never allow inventing or dropping.
HOUSE_RULES = [
    "When a table out of the source has to be carried as lines of text, separate the "
    "columns with two spaces on each side of the upright bar ( value  |  value ).",
    "A hyphen that the pdf prints with spaces around it belongs to the word or the range "
    "it is part of: write it closed up (CAS, EC, UN numbers, hyphenated names, numeric "
    "ranges).",
    "On a certificate of analysis the supplier's standard declarations under the table "
    "are not wanted: build the sheet from the Analysis details and Parameter test results "
    "tables only.",
    "A label and its value stand on one line; where the pdf pushed the value onto the "
    "next line, join it back with ' : '.",
]

_NUM = r"[<>≤≥]?\d+(?:[.,]\d+)?"


def close_hyphens(s):
    """House rule 2. '4.0 - 10.0 %' -> '4.0-10.0 %', '64742 - 82 - 1' -> '64742-82-1',
    'Ethyl - hexyl' -> 'Ethyl-hexyl'. A dash next to a date ('Nº 04 - 14/10/2026') or
    between two sentences is a separator, not a hyphen, and stays as printed."""
    if not s:
        return s
    out = re.sub(r"(?<=\d) - (?=\d)(?![\d.,]*/)", "-", s)
    out = re.sub(r"(?<=[A-Za-z]) - (?=[a-z])", "-", out)
    return out


def table_line(cells):
    """House rule 1 for a table that has to be flattened into text."""
    return "  |  ".join(c for c in cells)


# ------------------------------------------------------------------------- yellow
def mark_yellow(run):
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def iter_highlighted(doc):
    """Every highlighted run in the whole document, in reading order, each once:
    body text, tables, tables inside tables, text boxes, and the headers and footers of
    every section (first page, even pages). Yields (run, paragraph).

    Extended 19/10/2026 (Willem): before, only body paragraphs and top-level table cells
    were read, so a yellow mark in a header, footer, nested table or text box could pass
    --finalise. A text box is stored twice (DrawingML + a VML fallback): only the version
    Word shows is read, the fallback copy is skipped."""
    from docx.text.paragraph import Paragraph

    containers = [(doc.element.body, doc)]
    seen_parts = set()
    for s in doc.sections:
        for hf in (s.header, s.footer, s.first_page_header, s.first_page_footer,
                   s.even_page_header, s.even_page_footer):
            if hf.is_linked_to_previous:
                continue
            el = hf._element
            if id(el) in seen_parts:
                continue
            seen_parts.add(id(el))
            containers.append((el, hf))

    for element, parent in containers:
        for p_el in element.xpath('.//w:p[not(ancestor::*[local-name()="Fallback"])]'):
            p = Paragraph(p_el, parent)
            for r in p.runs:
                if r.font.highlight_color not in (None, WD_COLOR_INDEX.AUTO) and r.text.strip():
                    yield r, p


def yellow_runs(docx_path):
    """Every highlighted run that is still in the document (see iter_highlighted)."""
    from docx import Document
    return [r.text for r, _p in iter_highlighted(Document(str(docx_path)))]


def black(run):
    rPr = run._r.get_or_add_rPr()
    for c in rPr.findall(qn("w:color")):
        rPr.remove(c)
    col = OxmlElement("w:color")
    col.set(qn("w:val"), "000000")
    rPr.append(col)


# ------------------------------------------------------------------------- naming
def safe_part(x):
    x = re.sub(r'[\\/:*?"<>|]', "", str(x or "").strip())
    return re.sub(r"\s+", "_", x)[:80]


def non_clobbering(path):
    """An existing output is never overwritten: report it and step aside."""
    p = Path(path)
    if not p.exists() and not p.with_suffix(".pdf").exists():
        return p, None
    n = 2
    while True:
        alt = p.with_name(f"{p.stem}_v{n}{p.suffix}")
        if not alt.exists() and not alt.with_suffix(".pdf").exists():
            return alt, p
        n += 1


def render_pdf(docx_path):
    lo = shutil.which("libreoffice") or shutil.which("soffice")
    if not lo:
        raise FileNotFoundError("LibreOffice not found in PATH")
    subprocess.run([lo, "--headless", "--convert-to", "pdf", "--outdir",
                    str(Path(docx_path).parent), str(docx_path)],
                   check=True, capture_output=True, timeout=600)
    pdf = Path(docx_path).with_suffix(".pdf")
    if not pdf.exists() or pdf.stat().st_size == 0:
        raise RuntimeError(f"PDF not produced for {Path(docx_path).name}")
    return pdf


def finalise(docx_path):
    """After the coordinator checked every yellow part in Word (corrected it where
    needed and set the highlight to none), this renders the PDF. While anything is
    still yellow, Final stays locked -- the same rule as on the platform."""
    left = yellow_runs(docx_path)
    if left:
        print(f"  x  still marked yellow ({len(left)}): ")
        for t in left[:20]:
            print(f"       - {t[:90]}")
        print("     Check each one in Word, remove the highlight, then run --finalise again.")
        return None
    pdf = render_pdf(docx_path)
    print(f"  ok PDF  -> {pdf.name}")
    return pdf
