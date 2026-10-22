"""Generate assets/TMC_LABEL_SAMPLE.docx: a SAMPLE label template to test the label flow.

It is NOT a TMC label. Kars' real template replaces it (BRIEF phase 3); until then a label
made with it can never become Final. One label per A4 landscape sheet (Kars, via Willem
20/10/2026).

Placeholders, each in one run:
  {{t.<key>}}       fixed text in the label language, from assets/label_texts.json
  {{<field>}}       a value typed in the form (or taken from a Final COA)
  {{entity.<key>}}  the entity record (Settings > Entities)
A row or paragraph whose value is empty is left out, never filled with anything else.

Run: python scripts/make_label_template.py assets assets/TMC_LABEL_SAMPLE.docx
"""
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY = RGBColor(0x21, 0x33, 0x71)
RED = RGBColor(0xC0, 0x00, 0x00)

assets = Path(sys.argv[1])
out = Path(sys.argv[2])

doc = Document()
s = doc.sections[0]
s.orientation = WD_ORIENT.LANDSCAPE
s.page_width, s.page_height = Cm(29.7), Cm(21.0)
s.left_margin = s.right_margin = Cm(1.5)
s.top_margin = s.bottom_margin = Cm(1.2)

st = doc.styles["Normal"]
st.font.name = "Arial"
st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")


def par(text="", size=12, bold=False, color=None, align=None, after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    if align:
        p.alignment = align
    if text:
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.bold = bold
        if color:
            r.font.color.rgb = color
    return p


def cell_text(cell, text, bold=False, size=12):
    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)


def no_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "nil")
        borders.append(e)
    tbl_pr.append(borders)


par("SAMPLE TEMPLATE, NOT FOR USE", size=20, bold=True, color=RED, align=WD_ALIGN_PARAGRAPH.CENTER)
logo = doc.add_paragraph()
logo.add_run().add_picture(str(assets / "logo.png"), width=Cm(12))

par("{{t.productIdentifier}}", size=10, color=NAVY, after=0)
par("{{productName}}", size=26, bold=True, color=NAVY, after=10)

rows = [
    ("batch", "{{batch}}"),
    ("productionDate", "{{productionDate}}"),
    ("expiryDate", "{{expiryDate}}"),
    ("netWeight", "{{netWeight}} kg"),
    ("grossWeight", "{{grossWeight}} kg"),
]
table = doc.add_table(rows=0, cols=2)
no_borders(table)
for key, value in rows:
    cells = table.add_row().cells
    cells[0].width, cells[1].width = Cm(7), Cm(12)
    cell_text(cells[0], "{{t.%s}}" % key)
    cell_text(cells[1], value, bold=True)

par(after=10)
par("{{t.sdsNote}}", size=10, after=10)
par("{{entity.naam}}", size=10, bold=True, after=0)
par("{{entity.adres}}", size=10, after=0)
par("T {{entity.telefoon}}", size=10, after=0)
par("SAMPLE TEMPLATE, NOT FOR USE", size=14, bold=True, color=RED, align=WD_ALIGN_PARAGRAPH.CENTER)

doc.save(str(out))
print("template ->", out)
