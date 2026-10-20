"""Make one label (Word + PDF) from a label template and the values of the form.

    python make_label.py --template=TMC_LABEL_SAMPLE.docx --fields=fields.json
        --entity=entity.json --lang=EN --name=TMC_LABEL_840.S0850_Betaine_BT-3877_EN
        --output-dir=out

Rules (BRIEF 7.1, 7.10):
  - every text on the label comes from the template, the translation file
    (assets/label_texts.json), the entity record or a value the person typed;
    the provenance gate checks the result and there is no way to switch it off
  - a fixed text without a translation in the label language stops the run
    (never translated or guessed here)
  - an empty value leaves its row or paragraph out; nothing is filled in instead
  - dates are printed dd/mm/yyyy, numbers with the decimal sign of the language
Exit 0 = Word + PDF written, 1 = refused (nothing is left in the output folder).
"""
import argparse
import subprocess
import json
import re
import sys
from datetime import date
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from provenance import assert_clean  # noqa: E402
from tmc_common import load_entity, render_pdf  # noqa: E402

ASSETS = HERE.parent / "assets"
TEXTS = ASSETS / "label_texts.json"
PLACEHOLDER = re.compile(r"\{\{\s*([A-Za-z0-9_.]+)\s*\}\}")
# Fields the form sends and how each is printed. Anything else is refused.
FIELDS = {
    "productName": "text",
    "batch": "text",
    "productionDate": "date",
    "expiryDate": "date",
    "netWeight": "number",
    "grossWeight": "number",
}
ENTITY_KEYS = ("naam", "adres", "telefoon")
SAFE_NAME = re.compile(r"^[A-Za-z0-9][A-Za-z0-9_.\-]{0,200}$")


class LabelError(Exception):
    pass


def load_texts(path=TEXTS):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def fmt(kind, value, decimal):
    value = (value or "").strip()
    if not value:
        return ""
    if kind == "date":
        try:
            d = date.fromisoformat(value)
        except ValueError:
            raise LabelError(f"not a date (yyyy-mm-dd): {value}")
        return d.strftime("%d/%m/%Y")
    if kind == "number":
        if not re.fullmatch(r"\d{1,7}(\.\d{1,3})?", value):
            raise LabelError(f"not a number: {value}")
        return value.replace(".", decimal)
    return value


def resolve(fields, entity, lang, texts):
    """-> {placeholder: printed text}; '' means: leave its row/paragraph out."""
    unknown = sorted(set(fields) - set(FIELDS))
    if unknown:
        raise LabelError(f"unknown field(s): {', '.join(unknown)}")
    code = lang.lower()
    if code not in texts["languages"]:
        raise LabelError(f"no label texts for language {lang}")
    decimal = texts["languages"][code]["decimal"]
    values = {k: fmt(kind, fields.get(k), decimal) for k, kind in FIELDS.items()}
    for k in ENTITY_KEYS:
        values[f"entity.{k}"] = (entity.get(k) or "").strip()
    for key, tr in texts["texts"].items():
        if (tr.get(code) or "").strip():
            values[f"t.{key}"] = tr[code].strip()
    return values


def _paragraphs(doc):
    def walk(container):
        for p in container.paragraphs:
            yield p, None
        for table in container.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p, _ in walk(cell):
                        yield p, row
    yield from walk(doc)
    for section in doc.sections:
        for part in (section.header, section.footer):
            yield from walk(part)


def fill(template, values, out_docx):
    """Fill every placeholder. A placeholder without a value removes its table row (or its
    paragraph); a fixed text (t.*) without a translation stops the run."""
    from docx import Document

    doc = Document(str(template))
    drop_rows, drop_pars = [], []
    for p, row in list(_paragraphs(doc)):
        text = "".join(r.text for r in p.runs)
        keys = PLACEHOLDER.findall(text)
        if not keys:
            continue
        for k in keys:
            if k not in values:
                if k.startswith("t."):
                    raise LabelError(f"no translation for '{k[2:]}' in this language")
                raise LabelError(f"the template asks for an unknown value: {k}")
        if any(not values[k] for k in keys):
            (drop_rows if row is not None else drop_pars).append(row if row is not None else p)
            continue
        new = PLACEHOLDER.sub(lambda m: values[m.group(1)], text)
        # the paragraph keeps the format of its first run
        p.runs[0].text = new
        for r in p.runs[1:]:
            r.text = ""
    for p in drop_pars:
        p._p.getparent().remove(p._p)
    for row in {id(r): r for r in drop_rows}.values():
        row._tr.getparent().remove(row._tr)
    doc.save(str(out_docx))


def make_label(template, fields, entity, lang, name, out_dir, texts=None):
    if not SAFE_NAME.fullmatch(name) or ".." in name:
        raise LabelError("unusable document name")
    texts = texts or load_texts()
    values = resolve(fields, entity, lang, texts)
    for required in ("productName", "batch"):
        if not values[required]:
            raise LabelError(f"{required} is required")
    out_dir = Path(out_dir)
    docx_path = out_dir / f"{name}.docx"
    fill(template, values, docx_path)
    from docx import Document

    left = [t for t in ("".join(r.text for r in p.runs) for p, _ in _paragraphs(Document(str(docx_path)))) if "{{" in t]
    if left:
        docx_path.unlink()
        raise LabelError(f"placeholder left on the label: {left[0][:80]}")
    # HARD RULE: every text traceable to the template, the translation file, the entity or
    # the form. Raises (and deletes the output) when something cannot be traced.
    allowed = [v for v in values.values() if v]
    assert_clean(str(docx_path), [], allowed, on_fail=lambda: docx_path.unlink(missing_ok=True), template_path=str(template))
    pdf = render_pdf(docx_path)
    return docx_path, pdf


def main(argv=None):
    ap = argparse.ArgumentParser(description="Make one TMC label (Word + PDF).")
    ap.add_argument("--template", required=True, help="template file name in engine/assets")
    ap.add_argument("--fields", required=True)
    ap.add_argument("--entity", required=True)
    ap.add_argument("--lang", required=True)
    ap.add_argument("--name", required=True)
    ap.add_argument("--output-dir", required=True)
    a = ap.parse_args(argv)
    out_dir = Path(a.output_dir)
    try:
        # only a template that ships with the engine, never a path from outside
        if not re.fullmatch(r"TMC_LABEL_[A-Z0-9_]+\.docx", a.template):
            raise LabelError("unknown template")
        template = ASSETS / a.template
        if not template.is_file():
            raise LabelError("unknown template")
        fields = json.loads(Path(a.fields).read_text(encoding="utf-8"))
        if not isinstance(fields, dict) or not all(isinstance(v, str) for v in fields.values()):
            raise LabelError("fields must be text values")
        docx_path, pdf = make_label(template, fields, load_entity(a.entity), a.lang, a.name, out_dir)
    except (LabelError, RuntimeError, OSError, ValueError, subprocess.SubprocessError) as err:
        for f in out_dir.glob("*"):
            f.unlink(missing_ok=True)
        print(f"  x  label not made: {err}")
        return 1
    print(f"  ok Word -> {docx_path.name}")
    print(f"  ok PDF  -> {pdf.name}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
