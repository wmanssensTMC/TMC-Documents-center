"""LAPTOP ONLY: run the engine on real dossiers and compare with the TMC versions colleagues made
by hand. Real supplier documents never go into git, GitHub or Vercel (Willem 20/10/2026).

  python tests/compare_real_corpus.py "<OneDrive>/03 COA-SDS-Neutralizer/COA's/1. COA's" <temp output folder>
  python tests/compare_real_corpus.py "<...>/1. COA's" <temp output folder> --ai [--max-calls 30]   (AI off vs on;
      needs AI_READER=on and ANTHROPIC_API_KEY in the environment, and costs real money)

Each dossier folder holds "COA - EXTERN" (supplier COAs), "COA - INTERN" (TMC versions) and the
commercial document. Output and report.json go to the output folder only.
"""
import json, os, re, subprocess, sys, unicodedata
from pathlib import Path

ENGINE = Path(__file__).resolve().parent.parent
PY = sys.executable
FIELDS = [("lot", r"Lot\s*/\s*batch\s*nr"), ("production", r"Production\s*dat[ae]"),
          ("expiry", r"Expiry\s*date"), ("file_nr", r"TMC\s*file\s*nr")]
TOKEN = re.compile(r"[^\W_]+(?:[.,/:-][^\W_]+)*")


def text_of(path):
    path = Path(path)
    if path.suffix.lower() == ".docx":
        from docx import Document
        d = Document(str(path))
        lines = [p.text for p in d.paragraphs]
        for t in d.tables:
            for r in t.rows:
                seen = []
                for c in r.cells:
                    if c.text not in seen:
                        seen.append(c.text)
                lines.append("  ".join(seen))
        return "\n".join(lines)
    import pdfplumber
    with pdfplumber.open(str(path)) as pdf:
        return "\n".join(p.extract_text() or "" for p in pdf.pages)


def fields(text):
    out = {}
    m = re.search(r"Certificate of Analysis\s*[-–]\s*(.+)", text)
    out["product"] = m.group(1).strip() if m else ""
    for key, pat in FIELDS:
        m = re.search(pat + r"\s*[:]?\s*(.*)", text)
        out[key] = m.group(1).strip() if m else ""
    return out


LAYOUT_WORDS = {"page", "of"}


def results_tokens(text):
    """Words and numbers of the results part, split at every punctuation mark, so "PC-V-011" (Word)
    and "PC - V - 011" (read back from a PDF) count the same (26/10/2026)."""
    i = text.find("Parameter test results")
    body = unicodedata.normalize("NFKC", text[i:] if i >= 0 else text)
    return {t.lower() for t in re.findall(r"[^\W_]+", body)} - LAYOUT_WORDS


def norm_name(s):
    # NFKC: a superscript "ᵀᴴ" in a colleague's Word file is the same "TH" (13ᵀᴴ AUGUST = 13TH AUGUST)
    return re.sub(r"[^a-z0-9]", "", unicodedata.normalize("NFKC", s).lower())


# what the engine needs from the laptop's environment for an --ai run; nothing else goes along
# (no ANTHROPIC_BASE_URL or other settings that could send the key elsewhere; security review F11)
_BASE_ENV = ("PATH", "PATHEXT", "SYSTEMROOT", "WINDIR", "COMSPEC", "TEMP", "TMP", "HOME", "USERPROFILE",
             "APPDATA", "LOCALAPPDATA", "PYTHONIOENCODING", "LANG", "LC_ALL", "TESSDATA_PREFIX")
_AI_ENV = ("AI_READER", "ANTHROPIC_API_KEY", "AI_READER_MODEL", "AI_READER_EFFORT", "AI_READER_MAX_CALLS",
           "AI_READER_TIMEOUT_S", "AI_READER_RUN_DEADLINE_S")


def ai_env(env, on):
    keep = _BASE_ENV + (_AI_ENV if on else ())
    out = {k: env[k] for k in keep if env.get(k)}
    if not on:
        out["AI_READER"] = "off"
    return out


def main(root, out_root, env=None, quiet=False, max_calls=None):
    """`max_calls`: the AI reader's calls for the whole run together (--ai); each dossier may ask
    at most what is left, and only its supplier COAs (EXTERN) are named as files the AI may read."""
    root, out_root = Path(root), Path(out_root)
    out_root.mkdir(parents=True, exist_ok=True)
    report = []
    left = max_calls
    for dossier in sorted(p for p in root.iterdir() if p.is_dir()):
        subs = {p.name.lower(): p for p in dossier.iterdir() if p.is_dir()}
        ext = next((p for n, p in subs.items() if "extern" in n), None)
        intern = next((p for n, p in subs.items() if "intern" in n), None)
        commercial = [p for p in dossier.iterdir() if p.is_file() and p.suffix.lower() == ".pdf"]
        if not ext:
            report.append({"dossier": dossier.name, "note": "no supplier COA folder"})
            continue
        inputs = sorted(p for p in ext.rglob("*") if p.suffix.lower() == ".pdf")
        out = out_root / dossier.name
        out.mkdir(parents=True, exist_ok=True)
        run_env = env
        if env is not None and left is not None:
            # only the supplier COAs may go to the AI, never the PFI or order confirmation (F1)
            run_env = dict(env, AI_READER_SOURCES=os.pathsep.join(map(str, inputs)),
                           AI_READER_MAX_CALLS=str(max(0, min(left, int(env.get("AI_READER_MAX_CALLS") or 3)))))
        run = subprocess.run([PY, "-I", str(ENGINE / "scripts" / "neutralize_docx.py"), *map(str, inputs),
                              *map(str, commercial), "--output-dir", str(out)],
                             env=run_env, capture_output=True, text=True, encoding="utf-8", errors="replace")
        (out / "_engine.log").write_text(run.stdout + run.stderr, encoding="utf-8")
        if left is not None and (out / "ai_usage.json").exists():
            try:
                left -= len(json.loads((out / "ai_usage.json").read_text(encoding="utf-8")))
            except ValueError:
                left = 0                               # unreadable count: stop asking
        made = sorted(out.glob("*.docx"))
        humans = sorted(p for p in (intern.rglob("*") if intern else []) if p.suffix.lower() in (".pdf", ".docx"))
        entry = {"dossier": dossier.name, "inputs": [p.name for p in inputs], "exit": run.returncode,
                 "engine": [p.stem for p in made], "tmc": [p.name for p in humans], "pairs": []}
        if run.returncode not in (0, 1, 2) or (run.returncode == 1 and not made):
            entry["error"] = (run.stdout + run.stderr)[-1500:]
        # colleagues' file names often lack the batch: pair on the batch printed in their document
        # first, then on the file name and product words (26/10/2026)
        human_docs = []
        for h in humans:
            try:
                ht = text_of(h)
            except Exception:                      # an unreadable file is left out, not fatal
                continue
            human_docs.append((h, ht, fields(ht)))
        for m in made:
            mt = text_of(m)
            mf = fields(mt)
            best, score = None, -1
            lot = norm_name(mf["lot"])
            for h, ht, hf in human_docs:
                hs = norm_name(h.stem)
                s = ((5 if lot and (lot in norm_name(hf["lot"]) or norm_name(hf["lot"]) in lot) and hf["lot"] else 0)
                     + (3 if lot and lot in hs else 0)
                     + sum(1 for w in re.findall(r"[a-z]{4,}", m.stem.lower()) if w in hs))
                if s > score:
                    best, score, best_text, best_fields = h, s, ht, hf
            pair = {"engine": m.stem, "engine_fields": mf}
            if best is not None and score > 0:
                ht, hf = best_text, best_fields
                he, me = results_tokens(ht), results_tokens(mt)
                pair.update({
                    "tmc": best.name, "tmc_fields": hf,
                    "same_fields": {k: norm_name(mf[k]) == norm_name(hf[k]) for k in mf},
                    "results_recall": round(len(he & me) / len(he), 2) if he else None,
                    "missing_in_engine": sorted(he - me)[:25],
                    "extra_in_engine": sorted(me - he)[:25],
                })
            entry["pairs"].append(pair)
        report.append(entry)
    (out_root / "report.json").write_text(json.dumps(report, indent=1, ensure_ascii=False), encoding="utf-8")
    if quiet:
        return report
    for e in report:
        print(f"\n== {e['dossier']}  exit {e.get('exit')}  engine {len(e.get('engine', []))} doc(s), tmc {len(e.get('tmc', []))}")
        if e.get("error"):
            print("   ERROR", e["error"][-400:])
        for p in e.get("pairs", []):
            print(f"   {p['engine']}\n     vs {p.get('tmc', '(no TMC version found)')}")
            if "same_fields" in p:
                diff = {k: (p['engine_fields'][k], p['tmc_fields'][k]) for k, ok in p["same_fields"].items() if not ok}
                print(f"     results found: {p['results_recall']}  field differences: {diff or 'none'}")
                if p["missing_in_engine"]:
                    print(f"     missing in engine: {p['missing_in_engine'][:12]}")
    pairs = [p for e in report for p in e.get("pairs", []) if p.get("results_recall") is not None]
    drafts = sum(len(e.get("engine", [])) for e in report)
    if pairs:
        same = {k: sum(1 for p in pairs if p["same_fields"].get(k)) for k in ("lot", "production", "expiry", "file_nr")}
        avg = sum(p["results_recall"] for p in pairs) / len(pairs)
        print(f"\nSUMMARY: {drafts} draft(s), {len(pairs)} paired with a TMC version; "
              f"results found {avg:.0%} on average; fields equal (of {len(pairs)}): {same}")
    return report


def compare_ai(root, out_root, max_calls=30):
    """--ai (27/10/2026): every dossier twice, AI reader off (out/rules) and on (out/ai), against the
    TMC versions. Needs AI_READER=on and ANTHROPIC_API_KEY in the laptop's environment; it spends
    real money (counted below), at most `max_calls` calls for the whole run (--max-calls, default 30).
    The key is never printed."""
    env = dict(os.environ)
    if (env.get("AI_READER") or "").strip().lower() != "on" or not (env.get("ANTHROPIC_API_KEY") or "").strip():
        sys.exit("set AI_READER=on and ANTHROPIC_API_KEY first (the AI run costs real money)")
    out_root = Path(out_root)
    rules = main(root, out_root / "rules", ai_env(env, on=False), quiet=True)
    ai = main(root, out_root / "ai", ai_env(env, on=True), quiet=True, max_calls=max_calls)
    keys = ("lot", "production", "expiry")
    gained = lost = differs = paired = 0
    for r_entry, a_entry in zip(rules, ai):
        r_pairs = {p["engine"]: p for p in r_entry.get("pairs", [])}
        for a in a_entry.get("pairs", []):
            r = r_pairs.get(a["engine"])
            if not r or "same_fields" not in a or "same_fields" not in r:
                continue
            paired += 1
            for k in keys:
                filled = a["engine_fields"][k] != r["engine_fields"][k]     # the AI reader wrote this one
                if not r["same_fields"][k] and a["same_fields"][k]:
                    gained += 1
                if r["same_fields"][k] and not a["same_fields"][k]:
                    lost += 1
                if filled and not a["same_fields"][k]:
                    differs += 1
                    print(f"   {a_entry['dossier']} / {a['engine']}: AI {k} differs from the TMC version")
    calls = tokens_in = tokens_out = 0
    usd = 0.0
    for f in (out_root / "ai").glob("*/ai_usage.json"):
        for u in json.loads(f.read_text(encoding="utf-8")):
            calls += 1
            tokens_in += u.get("input_tokens", 0)
            tokens_out += u.get("output_tokens", 0)
            usd += u.get("usd") or 0
    print(f"\nAI SUMMARY: {paired} paired draft(s); fields gained {gained}, lost {lost}; "
          f"AI value differs from the TMC version {differs}; {calls} call(s), {tokens_in} in / {tokens_out} out tokens, "
          f"~${usd:.2f}" + (f" (~${usd / calls:.3f} per call)" if calls else ""))


if __name__ == "__main__":
    if "--ai" in sys.argv:
        rest = [a for a in sys.argv[1:] if a != "--ai"]
        cap = 30
        if "--max-calls" in rest:
            i = rest.index("--max-calls")
            cap = max(0, int(rest[i + 1]))
            del rest[i:i + 2]
        compare_ai(*rest[:2], max_calls=cap)
    else:
        main(sys.argv[1], sys.argv[2])
