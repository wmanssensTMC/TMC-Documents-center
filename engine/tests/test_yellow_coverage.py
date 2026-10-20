"""Yellow marks are found everywhere in the Word draft (19/10/2026, Willem: security review L2).

Before, yellow_runs() read only body paragraphs and top-level table cells, so a yellow mark
in a header, footer, nested table or text box could pass --finalise.
Run:  python tests/test_yellow_coverage.py
"""
import sys, tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / "scripts"))
from docx import Document                              # noqa: E402
from docx.enum.text import WD_COLOR_INDEX              # noqa: E402
from docx.oxml import parse_xml                        # noqa: E402
from tmc_common import finalise, iter_highlighted, yellow_runs  # noqa: E402

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"


def yellow(paragraph, text):
    paragraph.add_run(text).font.highlight_color = WD_COLOR_INDEX.YELLOW


def text_box(paragraph, text):
    """A text box as Word stores it: the shown version (Choice) and a fallback copy."""
    inner = (f'<w:p><w:r><w:rPr><w:highlight w:val="yellow"/></w:rPr><w:t>{text}</w:t></w:r></w:p>')
    paragraph._p.append(parse_xml(
        f'<w:r xmlns:w="{W}" xmlns:mc="{MC}"><mc:AlternateContent>'
        f'<mc:Choice Requires="wps"><w:txbxContent>{inner}</w:txbxContent></mc:Choice>'
        f'<mc:Fallback><w:txbxContent>{inner}</w:txbxContent></mc:Fallback>'
        f'</mc:AlternateContent></w:r>'))


def build(path):
    doc = Document()
    yellow(doc.add_paragraph("Product: "), "BODY")
    outer = doc.add_table(rows=1, cols=1)
    inner = outer.cell(0, 0).add_table(rows=1, cols=1)
    yellow(inner.cell(0, 0).paragraphs[0], "NESTED")
    text_box(doc.add_paragraph(), "TEXTBOX")
    # inside a hyperlink and inside inserted (tracked) text: not direct children of the paragraph
    p = doc.add_paragraph()
    p._p.append(parse_xml(
        f'<w:hyperlink xmlns:w="{W}"><w:r><w:rPr><w:highlight w:val="yellow"/></w:rPr><w:t>LINKED</w:t></w:r></w:hyperlink>'))
    p._p.append(parse_xml(
        f'<w:ins xmlns:w="{W}" w:id="1" w:author="x"><w:r><w:rPr><w:highlight w:val="yellow"/></w:rPr><w:t>INSERTED</w:t></w:r></w:ins>'))
    s = doc.sections[0]
    yellow(s.header.paragraphs[0], "HEADER")
    yellow(s.footer.paragraphs[0], "FOOTER")
    s.different_first_page_header_footer = True
    yellow(s.first_page_header.paragraphs[0], "FIRSTPAGE")
    doc.save(path)


ok = bad = 0
def check(name, cond, info=""):
    global ok, bad
    print(("  ok  " if cond else "  XX  ") + name + ("" if cond else f"\n{info}"))
    ok, bad = ok + bool(cond), bad + (not cond)


with tempfile.TemporaryDirectory() as d:
    d = Path(d)
    build(d / "draft.docx")
    found = yellow_runs(d / "draft.docx")
    expected = ["BODY", "NESTED", "TEXTBOX", "LINKED", "INSERTED", "HEADER", "FOOTER", "FIRSTPAGE"]
    check("finds yellow in body, nested table, text box, hyperlink, tracked insert, header, footer, first page",
          sorted(found) == sorted(expected), str(found))
    check("a text box counts once (its fallback copy is skipped)", found.count("TEXTBOX") == 1, str(found))
    check("finalise refuses while any of them is yellow", finalise(d / "draft.docx") is None)

    doc = Document(str(d / "draft.docx"))
    for r, _p in list(iter_highlighted(doc)):
        if r.text != "HEADER":
            r.font.highlight_color = None
    doc.save(d / "header_left.docx")
    check("a yellow mark left only in the header still blocks", yellow_runs(d / "header_left.docx") == ["HEADER"],
          str(yellow_runs(d / "header_left.docx")))

    doc = Document(str(d / "draft.docx"))
    for r, _p in list(iter_highlighted(doc)):
        r.font.highlight_color = None
    doc.save(d / "clean.docx")
    check("after every mark is cleared nothing is left", yellow_runs(d / "clean.docx") == [])

print(f"\n{ok} passed, {bad} failed")
sys.exit(1 if bad else 0)
