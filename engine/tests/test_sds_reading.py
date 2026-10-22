#!/usr/bin/env python3
"""How the MSDS reader reads real-world PDFs (20/10/2026, learned on Willem's real MSDS on the
laptop; every case here is a synthetic copy, no real document is in git).

  1. letters printed twice to fake bold are read once
  2. a letter in another font, a point lower, stays in its word ("l'écart")
  3. a page number line is not glued into a paragraph
  4. a box drawn around a whole section is a frame, not a table: 1.3 and 1.4 are found
  5. "SECTION 2: ..." is a section heading also when it is not bold
  6. the emergency number on the 1.4 heading line itself is replaced by TMC's
Run:  python tests/test_sds_reading.py      (needs reportlab)
"""
import json, sys, tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(ROOT / "scripts"))
try:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
except ImportError:
    sys.exit("SKIP: reportlab not installed")
import pdfplumber  # noqa: E402
from extract_sds import _lines, extract  # noqa: E402
from neutralise_sds import find_supplier_block, neutralise  # noqa: E402

ok = bad = 0
def check(name, cond, info=""):
    global ok, bad
    print(("  ok  " if cond else "  XX  ") + name + ("" if cond else f"\n{str(info)[:1500]}"))
    ok, bad = ok + bool(cond), bad + (not cond)


W, H = A4
with tempfile.TemporaryDirectory() as d:
    d = Path(d)

    # 1-3: one page with fake bold, a split French word and a page number between paragraphs
    p = d / "reading.pdf"
    c = canvas.Canvas(str(p), pagesize=A4)
    c.setFont("Helvetica-Bold", 12)
    for dx in (0, 0.3):                                   # the same title drawn twice, 0.3 pt apart
        c.drawString(40 + dx, H - 150, "ORGAWHITE 2000")
    c.setFont("Helvetica", 10)
    c.drawString(40, H - 200, "Tenir à l")
    x = 40 + c.stringWidth("Tenir à l", "Helvetica", 10)
    c.setFont("Times-Roman", 10); c.drawString(x, H - 201, "’é")         # other font, 1 pt lower
    x += c.stringWidth("’é", "Times-Roman", 10)
    c.setFont("Helvetica", 10); c.drawString(x, H - 200, "cart de la chaleur.")
    c.drawString(40, H - 260, "1 / 1")                       # its own page count
    c.save()
    with pdfplumber.open(str(p)) as pdf:
        texts = [ln["text"] for ln in _lines(pdf.pages[0])]
    check("fake bold is read once", "ORGAWHITE 2000" in texts, texts)
    check("a letter in another font stays in its word", any("l’écart de la chaleur." in t for t in texts), texts)
    check("a page number line is not read as content", "1 / 1" not in texts, texts)

    # 4-6: every section in a drawn box, 1.4 number on the heading line, plain SECTION lines
    p = d / "boxed.pdf"
    c = canvas.Canvas(str(p), pagesize=A4)
    c.setFont("Helvetica", 9)
    y = H - 130
    lines = ["SECTION 1: IDENTIFICATION OF THE SUBSTANCE/MIXTURE AND OF THE COMPANY/UNDERTAKING",
             "1.1 Product identifier: BEST FRIEND",
             "1.3 Details of the supplier of the safety data sheet:",
             "EXAMPLE FRAGRANCE S.L", "C/ Example 19, 08640 Olesa", "Phone: +34 900 000 000 info@example-fragrance.com",
             "1.4 Emergency telephone number: +34 900000000 (Mon-Thu: 8-17h)"]
    top = y + 12
    for t in lines:
        c.drawString(46, y, t); y -= 14
    c.rect(40, y + 8, W - 80, top - y - 8)                  # the box around section 1
    y -= 20
    top2 = y + 12
    c.drawString(46, y, "SECTION 2: HAZARDS IDENTIFICATION"); y -= 14
    c.drawString(46, y, "2.1 Classification of the substance or mixture: not classified"); y -= 14
    c.rect(40, y + 8, W - 80, top2 - y - 8)
    c.save()
    els, meta = extract(str(p))
    kinds = [(e["kind"], (e.get("text") or "")[:30]) for e in els]
    check("a box around a section is not a table", not any(e["kind"] == "table" for e in els), kinds)
    check("'SECTION 2: ...' is a section heading", any(e["kind"] == "section" and e["text"].startswith("SECTION 2") for e in els), kinds)
    lo, hi = find_supplier_block(els)
    check("1.3 found inside the box, block ends at 1.4", lo is not None and [e.get("text") for e in els[lo:hi]] ==
          ["EXAMPLE FRAGRANCE S.L", "C/ Example 19, 08640 Olesa", "Phone: +34 900 000 000 info@example-fragrance.com"], kinds)
    ent = json.loads((ROOT / "entity.example.json").read_text(encoding="utf-8"))
    out, _meta, rep = neutralise(els, meta, ent, "Best Friend", None)
    body = "\n".join(e.get("text", "") for e in out)
    check("the supplier's emergency number on the 1.4 line is replaced", "+34 900000000" not in body and ent["nood"] in body, body)
    check("the 1.4 heading and TMC's number are two lines (each traceable)",
          any(e.get("text") == "1.4 Emergency telephone number:" for e in out), body)
    check("no supplier identity left", "example-fragrance" not in body.lower() and "EXAMPLE FRAGRANCE" not in body, body)

# MSDS round 2 (22/10/2026), each a synthetic copy of a real case
from extract_sds import _tidy_number, _furniture, _open_sentence  # noqa: E402
from neutralise_sds import harvest_identity  # noqa: E402
check("'1 .3 Details' is read as '1.3 Details'; '28 Rue' and '12.5 kg' stay",
      [_tidy_number(t) for t in ("1 .3 Details", "2 .1 .1 X", "28 Rue de Renory", "12.5 kg")]
      == ["1.3 Details", "2.1.1 X", "28 Rue de Renory", "12.5 kg"])
els = [{"kind": "section", "text": "1 .3 Details of the supplier of the safety data sheet"},
       {"kind": "para", "text": "Supplier : Example Fils"}, {"kind": "section", "text": "1 .4 Emergency telephone number"},
       {"kind": "para", "text": "Emergency telephone number : NCEC (regional) +44 1235 000000"},
       {"kind": "para", "text": "Other Emergency Numbers : NCEC (Europe) +44 1235 000001"},
       {"kind": "section", "text": "[x] 2 HAZARDS IDENTIFICATION"}]
out, _m, _r = neutralise(els, {}, json.loads((ROOT / "entity.example.json").read_text(encoding="utf-8")), None, None)
body = "\n".join(e.get("text", "") for e in out)
check("'1 .3' with a space before the dot is found", "Example Fils" not in body, body)
check("all of 1.4 becomes TMC's number only, the supplier's other numbers go too",
      "NCEC" not in body and "Emergency telephone number :" in body and "58 45 45" in body, body)
strong, _weak = harvest_identity([{"kind": "para", "text": "Company name Example PharmaChemikalien GmbH"},
                                  {"kind": "kv", "label": "Telephone", "value": "+49 000 111-0"}], 0, 2)
check("the company name is taken without its label, so it is also found in a disclaimer",
      "Example PharmaChemikalien GmbH" in strong, strong)
lines = [[{"text": "Product X SDS UK", "top": 40, "x0": 40}, {"text": f"text {i}", "top": 400, "x0": 40}] for i in range(4)]
check("a line repeated at the top of most pages is page furniture", "product x sds uk" in _furniture(lines, [842] * 4))
check("a table value at the bottom of a page does not continue on the next page",
      not _open_sentence("Marine water 0.01 mg/l"))
check("an unfinished sentence does", _open_sentence("If exposure limits have not been established and"))

# security review of round 2 (22/10/2026)
els = [{"kind": "para", "text": "1.3 Supplier"}, {"kind": "para", "text": "Example SA"},
       {"kind": "para", "text": "1.4 Emergency telephone number"},
       {"kind": "para", "text": "+32 4 000 00 00 (supplier)"},
       {"kind": "para", "text": "Available outside office hours."},
       {"kind": "table", "rows": [["Poison centre", "+44 1235 000 000"]]},
       {"kind": "section", "text": "SECTION 2: HAZARDS"}]
out, _m, rep = neutralise(els, {}, json.loads((ROOT / "entity.example.json").read_text(encoding="utf-8")), None, None)
texts = [e.get("text", "") for e in out]
check("1.4: lines without a phone number or e-mail stay (review M2)", "Available outside office hours." in texts, texts)
check("1.4: a phone number in a 1.4 table is removed too (review M3)",
      not any("1235" in c for e in out if e["kind"] == "table" for r in e["rows"] for c in r), out)
els = [{"kind": "para", "text": "1.3 Supplier"}, {"kind": "para", "text": "Example SA"},
       {"kind": "para", "text": "1.4 Emergency telephone number"}] + \
      [{"kind": "para", "text": f"Country {n}: +1 555 000 00{n:02d}"} for n in range(12)]
out, _m, rep = neutralise(els, {}, json.loads((ROOT / "entity.example.json").read_text(encoding="utf-8")), None, None)
check("1.4 with no section 2 after it and many lines: kept and marked yellow, never cleared (review M1)",
      sum(1 for e in out if e.get("yellow")) == 12, [(e.get("text"), e.get("yellow")) for e in out])
lines = [[{"text": "Product X SDS UK", "top": 40 + (i * 7 if i % 2 else 0), "x0": 40}, {"text": f"t {i}", "top": 400, "x0": 40}]
         for i in range(4)]
check("a repeated line at different places is not page furniture (review M4)", "product x sds uk" not in _furniture(lines, [842] * 4))
check("'1. 2-Butanone' is left as printed (review L2)", _tidy_number("1. 2-Butanone") == "1. 2-Butanone")
strong, _weak = harvest_identity([{"kind": "para", "text": "Company name:"}, {"kind": "para", "text": "Name Chemicals Ltd"}], 0, 2)
check("a label-only line is skipped and the real name is taken, also as written (review L1)",
      "Name Chemicals Ltd" in strong, strong)

try:
    neutralise([{"kind": "para", "text": "1.3 Supplier"}, {"kind": "para", "text": "Example SA"},
                {"kind": "section", "text": "SECTION 2: HAZARDS"}], {},
               json.loads((ROOT / "entity.example.json").read_text(encoding="utf-8")), None, None)
    check("no section 1.4 found: refused for a person (fail closed)", False)
except ValueError as exc:
    check("no section 1.4 found: refused for a person (fail closed)", "1.4" in str(exc), exc)

# a long product name is never cut off in the page header (22/10/2026, Willem)
from build_sds_docx import build  # noqa: E402
from docx import Document  # noqa: E402
with tempfile.TemporaryDirectory() as d:
    tpl = str(ROOT / "assets" / "TMC_SDS_template.docx")
    base = Document(tpl).sections[0].top_margin
    short, long_ = Path(d) / "short.docx", Path(d) / "long.docx"
    build([{"kind": "para", "text": "x"}], {"product_name": "Best Friend"}, tpl, str(short))
    build([{"kind": "para", "text": "x"}], {"product_name": "TENSAGEX EOC370 / EOC470 / EOC670 / EOC670B / EOC670(MB) / "
                                            "EOC670N / EOC670N(MB) Lauryl Ether Carboxylic Acid"}, tpl, str(long_))
    def head(p):
        sec = Document(str(p)).sections[0]
        run = next(r for par in sec.header.paragraphs for r in par.runs if r.text.strip())
        return run.font.size.pt, sec.top_margin
    s_size, s_margin = head(short)
    l_size, l_margin = head(long_)
    check("a short name keeps 16 pt and the page margin", s_size == 16 and s_margin == base, (s_size, s_margin, base))
    check("a long name shrinks (11 pt at least) and the page text moves down to fit it",
          11 <= l_size < 16 and l_margin > base, (l_size, l_margin, base))

# review F2: the overprint rule follows the letter size; real repeated letters are never merged
with tempfile.TemporaryDirectory() as d:
    p = Path(d) / "sizes.pdf"
    c = canvas.Canvas(str(p), pagesize=A4)
    for size, y in ((22, 700), (6, 650)):
        c.setFont("Helvetica-Bold", size)
        for dx in (0, size * 0.044):                      # fake bold, as measured on a real PDF
            c.drawString(40 + dx, y - dx, "ORGAL")
    c.setFont("Helvetica", 4); c.drawString(40, 600, "Assay 1100 ll 00")   # real repeats, tiny font
    c.save()
    with pdfplumber.open(str(p)) as pdf:
        texts = [ln["text"] for ln in _lines(pdf.pages[0])]
    check("fake bold read once at size 22 and at size 6", texts.count("ORGAL") == 2, texts)
    check("real repeated letters in a 4 pt font stay ('1100', 'll', '00')", "Assay 1100 ll 00" in texts, texts)

# review fixes: real content is never dropped or cut short
from extract_sds import _is_page_number  # noqa: E402
check("'2 / 21' on a 21-page MSDS is a page number", _is_page_number("2 / 21", 21))
check("'1 / 10' on a 21-page MSDS is content (review F1)", not _is_page_number("1 / 10", 21))
check("'Page 3 of 8' is a page number", _is_page_number("Page 3 of 8", 21))
els = [{"kind": "para", "text": "1.3 Supplier"}, {"kind": "table", "rows": [["Company", "Example SA"]]}, {"kind": "para", "text": "1.4 Emergency"}]
try:
    neutralise(els, {}, json.loads((ROOT / "entity.example.json").read_text(encoding="utf-8")), None, None)
    check("a table inside 1.3 is refused for a person (review F4)", False)
except ValueError:
    check("a table inside 1.3 is refused for a person (review F4)", True)
els = [{"kind": "para", "text": "1.3 Supplier"}, {"kind": "para", "text": "Example SA"},
       {"kind": "para", "text": "1.4 Emergency telephone number: (see section 16)"}, {"kind": "section", "text": "SECTION 2: X"}]
out, _m, _r = neutralise(els, {}, json.loads((ROOT / "entity.example.json").read_text(encoding="utf-8")), None, None)
check("'(see section 16)' after 1.4 is not taken for a phone number (review F5)",
      any(e.get("text") == "1.4 Emergency telephone number: (see section 16)" for e in out), [e.get("text") for e in out])

# round 3 (22/10/2026, a real French MSDS): title below the header band, body text that starts
# high on a page, a value that runs over a page break, a row that wraps in both columns at once
from neutralize_sds_docx import _source_corpus  # noqa: E402
from provenance import build_corpus, norm  # noqa: E402
with tempfile.TemporaryDirectory() as d:
    p = Path(d) / "french.pdf"
    c = canvas.Canvas(str(p), pagesize=A4)
    for page in (1, 2, 3):
        c.setFont("Helvetica", 9)
        c.drawString(71, H - 50, "PRODUIT X")                          # running header, every page
        c.drawString(67, H - 78, "Version 2 Date de révision 09.10.2018")
        if page == 1:
            c.setFont("Helvetica-Bold", 18); c.drawString(71, H - 125, "FICHE DE DONNÉES DE SÉCURITÉ")
            c.setFont("Helvetica", 9); c.drawString(71, H - 160, "1.1 Identificateur de produit")
        if page == 2:
            c.drawString(92, H - 700, "Indications pour la protection"); c.drawString(234, H - 700, ":")
            c.drawString(248, H - 700, "Utilisez un équipement protégé.")
            c.drawString(92, H - 712, "contre l'incendie")     # the label wraps under its own value
            c.drawString(248, H - 712, "Utiliser des outils anti-étincelles.")
        if page == 3:
            c.drawString(248, H - 103, "Tenez éloigné des agents réducteurs.")   # starts above 100 pt
            c.drawString(92, H - 300, "Sensibilisation respiratoire"); c.drawString(211, H - 300, "ou")
            c.drawString(234, H - 300, ":")                          # label column wider than 200 pt
            c.drawString(248, H - 300, "Non classé sur la base des")
            c.drawString(92, H - 312, "cutanée"); c.drawString(248, H - 312, "informations disponibles.")
        c.showPage()
    c.save()
    els, meta = extract(str(p))
    texts = [e.get("text") or (e.get("label", "") + " : " + e.get("value", "")) for e in els]
    check("the title printed below the header band is the document's own title",
          meta.get("doc_kind") == "FICHE DE DONNÉES DE SÉCURITÉ", meta.get("doc_kind"))
    check("body text that starts above the usual header band is not dropped",
          any("Tenez éloigné des agents réducteurs." in t for t in texts), texts)
    kv = next((e for e in els if e["kind"] == "kv" and e["label"].startswith("Sensibilisation")), {})
    check("a row that wraps in both columns: each part goes to its own column",
          kv.get("label") == "Sensibilisation respiratoire ou cutanée"
          and kv.get("value") == "Non classé sur la base des informations disponibles.", kv)
    corpus = build_corpus(_source_corpus(str(p)))     # as the gate builds it (review L1)
    check("the gate finds a label wider than the usual label column",
          any(norm("Sensibilisation respiratoire ou cutanée") in t for t in corpus))
    check("the gate finds a value that runs over a page break, without the running header",
          any(norm("Utilisez un équipement protégé. Utiliser des outils anti-étincelles. Tenez éloigné des agents "
                   "réducteurs.") in t for t in corpus))
    check("the gate does not find words joined across the two columns",
          not any(norm("Sensibilisation respiratoire ou cutanée informations disponibles") in t for t in corpus))

from extract_sds import _split_kv  # noqa: E402
words = [{"text": t, "x0": x, "x1": x + 10} for t, x in
         (("Precautionary", 78), ("statements", 136), ("(full", 182), ("text", 199), (":", 216), ("see", 221),
          ("section", 238), ("16)", 269))]
check("a colon inside an open bracket is not a label : value row", _split_kv({"words": words}) is None)
words = [{"text": t, "x0": x, "x1": x + 10} for t, x in (("Flash", 92), ("point", 120), ("(closed", 150), ("cup)", 180),
                                                         (":", 234), ("85", 248), ("°C", 262))]
check("a closed bracket before the colon still makes a label : value row", (_split_kv({"words": words}) or [""])[0] == "Flash point (closed cup)")

# security review of round 3
ent = json.loads((ROOT / "entity.example.json").read_text(encoding="utf-8"))
with tempfile.TemporaryDirectory() as d:
    p = Path(d) / "h1.pdf"
    c = canvas.Canvas(str(p), pagesize=A4)
    c.setFont("Helvetica", 9)
    for y, x, t in ((130, 46, "SECTION 1: IDENTIFICATION"), (150, 46, "1.3 Details of the supplier of the safety data sheet"),
                    (165, 92, "Société"), (165, 234, ":"), (165, 248, "ACME SAS"),
                    (177, 92, "Adresse"), (177, 248, "12 rue X"),            # no colon: not part of the name
                    (195, 46, "1.4 Emergency telephone number: +33 1 23 45 67 89"),
                    (215, 46, "SECTION 2: HAZARDS IDENTIFICATION"), (235, 46, "SECTION 16: OTHER INFORMATION"),
                    (250, 46, "Prepared by ACME SAS regulatory affairs.")):
        c.drawString(x, H - y, t)
    c.save()
    els, meta = extract(str(p))
    out, _m, rep = neutralise(els, meta, ent, None, None)
    body = " ".join(str(e.get(k, "")) for e in out for k in ("text", "label", "value"))
    check("a left-margin line after the company row does not widen the name; the bare name goes everywhere (H1)",
          "ACME" not in body and "ACME SAS" in rep["supplier_tokens"], (body, rep["supplier_tokens"]))
els = [{"kind": "para", "text": "1.3 Supplier"}, {"kind": "para", "text": "Example SA"},
       {"kind": "para", "text": "1.4 Emergency telephone number: +32 3 000 00 00"}, {"kind": "section", "text": "SECTION 2: X"},
       {"kind": "para", "text": "Tenez éloigné des agents réducteurs."}]
out, _m, rep = neutralise(els, {"header_band": ["Tenez éloigné des agents réducteurs."]}, ent, None, None)
check("body text from the top band of a page is marked yellow for a person (M3)",
      any(e.get("yellow") and "Tenez" in e.get("text", "") for e in out), out)

print(f"\n{ok} passed, {bad} failed")
sys.exit(1 if bad else 0)
