"""Labels (phase 3 frame): the SAMPLE template filled from the form, Word + PDF.
Synthetic values only. Run:  python tests/test_labels.py
"""
import json
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
ENGINE = HERE.parent
sys.path.insert(0, str(ENGINE / "scripts"))
import make_label  # noqa: E402
from make_label import LabelError, load_texts, make_label as make  # noqa: E402
from provenance import docx_strings  # noqa: E402
from tmc_common import load_entity  # noqa: E402

TEMPLATE = ENGINE / "assets" / "TMC_LABEL_SAMPLE.docx"
ENTITY = load_entity(ENGINE / "entity.example.json")
FIELDS = {
    "productName": "Betaine AI test",
    "batch": "BT-3877",
    "productionDate": "2026-03-01",
    "expiryDate": "2028-02-29",
    "netWeight": "25.5",
    "grossWeight": "26",
}


def texts_of(docx):
    return [t for _, t in docx_strings(str(docx))]


def run(fields=FIELDS, lang="EN", texts=None):
    out = Path(tempfile.mkdtemp())
    docx, pdf = make(TEMPLATE, dict(fields), ENTITY, lang, f"TMC_LABEL_test_{lang}", out, texts)
    return docx, pdf, out


def refused(fn, words):
    try:
        fn()
    except (LabelError, RuntimeError) as err:
        assert words in str(err), f"refused for another reason: {err}"
        return
    raise AssertionError(f"not refused ({words})")


def test_english_label():
    docx, pdf, _ = run()
    t = texts_of(docx)
    for want in ("Product identifier", "Betaine AI test", "Batch", "BT-3877", "01/03/2026", "29/02/2028",
                 "25.5 kg", "26 kg", "Transmare Chemie NV", "SAMPLE TEMPLATE, NOT FOR USE"):
        assert want in t, f"missing on the label: {want}"
    assert not any("{{" in x for x in t)
    assert pdf.exists() and pdf.stat().st_size > 0


def test_french_label_uses_translation_file_and_decimal_comma():
    docx, _, _ = run(lang="FR")
    t = texts_of(docx)
    for want in ("Identificateur de produit", "Lot", "Date de fabrication", "Date de péremption", "Poids net", "25,5 kg"):
        assert want in t, f"missing on the FR label: {want}"
    assert "Batch" not in t and "Net weight" not in t


def test_empty_value_leaves_its_row_out():
    docx, _, _ = run(dict(FIELDS, expiryDate="", grossWeight=""))
    t = texts_of(docx)
    assert "Expiry date" not in t and "Gross weight" not in t
    assert "Production date" in t


def test_no_language_without_texts():
    # Arabic follows when TMC supplies its texts; never translated here
    refused(lambda: run(lang="AR"), "no label texts for language AR")


def test_missing_translation_stops_the_run():
    texts = load_texts()
    del texts["texts"]["batch"]["fr"]
    refused(lambda: run(lang="FR", texts=texts), "no translation for 'batch'")


def test_bad_input_refused():
    refused(lambda: run(dict(FIELDS, supplier="ACME")), "unknown field")
    refused(lambda: run(dict(FIELDS, productionDate="01/03/2026")), "not a date")
    refused(lambda: run(dict(FIELDS, netWeight="25 kg")), "not a number")
    refused(lambda: run(dict(FIELDS, batch="")), "batch is required")


def test_provenance_gate_blocks_invented_text():
    real_fill = make_label.fill

    def fill_and_invent(template, values, out_docx):
        real_fill(template, values, out_docx)
        from docx import Document
        d = Document(str(out_docx))
        d.add_paragraph("Manufactured by Acme Chemicals GmbH")
        d.save(str(out_docx))

    make_label.fill = fill_and_invent
    try:
        out = Path(tempfile.mkdtemp())
        refused(lambda: make(TEMPLATE, dict(FIELDS), ENTITY, "EN", "TMC_LABEL_x_EN", out), "HARD RULE VIOLATED")
        assert not list(out.iterdir()), "output kept after a provenance failure"
    finally:
        make_label.fill = real_fill


def test_command_line_refuses_outside_template_and_cleans_up():
    with tempfile.TemporaryDirectory() as d:
        d = Path(d)
        out = d / "out"
        out.mkdir()
        (d / "f.json").write_text(json.dumps(FIELDS), encoding="utf-8")
        args = [f"--fields={d / 'f.json'}", f"--entity={ENGINE / 'entity.example.json'}", "--lang=EN",
                "--name=TMC_LABEL_x_EN", f"--output-dir={out}"]
        assert make_label.main(["--template=../entity.example.json", *args]) == 1
        assert make_label.main(["--template=TMC_COA_template.docx", *args]) == 1
        assert make_label.main(["--template=TMC_LABEL_SAMPLE.docx", *args[:3], "--name=../evil", args[4]]) == 1
        assert not list(out.iterdir())
        assert make_label.main(["--template=TMC_LABEL_SAMPLE.docx", *args]) == 0
        assert sorted(p.name for p in out.iterdir()) == ["TMC_LABEL_x_EN.docx", "TMC_LABEL_x_EN.pdf"]


def custom_template(*paragraphs):
    """A copy of the SAMPLE template with extra paragraphs; each paragraph is a list of runs."""
    from docx import Document
    d = Document(str(TEMPLATE))
    for runs in paragraphs:
        p = d.add_paragraph()
        for r in runs:
            p.add_run(r)
    path = Path(tempfile.mkdtemp()) / "TMC_LABEL_TEST.docx"
    d.save(str(path))
    return path


def test_placeholder_split_over_runs_is_filled():
    # Word often splits {{batch}} over several runs after an edit
    template = custom_template(["Lot no. {{ba", "tch}} / {{product", "Name}}"])
    out = Path(tempfile.mkdtemp())
    docx, _ = make(template, dict(FIELDS), ENTITY, "EN", "TMC_LABEL_split_EN", out)
    t = texts_of(docx)
    assert "Lot no. BT-3877 / Betaine AI test" in t, t
    assert not any("{{" in x or "}}" in x for x in t)


def test_unknown_template_placeholder_stops_the_run():
    template = custom_template(["Supplier: {{supplierName}}"])
    out = Path(tempfile.mkdtemp())
    refused(lambda: make(template, dict(FIELDS), ENTITY, "EN", "TMC_LABEL_x_EN", out), "unknown value: supplierName")
    assert not list(out.iterdir()), "output kept after an unknown placeholder"


def test_typed_placeholder_is_not_expanded():
    # defence in depth: the web form refuses {{ }}, the engine must not fill it in either
    out = Path(tempfile.mkdtemp())
    refused(lambda: make(TEMPLATE, dict(FIELDS, batch="{{entity.naam}}"), ENTITY, "EN", "TMC_LABEL_x_EN", out),
            "placeholder left")
    assert not list(out.iterdir()), "output kept after a typed placeholder"


def test_unicode_and_special_characters_in_values_print_as_typed():
    docx, _, _ = run(dict(FIELDS, productName="Bétaïne <AI> & \"test\" 25%"))
    assert "Bétaïne <AI> & \"test\" 25%" in texts_of(docx)


if __name__ == "__main__":
    failed = 0
    for name, fn in list(globals().items()):
        if name.startswith("test_") and callable(fn):
            try:
                fn()
                print(f"ok  {name}")
            except Exception as err:  # noqa: BLE001
                failed += 1
                print(f"XX  {name}: {err}")
    print("labels:", "FAIL" if failed else "ok")
    sys.exit(1 if failed else 0)
